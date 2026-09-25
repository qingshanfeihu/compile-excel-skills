# 生成：tools/extract_engine.py ← InfoTest main/case_compiler/excel_capability_samples.py（sha256 83ccf0bd58978391）。不在这里手改。
from __future__ import annotations
from cex_core.engine._root import _cex_data_path
import ast
import io
import json
import os
import re
import secrets
import string
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping
from openpyxl import load_workbook
from cex_core.engine.case_compiler._sealed_io import canonical_json as _sealed_canonical_json, directory_flags, open_directory_nofollow, read_regular_nofollow as _sealed_read_regular_nofollow, sha256_bytes as _sha256, validate_json_budget, validate_xlsx_zip_budget
from cex_core.engine.case_compiler.excel_contract import validate_g_arguments_for_entry, ValidatedExcelContract, resolve_execution_sheet, validate_excel_contract, validate_g_for_entry, parse_g_arguments
SCHEMA = 'ist.excel.capability-sample-set'
_ROOT = _cex_data_path('')
_RUNTIME_ROOT = _ROOT / 'runtime'
_AUTOID_RE = re.compile('^[0-9]{18}$')
_MAX_CONTRACT_BYTES = 16 * 1024 * 1024
_MAX_TEMPLATE_BYTES = 128 * 1024 * 1024
_CHECKPOINT_DISPATCHES = {'checkpoint_two_argument', 'checkpoint_v2_three_argument'}
_PURE_ACTION_INPUTS = {'action.func_208': '00:00:00.000\n00:00:01.000', 'action.func_220': '2026 Jan 1 00:00:00 Threshold\n2026 Jan 1 00:00:01 Threshold', 'action.func_300': 'ether 00:11:22:33:44:55', 'action.func_223': 'Session ID: 0a1b2c3d4e5f', 'client_action.func_207': '\nRS=RS231'}
_SSL_RELEASE_SYMBOLS = frozenset({'ssl_comm.importRootCA', 'ssl_comm.importKey', 'ssl_comm.importCert', 'ssl_comm.activeCert'})
_SSL_RELEASE_KEY = 'cert/epolicy_ssl/rsaca/1024rsa.key'
_SSL_RELEASE_CERT = 'cert/epolicy_ssl/rsaca/1024rsa.crt'
_SSL_RELEASE_ROOT = 'cert/epolicy_ssl/rsarootca.crt'
_HA_RELEASE_SYMBOL = 'ha_comm.ha_default'
_HA_RELEASE_STATUS_PATTERN = 'Group\\s+1\\s+(?:Active\\s+Standby|Standby\\s+Active)'

def certified_ssl_release_fixture() -> dict[str, Any]:
    return {'schema': 'ist.ssl-release-fixture', 'authority': 'promoted_release_replay', 'cert_group': 'epolicy_ssl', 'rootca_file': 'rsarootca.crt', 'pairs': [{'key_file': 'rsaca/1024rsa.key', 'cert_file': 'rsaca/1024rsa.crt', 'index': 1}], 'resolved_paths': {'rootca': _SSL_RELEASE_ROOT, 'keys': [_SSL_RELEASE_KEY], 'certs': [_SSL_RELEASE_CERT]}}
_NON_RECOVERABLE_SAMPLE_TEXT = re.compile('(?i)(?:clear\\s+conf(?:ig)?\\s+all|system\\s+reboot|\\b(?:ip|ip\\s+-6)\\s+route\\s+(?:add|del|delete)\\b|\\bip\\s+addr\\s+(?:add|del|delete)\\b|\\bwrite\\s+(?:memory|file)\\b)')

class ExcelCapabilitySampleError(RuntimeError):
    pass

@dataclass(frozen=True)
class CapabilitySampleArtifacts:
    workbook_path: Path
    manifest_path: Path
    workbook_sha256: str
    manifest_file_sha256: str
    manifest_canonical_sha256: str
    row_count: int
    enabled_entry_count: int
    enabled_action_binding_count: int

@dataclass(frozen=True)
class _SampleRow:
    description: str
    e: str
    f: str
    g: str
    h: str = ''
    i: str = ''
    kind: str = 'entry'
    python_symbol: str = ''
    dispatcher: str = ''
    normalized: str = ''
    canonical: str = ''
    provenance_note: str = ''

def _canonical_json(payload: Mapping[str, Any], *, omit: str | None=None) -> bytes:
    return _sealed_canonical_json(payload, ensure_ascii=False, omit=omit)

def _read_regular_nofollow(path: Path, *, max_bytes: int, label: str) -> bytes:
    return _sealed_read_regular_nofollow(path, error_type=ExcelCapabilitySampleError, invalid_message=f'{label} path is invalid', directory_message='sample output root contains an unavailable or symlink directory', open_message=f'{label} cannot be opened safely', bounds_message=f'{label} must be a bounded single-link regular file', changed_message=f'{label} changed while being read', max_bytes=max_bytes, min_bytes=1)

def _directory_flags() -> int:
    return directory_flags(error_type=ExcelCapabilitySampleError, unavailable_message='sample output requires O_NOFOLLOW')

def _open_directory_path(path: Path) -> int:
    return open_directory_nofollow(path, error_type=ExcelCapabilitySampleError, invalid_message='sample output directory is invalid', unavailable_message='sample output root contains an unavailable or symlink directory')

def _open_output_directory(parent: Path, trusted_root: Path) -> int:
    root = Path(os.path.abspath(os.fspath(trusted_root)))
    selected = Path(os.path.abspath(os.fspath(parent)))
    try:
        relative = selected.relative_to(root)
    except ValueError as exc:
        raise ExcelCapabilitySampleError('sample outputs must stay inside the project runtime directory') from exc
    descriptor = _open_directory_path(root)
    try:
        for part in relative.parts:
            try:
                os.mkdir(part, 448, dir_fd=descriptor)
            except FileExistsError:
                pass
            next_descriptor = os.open(part, _directory_flags(), dir_fd=descriptor)
            os.close(descriptor)
            descriptor = next_descriptor
        return descriptor
    except OSError as exc:
        os.close(descriptor)
        raise ExcelCapabilitySampleError('sample output directory cannot be created safely') from exc

def _write_all(descriptor: int, payload: bytes) -> None:
    offset = 0
    while offset < len(payload):
        written = os.write(descriptor, payload[offset:])
        if written <= 0:
            raise OSError('short write')
        offset += written

def _write_new_pair(workbook_path: Path, workbook_bytes: bytes, manifest_path: Path, manifest_bytes: bytes, *, trusted_root: Path) -> None:
    if workbook_path.parent != manifest_path.parent:
        raise ExcelCapabilitySampleError('sample workbook and manifest must share one transaction directory')
    if workbook_path.name == manifest_path.name:
        raise ExcelCapabilitySampleError('sample workbook and manifest must use different filenames')
    for path in (workbook_path, manifest_path):
        if path.name in {'', '.', '..'} or '/' in path.name or '\\' in path.name:
            raise ExcelCapabilitySampleError('sample output filename is invalid')
    directory_fd = _open_output_directory(workbook_path.parent, trusted_root)
    token = secrets.token_hex(12)
    temporary_names = {workbook_path.name: f'.{workbook_path.name}.{token}.tmp', manifest_path.name: f'.{manifest_path.name}.{token}.tmp'}
    published: list[str] = []
    try:
        for name in (workbook_path.name, manifest_path.name):
            try:
                os.stat(name, dir_fd=directory_fd, follow_symlinks=False)
            except FileNotFoundError:
                continue
            raise ExcelCapabilitySampleError('sample output already exists; use a fresh evidence directory')
        for name, payload in ((workbook_path.name, workbook_bytes), (manifest_path.name, manifest_bytes)):
            temporary = temporary_names[name]
            descriptor = os.open(temporary, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW | getattr(os, 'O_CLOEXEC', 0), 384, dir_fd=directory_fd)
            try:
                _write_all(descriptor, payload)
                os.fsync(descriptor)
            finally:
                os.close(descriptor)
        for name in (workbook_path.name, manifest_path.name):
            temporary = temporary_names[name]
            os.link(temporary, name, src_dir_fd=directory_fd, dst_dir_fd=directory_fd, follow_symlinks=False)
            published.append(name)
            os.unlink(temporary, dir_fd=directory_fd)
        os.fsync(directory_fd)
    except Exception as exc:
        for temporary in temporary_names.values():
            try:
                os.unlink(temporary, dir_fd=directory_fd)
            except FileNotFoundError:
                pass
        if len(published) != 2:
            for name in published:
                try:
                    os.unlink(name, dir_fd=directory_fd)
                except FileNotFoundError:
                    pass
            os.fsync(directory_fd)
        if isinstance(exc, FileExistsError):
            raise ExcelCapabilitySampleError('sample output appeared during publication; nothing was overwritten') from exc
        raise
    finally:
        os.close(directory_fd)

def _octal_printf(text: str) -> str:
    encoded = ''.join((f'\\{ord(char):03o}' for char in text))
    return f"printf '{encoded}'"

def _shell_sample(label: str, *, repeats: int=1) -> str:
    token = f'IST_{label}_SAFE'
    return _octal_printf((token + '\n') * repeats)

def _placeholder_printf(prefix: str) -> str:
    encoded = ''.join((f'\\{ord(char):03o}' for char in prefix))
    return f"printf '{encoded}{{}}\\012'"

def _entry_key(entry: Mapping[str, Any]) -> tuple[str, str, str]:
    return (str(entry['e']), str(entry['f']), str(entry['python_symbol']))

def _seated_apv_es() -> frozenset[str]:
    """当前床拓扑在席的 APV 席位（与 scripts.gen_capability_atlas._bed_apv_seats
    同源同口径：读 network_topology.json 的 APV\\d 命名设备；口径漂移由
    test_seated_scope_matches_generator 锁）。

    认证样本只给在席席位构建序列；拓扑不可读时 fail 硬——fail-open 会把
    缺席席位的行放进认证卷，设备上必挂。
    """
    topology = _cex_data_path('') / 'knowledge/data/auto_env/network_topology.json'
    try:
        payload = json.loads(topology.read_text(encoding='utf-8'))
    except (OSError, ValueError) as exc:
        raise ExcelCapabilitySampleError('bed topology is unreadable; recoverable certification candidates cannot be seated safely') from exc
    seats: set[str] = set()
    for device in payload.get('devices', []) or []:
        if not isinstance(device, dict):
            continue
        match = re.fullmatch('APV(\\d+)', str(device.get('name') or ''))
        if match:
            seats.add(f'APV_{match.group(1)}')
    return frozenset(seats)

def _recoverable_release_entries(contract: Mapping[str, Any], exclude_es: frozenset[str]=frozenset()) -> list[dict[str, Any]]:
    """release 序列闭集（SSL 四装载 + HA 默认）里因无收据而禁、且床位在席
    的候选条目——认证样本与覆盖闭集的同一取面，两处共用防止口径漂移。"""
    seated = _seated_apv_es()
    return [dict(entry) for entry in contract['entries'] if _recoverable_certification_candidate(entry) and entry['e'] not in exclude_es and (entry['dispatch'] not in _CHECKPOINT_DISPATCHES) and (entry['dispatch'] != 'execute_registry') and re.fullmatch('APV_\\d+', str(entry['e'])) and (str(entry['e']) in seated) and (str(entry.get('python_symbol') or '') in _SSL_RELEASE_SYMBOLS or str(entry.get('python_symbol') or '') == _HA_RELEASE_SYMBOL)]

def _recoverable_certification_candidate(entry: Mapping[str, Any]) -> bool:
    """「因无收据而禁、语义上可认证」的条目：认证样本必须覆盖它。

    席位被床事实禁用的条目（authority=bed_topology:*）不在此列——认证
    通过了也无法晋升（席位缺席是环境事实不是认证缺口），强行进样本只会
    白烧一轮设备并让 release 卷无谓失败。2026-09-22 事故回归点：
    ssl_comm/ha_comm 掉回 disabled 后样本构建只看 enabled，重认证通道
    自锁，认证永久丢失。
    """
    return entry.get('status') == 'disabled' and entry.get('certification_candidate') is True and str(entry.get('status_authority') or '').startswith('project_policy:')

def _action_key(action: Mapping[str, Any], e_value: str) -> tuple[str, ...]:
    return (e_value, str(action['dispatcher']), str(action['normalized']), str(action['canonical']), str(action['python_symbol']))

def _placeholder_fields(raw_value: Any) -> list[str]:
    try:
        parsed = list(string.Formatter().parse(str(raw_value)))
    except ValueError as exc:
        raise ExcelCapabilitySampleError('sample G placeholder braces are not balanced') from exc
    fields: list[str] = []
    for _literal, field_name, format_spec, conversion in parsed:
        if field_name is None:
            continue
        if field_name not in {'', '0'} or format_spec or conversion:
            raise ExcelCapabilitySampleError('sample G allows only an exact {} or {0} placeholder')
        fields.append(field_name)
    if '' in fields and len(fields) > 1:
        raise ExcelCapabilitySampleError('sample G automatic placeholder may appear only once and may not mix with {0}')
    return fields

def _validate_row_placeholders(row: _SampleRow) -> None:
    if row.e == 'check_point':
        return
    args, kwargs = parse_g_arguments(row.g, row.f)
    has_input = bool(str(row.i).strip())
    if not has_input:
        static_values = [*args, *kwargs.values()]
        if any((re.search('(?<!\\{)\\{(?:0)?\\}(?!\\})', str(value or '')) for value in static_values)):
            raise ExcelCapabilitySampleError('sample G injection placeholder requires a non-empty I register')
        return
    positional_fields = [_placeholder_fields(value) for value in args]
    keyword_fields = [_placeholder_fields(value) for value in kwargs.values()]
    if any(positional_fields[1:]) or any(keyword_fields):
        raise ExcelCapabilitySampleError('sample G placeholder is allowed only in the first positional argument')
    first = positional_fields[0] if positional_fields else []
    if has_input and (not args or not first):
        raise ExcelCapabilitySampleError("sample I register requires a placeholder in G's first positional argument")

def _entry_sample(entry: Mapping[str, Any], *, index: int, apv_read_command: str='') -> _SampleRow:
    e_value = str(entry['e'])
    f_value = str(entry['f'])
    dispatch = str(entry['dispatch'])
    symbol = str(entry['python_symbol'])
    register = f'sample_{index:03d}'
    note = 'Author fixed a no-state-change input before execution to prove dispatch.'
    if dispatch == 'cmd_primitive':
        if f_value == 'cmd':
            g_value = _shell_sample(f'{e_value}_CMD')
        elif f_value in {'cmd_enable', 'cmd_config'}:
            if not apv_read_command:
                raise ExcelCapabilitySampleError('build-bound read-only APV sample command is unavailable')
            g_value = apv_read_command
        elif f_value == 'cmds_config':
            if not apv_read_command:
                raise ExcelCapabilitySampleError('build-bound read-only APV sample command is unavailable')
            g_value = f'{apv_read_command}\n{apv_read_command}'
        else:
            raise ExcelCapabilitySampleError(f'no safe cmd_primitive sample for {e_value}::{f_value}')
        return _SampleRow(f'安全调用 {e_value}.{f_value}', e_value, f_value, g_value, h=register, python_symbol=symbol, provenance_note=note)
    if dispatch == 'direct_method_call' and e_value == 'test_env':
        if f_value == 'server213':
            g_value = _shell_sample('FT_TOKEN_7Q9', repeats=2)
            register = ''
        elif f_value == 'clientc':
            g_value = _placeholder_printf('IST_I_PLACEHOLDER_')
            register = f'sample_{index:03d}'
            return _SampleRow(f'安全调用 {e_value}.{f_value} 并验证 I 注入', e_value, f_value, g_value, h=register, i='sample_scalar', python_symbol=symbol, provenance_note=note)
        else:
            g_value = _shell_sample(f'{e_value}_{f_value}')
        return _SampleRow(f'安全调用 {e_value}.{f_value}', e_value, f_value, g_value, h=register, python_symbol=symbol, provenance_note=note)
    if dispatch == 'builtin_special_case' and e_value == 'time' and (f_value == 'sleep'):
        return _SampleRow('安全调用 time.sleep', e_value, f_value, '1', h=register, python_symbol=symbol, provenance_note='Author chose a bounded one-second duration before execution.')
    raise ExcelCapabilitySampleError(f'no independently safe release sample for enabled {e_value}::{f_value} ({symbol})')

def _checkpoint_rows(entries: list[Mapping[str, Any]]) -> list[_SampleRow]:
    by_f = {str(entry['f']): entry for entry in entries}
    expected = {'found', 'abs_found', 'not_found', 'found_times'}
    if set(by_f) != expected:
        raise ExcelCapabilitySampleError('enabled checkpoint set differs from the certified four-method sample')
    token = 'IST_FT_TOKEN_7Q9_SAFE'
    note = 'Author encoded the fixed token twice in the preceding printf command before execution; the expected count is not copied from device output.'
    rows = [_SampleRow('正则包含断言', 'check_point', 'found', token, python_symbol=str(by_f['found']['python_symbol']), provenance_note=note), _SampleRow('字面包含断言', 'check_point', 'abs_found', token, python_symbol=str(by_f['abs_found']['python_symbol']), provenance_note=note), _SampleRow('正则不包含断言', 'check_point', 'not_found', 'IST_NEVER_PRESENT_4F2', python_symbol=str(by_f['not_found']['python_symbol']), provenance_note=note), _SampleRow('出现次数断言', 'check_point', 'found_times', token, i='2', python_symbol=str(by_f['found_times']['python_symbol']), provenance_note=note)]
    return rows

def _build_rows(contract: ValidatedExcelContract, exclude_es: frozenset[str]=frozenset(), *, device_build: str='') -> tuple[list[_SampleRow], set[tuple[str, ...]], set[tuple[str, ...]]]:
    enabled_entries = [entry for entry in contract['entries'] if entry['status'] == 'enabled' and entry['e'] not in exclude_es]
    checkpoint_entries = [entry for entry in enabled_entries if entry['dispatch'] in _CHECKPOINT_DISPATCHES]
    rows: list[_SampleRow] = []
    entry_coverage: set[tuple[str, ...]] = set()
    action_coverage: set[tuple[str, ...]] = set()
    direct_entries = [entry for entry in enabled_entries if entry['dispatch'] not in _CHECKPOINT_DISPATCHES and entry['dispatch'] != 'execute_registry']
    recoverable_entries = _recoverable_release_entries(contract, exclude_es)
    ssl_entries = [entry for entry in direct_entries + recoverable_entries if str(entry.get('python_symbol') or '') in _SSL_RELEASE_SYMBOLS and (str(entry.get('status_authority') or '').startswith('device_receipt_set:') or _recoverable_certification_candidate(entry))]
    ssl_entry_ids = {id(entry) for entry in ssl_entries}
    ha_entries = [entry for entry in direct_entries + recoverable_entries if str(entry.get('python_symbol') or '') == _HA_RELEASE_SYMBOL and (str(entry.get('status_authority') or '').startswith('device_receipt_set:') or _recoverable_certification_candidate(entry))]
    ha_entry_ids = {id(entry) for entry in ha_entries}
    early_entries = [entry for entry in direct_entries if entry['e'] != 'test_env' and id(entry) not in ssl_entry_ids and (id(entry) not in ha_entry_ids)]
    deferred_entries = [entry for entry in direct_entries if entry['e'] == 'test_env']
    next_index = 1
    apv_read_command = ''
    if any((entry['dispatch'] == 'cmd_primitive' and str(entry['e']).startswith('APV') and (entry['f'] in {'cmd_enable', 'cmd_config', 'cmds_config'}) for entry in enabled_entries)):
        from cex_core.engine.ist_core.tools.device.emit_xlsx_tool import build_command_tree_resolver
        _vendor, _inventory, resolve = build_command_tree_resolver(device_build=device_build)
        verdict = resolve('show version')
        if verdict.get('kind') != 'hit' or verdict.get('parameters_valid') is not True:
            raise ExcelCapabilitySampleError('the build-bound XML tree does not certify the release read probe')
        apv_read_command = 'show version'
    ssl_by_e: dict[str, dict[str, Mapping[str, Any]]] = {}
    for entry in ssl_entries:
        ssl_by_e.setdefault(str(entry['e']), {})[str(entry['python_symbol'])] = entry
    entry_by_ef = {(str(entry['e']), str(entry['f'])): entry for entry in enabled_entries}
    for seat_index, (e_value, methods) in enumerate(sorted(ssl_by_e.items())):
        required_ssl = {'ssl_comm.importKey', 'ssl_comm.importCert', 'ssl_comm.activeCert'}
        if not required_ssl <= set(methods):
            raise ExcelCapabilitySampleError(f'SSL release sample for {e_value} lacks a certified method')
        cmd_entry = entry_by_ef.get((e_value, 'cmd_config'))
        found_entry = entry_by_ef.get(('check_point', 'found'))
        if cmd_entry is None or found_entry is None:
            raise ExcelCapabilitySampleError('SSL release sample requires enabled cmd_config and found')
        vhost = f'ist_release_cert_{seat_index}'
        note = 'The exact method order, certificate files and object-scoped interactive cleanup were signed by the current same-source certification receipts.'
        sequence: list[tuple[_SampleRow, Mapping[str, Any]]] = [(_SampleRow(f'创建 {e_value} 隔离 SSL 主机', e_value, 'cmd_config', f'ssl host virtual "{vhost}"', python_symbol=str(cmd_entry['python_symbol']), provenance_note=note), cmd_entry)]
        if 'ssl_comm.importRootCA' in methods:
            sequence.append((_SampleRow(f'{e_value} 导入认证 RSA 根证书', e_value, 'importRootCA', f'{vhost},{_SSL_RELEASE_ROOT}', python_symbol='ssl_comm.importRootCA', provenance_note=note), methods['ssl_comm.importRootCA']))
        sequence.extend([(_SampleRow(f'{e_value} 导入认证 RSA 私钥', e_value, 'importKey', f'{vhost},{_SSL_RELEASE_KEY}', python_symbol='ssl_comm.importKey', provenance_note=note), methods['ssl_comm.importKey']), (_SampleRow(f'{e_value} 导入认证 RSA 证书', e_value, 'importCert', f'{vhost},{_SSL_RELEASE_CERT}', python_symbol='ssl_comm.importCert', provenance_note=note), methods['ssl_comm.importCert']), (_SampleRow(f'{e_value} 激活认证证书', e_value, 'activeCert', vhost, python_symbol='ssl_comm.activeCert', provenance_note=note), methods['ssl_comm.activeCert']), (_SampleRow(f'读取 {e_value} 证书效果', e_value, 'cmd_config', f'show ssl certificate {vhost}', python_symbol=str(cmd_entry['python_symbol']), provenance_note=note), cmd_entry), (_SampleRow(f'{e_value} 证书 Issuer 可见', 'check_point', 'found', 'Issuer', python_symbol=str(found_entry['python_symbol']), provenance_note=note), found_entry), (_SampleRow(f'开始清理 {e_value} 隔离 SSL 主机', e_value, 'cmd_config', f'clear ssl host "{vhost}",prompt=abort:', python_symbol=str(cmd_entry['python_symbol']), provenance_note=note), cmd_entry), (_SampleRow(f'确认清理 {e_value} 隔离 SSL 主机', e_value, 'cmd_config', 'YES', python_symbol=str(cmd_entry['python_symbol']), provenance_note=note), cmd_entry)])
        for row, entry in sequence:
            rows.append(row)
            entry_coverage.add(_entry_key(entry))
    if ha_entries:
        ha_by_e = {str(entry['e']): entry for entry in ha_entries}
        cmd_entries = {e_value: entry_by_ef.get((e_value, 'cmd_config')) for e_value in sorted(ha_by_e)}
        found_entry = entry_by_ef.get(('check_point', 'found'))
        sleep_entry = entry_by_ef.get(('time', 'sleep'))
        if len(ha_by_e) < 2 or any((entry is None for entry in cmd_entries.values())) or found_entry is None or (sleep_entry is None):
            raise ExcelCapabilitySampleError('HA release sample requires two seated HA entries plus cmd_config, found and time.sleep')
        note = 'The manual declares that one enabled group across two HA nodes resolves to Active/Standby roles; the device run decides the observed ordering.'
        for e_value, entry in sorted(ha_by_e.items()):
            rows.append(_SampleRow(f'在 {e_value} 调用已认证 HA 默认配置', e_value, 'ha_default', '', python_symbol=str(entry['python_symbol']), provenance_note=note))
            entry_coverage.add(_entry_key(entry))
        rows.append(_SampleRow('等待双节点 HA 状态收敛', 'time', 'sleep', '30', python_symbol=str(sleep_entry['python_symbol']), provenance_note='Author chose a bounded wait before reading HA state.'))
        entry_coverage.add(_entry_key(sleep_entry))
        for e_value, cmd_entry in sorted(cmd_entries.items()):
            assert cmd_entry is not None
            rows.extend([_SampleRow(f'读取 {e_value} HA 分组状态', e_value, 'cmd_config', 'show ha status', python_symbol=str(cmd_entry['python_symbol']), provenance_note=note), _SampleRow(f'{e_value} 呈现一主一备角色', 'check_point', 'found', _HA_RELEASE_STATUS_PATTERN, python_symbol=str(found_entry['python_symbol']), provenance_note=note)])
            entry_coverage.add(_entry_key(cmd_entry))
            entry_coverage.add(_entry_key(found_entry))
    for entry in sorted(early_entries, key=lambda item: (item['e'], item['f'])):
        index = next_index
        next_index += 1
        row = _entry_sample(entry, index=index, apv_read_command=apv_read_command)
        rows.append(row)
        entry_coverage.add(_entry_key(entry))
    enabled_actions = [action for action in contract['execute_actions'] if action['status'] == 'enabled']
    action_index = 0
    for action in sorted(enabled_actions, key=lambda item: (item['dispatcher'], item['canonical'], item['python_symbol'])):
        symbol = str(action['python_symbol'])
        if symbol not in _PURE_ACTION_INPUTS:
            raise ExcelCapabilitySampleError(f'enabled execute action has no audited pure sample: {symbol}')
        for e_value in sorted(action['allowed_es']):
            if e_value in exclude_es:
                continue
            action_index += 1
            entry = next((item for item in enabled_entries if item['e'] == e_value and item['f'] == 'execute' and (item['dispatch'] == 'execute_registry')), None)
            if entry is None:
                declared = any((item['e'] == e_value and item['f'] == 'execute' and (item['dispatch'] == 'execute_registry') for item in contract['entries']))
                if declared:
                    continue
                raise ExcelCapabilitySampleError(f'enabled action has no enabled execute entry for {e_value}')
            save_as = 'sample_scalar' if symbol == 'action.func_220' and e_value == 'APV_0' else f'action_{action_index:03d}'
            row = _SampleRow(f"纯计算动作 {action['canonical']}（{e_value}）", e_value, 'execute', f"{action['name']}：{_PURE_ACTION_INPUTS[symbol]}", h=save_as, kind='execute_action', python_symbol=symbol, dispatcher=str(action['dispatcher']), normalized=str(action['normalized']), canonical=str(action['canonical']), provenance_note='Author supplied fixed synthetic text for a mirror-proven pure transformation; no device state is read or changed.')
            rows.append(row)
            entry_coverage.add(_entry_key(entry))
            action_coverage.add(_action_key(action, e_value))
    for entry in sorted(deferred_entries, key=lambda item: (item['f'] == 'server213', item['f'])):
        index = next_index
        next_index += 1
        row = _entry_sample(entry, index=index, apv_read_command=apv_read_command)
        rows.append(row)
        entry_coverage.add(_entry_key(entry))
    rows.extend(_checkpoint_rows(checkpoint_entries))
    entry_coverage.update((_entry_key(entry) for entry in checkpoint_entries))
    return (rows, entry_coverage, action_coverage)

def _validate_rows(rows: list[_SampleRow], contract: ValidatedExcelContract, entry_coverage: set[tuple[str, ...]], action_coverage: set[tuple[str, ...]], exclude_es: frozenset[str]=frozenset(), device_build: str='') -> dict[str, Any]:
    recoverable_keys = {_entry_key(entry) for entry in _recoverable_release_entries(contract, exclude_es)}
    entries = {_entry_key(entry): entry for entry in contract['entries'] if entry['e'] not in exclude_es and (entry['status'] == 'enabled' or _entry_key(entry) in recoverable_keys)}
    required_entries = set(entries)
    if entry_coverage != required_entries:
        raise ExcelCapabilitySampleError('sample entry coverage is incomplete or contains an undeclared capability')
    enabled_execute_es = {entry['e'] for entry in contract['entries'] if entry['status'] == 'enabled' and entry['dispatch'] == 'execute_registry'}
    required_actions = {_action_key(action, e_value) for action in contract['execute_actions'] if action['status'] == 'enabled' for e_value in action['allowed_es'] if e_value not in exclude_es and e_value in enabled_execute_es}
    if action_coverage != required_actions:
        raise ExcelCapabilitySampleError('sample execute-action coverage is incomplete')
    entry_by_ef = {(entry['e'], entry['f']): entry for entry in entries.values()}
    registers: set[str] = set()
    has_result = False
    for row in rows:
        if _NON_RECOVERABLE_SAMPLE_TEXT.search(row.g):
            raise ExcelCapabilitySampleError(f'sample contains a non-recoverable state-changing command: {row.e}::{row.f}')
        entry = entry_by_ef.get((row.e, row.f))
        if entry is None:
            raise ExcelCapabilitySampleError(f'sample row targets a disabled or unknown entry: {row.e}::{row.f}')
        if _recoverable_certification_candidate(entry):
            validate_g_arguments_for_entry(entry, row.g)
        else:
            validate_g_for_entry(entry, row.g, contract)
        if row.e == 'check_point':
            if row.f == 'found_times':
                if row.h or row.i != '2':
                    raise ExcelCapabilitySampleError('found_times sample must use legacy G expected/H blank/I count semantics')
                if not has_result:
                    raise ExcelCapabilitySampleError('found_times sample requires a preceding observation result')
            elif not row.i and (not has_result):
                raise ExcelCapabilitySampleError('checkpoint sample requires a preceding observation result')
            continue
        _validate_row_placeholders(row)
        if row.i and row.i not in registers:
            raise ExcelCapabilitySampleError(f'sample I register is not available before use: {row.i}')
        if row.h:
            if not re.fullmatch('[A-Za-z_]\\w*', row.h):
                raise ExcelCapabilitySampleError(f'sample H register name is invalid: {row.h!r}')
            registers.add(row.h)
        else:
            has_result = True
    return _validate_net_zero_state(rows, contract=contract, device_build=device_build)

def _literal_direct_method_commands(row: _SampleRow, contract: Mapping[str, Any]) -> list[str]:
    if not row.e.startswith('APV') or not row.python_symbol:
        return []
    entry = next((item for item in contract.get('entries', []) if str(item.get('e') or '') == row.e and str(item.get('f') or '') == row.f and (str(item.get('python_symbol') or '') == row.python_symbol)), None)
    if not isinstance(entry, Mapping) or entry.get('dispatch') != 'direct_method_call':
        return []
    source = entry.get('source')
    source_path = str(source.get('path') or '') if isinstance(source, Mapping) else ''
    try:
        owner, method = row.python_symbol.split('.', 1)
    except ValueError:
        return []
    from cex_core.engine.case_compiler.apv_lang import mirror_src
    text = mirror_src(source_path)
    if not text:
        return []
    try:
        tree = ast.parse(text)
    except SyntaxError:
        return []
    function = next((child for node in tree.body if isinstance(node, ast.ClassDef) and node.name == owner for child in node.body if isinstance(child, (ast.FunctionDef, ast.AsyncFunctionDef)) and child.name == method), None)
    if function is None:
        return []
    commands = []
    for call in ast.walk(function):
        if not isinstance(call, ast.Call) or not isinstance(call.func, ast.Attribute) or call.func.attr != 'cmd_config' or (not call.args) or (not isinstance(call.args[0], ast.Constant)) or (not isinstance(call.args[0].value, str)):
            continue
        command = call.args[0].value.strip()
        if command and command.upper() not in {'YES', 'NO'}:
            commands.append(command)
    return commands

def _validate_net_zero_state(rows: list[_SampleRow], *, contract: Mapping[str, Any] | None=None, device_build: str='') -> dict[str, Any]:
    commands: list[str] = []
    for index, row in enumerate(rows):
        if contract is not None:
            commands.extend(_literal_direct_method_commands(row, contract))
        if not row.e.startswith('APV') or row.f not in {'cmd_config', 'cmds_config', 'cmd_enable'}:
            continue
        if row.f == 'cmd_config':
            try:
                args, _kwargs = parse_g_arguments(row.g, row.f)
                values = [str(args[0])] if args else []
            except Exception:
                values = [str(row.g or '')]
            if values and values[0].strip().upper() in {'YES', 'NO'} and (index > 0):
                previous = rows[index - 1]
                if previous.e == row.e and previous.f == 'cmd_config':
                    try:
                        _args, kwargs = parse_g_arguments(previous.g, previous.f)
                    except Exception:
                        kwargs = {}
                    if str(kwargs.get('prompt') or '').strip():
                        values = []
        else:
            values = str(row.g or '').splitlines()
        commands.extend((line.strip() for line in values if line.strip()))
    forward_candidates = [command for command in commands if not command.lower().startswith(('show ', 'no ', 'clear '))]
    if not forward_candidates:
        return {'schema': 'ist.excel-capability-net-zero', 'status': 'no_state_change', 'device_build': str(device_build or ''), 'covered': []}
    if not str(device_build or '').strip():
        raise ExcelCapabilitySampleError('state-changing capability samples require a build-bound teardown atlas identity')
    from cex_core.engine.case_compiler.tau_coverage import TauAtlasUnavailableError, check_tau_coverage_lines
    try:
        report = check_tau_coverage_lines(commands, device_build=device_build)
    except TauAtlasUnavailableError as exc:
        raise ExcelCapabilitySampleError('state-changing capability samples cannot load the build-bound teardown atlas') from exc
    if report.out_of_scope:
        raise ExcelCapabilitySampleError('capability sample contains persistent state outside the teardown atlas recovery surface')
    if report.residual_config:
        heads = ', '.join(sorted({str(item.get('head') or '') for item in report.residual_config}))
        raise ExcelCapabilitySampleError(f'capability sample contains configuration with no mechanical teardown: {heads}')
    if report.missing:
        heads = ', '.join(sorted({str(item.get('head') or '') for item in report.missing}))
        raise ExcelCapabilitySampleError(f'capability sample does not return device state to zero; missing teardown: {heads}')
    classified = {str(item.get('cmd') or '') for item in report.covered if str(item.get('cmd') or '')}
    unknown = [command for command in forward_candidates if command not in classified]
    if unknown:
        raise ExcelCapabilitySampleError('capability sample contains a state-changing command outside the generated teardown-atlas closed set')
    covered = [{'head': str(item.get('head') or ''), 'class': str(item.get('class') or ''), 'coverage': str(item.get('coverage') or 'in_case_teardown')} for item in report.covered if item.get('head')]
    return {'schema': 'ist.excel-capability-net-zero', 'status': 'covered', 'device_build': str(report.device_build or device_build), 'atlas_identity': dict(report.atlas_identity or {}), 'covered': covered}

def _manifest_sample(row: _SampleRow, *, row_number: int, contract: ValidatedExcelContract) -> dict[str, Any]:
    source_sha = ''
    if row.kind == 'entry':
        target = next((item for item in contract['entries'] if item['e'] == row.e and item['f'] == row.f))
    else:
        target = next((item for item in contract['execute_actions'] if item['dispatcher'] == row.dispatcher and item['normalized'] == row.normalized and (item['python_symbol'] == row.python_symbol) and (row.e in item['allowed_es'])))
    source_sha = contract['source_hashes'][target['source']['path']]
    identity = {'kind': row.kind, 'e': row.e, 'f': row.f, 'python_symbol': row.python_symbol, 'dispatcher': row.dispatcher or None, 'normalized': row.normalized or None, 'canonical': row.canonical or None, 'row': row_number}
    sample_id = 'sample-' + _sha256(_canonical_json(identity))[:20]
    return {**identity, 'sample_id': sample_id, 'g_sha256': _sha256(row.g.encode('utf-8')), 'source_file_sha256': source_sha, 'provenance': {'expect_source': 'Author', 'note': row.provenance_note}}

def build_capability_sample_artifacts(*, contract_path: str | Path, template_path: str | Path, workbook_path: str | Path, manifest_path: str | Path, autoid: str='990000000000000001', trusted_runtime_root: str | Path=_RUNTIME_ROOT, exclude_es: tuple[str, ...]=(), device_build: str='') -> CapabilitySampleArtifacts:
    if not _AUTOID_RE.fullmatch(str(autoid)):
        raise ExcelCapabilitySampleError('sample autoid must be an 18-digit value')
    contract_file = Path(contract_path)
    template_file = Path(template_path)
    contract_bytes = _read_regular_nofollow(contract_file, max_bytes=_MAX_CONTRACT_BYTES, label='Excel contract')
    try:
        validate_json_budget(contract_bytes, error_type=ExcelCapabilitySampleError, message='Excel contract exceeds the JSON structure budget')
        contract_payload = json.loads(contract_bytes.decode('utf-8'))
    except (UnicodeError, json.JSONDecodeError, RecursionError) as exc:
        raise ExcelCapabilitySampleError('Excel contract is not valid JSON') from exc
    if not isinstance(contract_payload, Mapping):
        raise ExcelCapabilitySampleError('Excel contract is not a JSON object')
    contract = validate_excel_contract(contract_payload)
    template_bytes = _read_regular_nofollow(template_file, max_bytes=_MAX_TEMPLATE_BYTES, label='runtime template')
    validate_xlsx_zip_budget(template_bytes, error_type=ExcelCapabilitySampleError, message='runtime template exceeds the XLSX expansion budget')
    excluded = frozenset((str(item) for item in exclude_es if str(item).strip()))
    rows, entry_coverage, action_coverage = _build_rows(contract, exclude_es=excluded, device_build=device_build)
    net_zero_receipt = _validate_rows(rows, contract, entry_coverage, action_coverage, exclude_es=excluded, device_build=device_build)
    workbook = load_workbook(io.BytesIO(template_bytes), data_only=False)
    try:
        worksheet, layout = resolve_execution_sheet(workbook, allow_legacy=False, contract=contract)
        if workbook.worksheets != [worksheet]:
            raise ExcelCapabilitySampleError('runtime sample template must contain exactly one execution worksheet')
        for row in worksheet.iter_rows(min_row=layout.header_row, max_col=worksheet.max_column):
            for cell in row:
                if cell.column > layout.n_cols and cell.value is not None:
                    raise ExcelCapabilitySampleError('runtime sample template execution data area must contain only the A-I execution columns')
        gap_column = layout.n_cols + 1
        for row in worksheet.iter_rows(max_col=worksheet.max_column):
            for cell in row:
                if cell.column == gap_column and cell.value is not None:
                    raise ExcelCapabilitySampleError('runtime sample template must contain only the A-I execution columns')
        for row_number in range(layout.data_start, worksheet.max_row + 1):
            for column in range(1, layout.n_cols + 1):
                worksheet.cell(row=row_number, column=column).value = None
        samples: list[dict[str, Any]] = []
        for offset, row in enumerate(rows):
            row_number = layout.data_start + offset
            values = [autoid if offset == 0 else None, 'P1' if offset == 0 else None, offset + 2, row.description, row.e, row.f, row.g, row.h or None, row.i or None]
            for column, value in enumerate(values, start=1):
                worksheet.cell(row=row_number, column=column, value=value)
            samples.append(_manifest_sample(row, row_number=row_number, contract=contract))
        output = io.BytesIO()
        workbook.save(output)
        workbook_bytes = output.getvalue()
    finally:
        workbook.close()
    validate_xlsx_zip_budget(workbook_bytes, error_type=ExcelCapabilitySampleError, message='generated sample workbook exceeds the XLSX expansion budget')
    verified_workbook = load_workbook(io.BytesIO(workbook_bytes), data_only=False)
    try:
        verified_sheet, verified_layout = resolve_execution_sheet(verified_workbook, allow_legacy=False, contract=contract)
        verified_sheet_title = verified_sheet.title
        verified_gap_column = verified_layout.n_cols + 1
        verified_data_area_clean = all((cell.value is None for row in verified_sheet.iter_rows(min_row=verified_layout.header_row, max_col=verified_sheet.max_column) for cell in row if cell.column > verified_layout.n_cols))
        verified_gap_clean = all((cell.value is None for row in verified_sheet.iter_rows(max_col=verified_sheet.max_column) for cell in row if cell.column == verified_gap_column))
        if verified_workbook.worksheets != [verified_sheet] or not verified_data_area_clean or (not verified_gap_clean) or (verified_layout.header_row != layout.header_row) or (verified_layout.data_start != layout.data_start):
            raise ExcelCapabilitySampleError('saved sample workbook changed the execution-sheet identity')
    finally:
        verified_workbook.close()
    manifest: dict[str, Any] = {'schema': SCHEMA, 'contract_sha256': contract['contract_sha256'], 'contract_file_sha256': _sha256(contract_bytes), 'runtime_version': contract['runtime']['minimum_version'], 'template_sha256': _sha256(template_bytes), 'workbook_sha256': _sha256(workbook_bytes), 'autoid': autoid, 'worksheet': verified_sheet_title, 'header_row': layout.header_row, 'data_start': layout.data_start, 'row_count': len(rows), 'enabled_entry_count': len(entry_coverage), 'enabled_action_binding_count': len(action_coverage), 'excluded_es': sorted(excluded), 'net_zero_state': net_zero_receipt, 'samples': samples}
    manifest['manifest_sha256'] = _sha256(_canonical_json(manifest, omit='manifest_sha256'))
    manifest_bytes = json.dumps(manifest, ensure_ascii=False, sort_keys=True, indent=2).encode('utf-8') + b'\n'
    manifest_file_sha256 = _sha256(manifest_bytes)
    selected_workbook = Path(workbook_path)
    selected_manifest = Path(manifest_path)
    _write_new_pair(selected_workbook, workbook_bytes, selected_manifest, manifest_bytes, trusted_root=Path(trusted_runtime_root))
    return CapabilitySampleArtifacts(workbook_path=selected_workbook, manifest_path=selected_manifest, workbook_sha256=manifest['workbook_sha256'], manifest_file_sha256=manifest_file_sha256, manifest_canonical_sha256=manifest['manifest_sha256'], row_count=len(rows), enabled_entry_count=len(entry_coverage), enabled_action_binding_count=len(action_coverage))
__all__ = ['SCHEMA', 'ExcelCapabilitySampleError', 'CapabilitySampleArtifacts', 'build_capability_sample_artifacts']
