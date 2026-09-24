# 生成：tools/extract_engine.py ← InfoTest main/ist_core/tools/device/structural_gate.py（sha256 8a4f86783bc25e17）。不在这里手改。
from __future__ import annotations
import logging
import re
import threading
from collections import OrderedDict
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from cex_core.engine.case_compiler.observe_ops import config_existence_check, observe_kind
from cex_core.engine.case_compiler.case_ir import IInjectionSyntaxError, parse_found_times_cells, validate_i_injection_syntax
from cex_core.engine.case_compiler.excel_contract import ExcelContractError, contract_entry, load_excel_contract, resolve_execution_sheet, validate_execute_action, validate_g_for_entry
logger = logging.getLogger(__name__)
_OP_PREFIXES = ('no', 'show', 'clear')
_VALUE_TOK_RE = re.compile('^([<\\[{].*|.*[|].*|\\d.*|[\\"\'].*|.*\\.\\d+.*)$')
_OBSERVE_RE = re.compile('\\b(show|statistics|stat|display|dig|nslookup|get|list)\\b', re.IGNORECASE)
from cex_core.engine.case_compiler.apv_lang import APV_ACTION_SRC, CLIENT_ACTION_SRC, _PARAM_SPLIT_RE, mirror_src, nearest_candidates

def _candidate_signature_hint(e: str, name: str, contract: dict) -> str:
    entry = contract_entry(e, name, contract)
    if entry is None:
        return ''
    sig = entry.get('signature') or {}
    req = ', '.join(sig.get('required') or []) or '(none)'
    opt = ', '.join(sig.get('optional') or []) or '(none)'
    return f'{name}(required=[{req}], optional=[{opt}])'

def _load_contract_for_gate(result: StructuralResult | None=None):
    try:
        return load_excel_contract()
    except ExcelContractError as exc:
        if result is not None and (not any((item.code == 'excel_contract_unavailable' for item in result.violations))):
            result.add('excel_contract_unavailable', f"Excel function contract is unavailable or stale: {exc}; no E/F/G row in any case can be certified. This is NOT an authoring defect and no edit to this case's steps can clear it — the contract is a generated projection and has to be regenerated engine-side from the mirror. Do not reshape the case around this error; report the blocked state.")
        return None

def _contract_host_slot_es(contract: dict) -> frozenset[str]:
    return frozenset((str(item['e']) for item in contract['objects'] if item['status'] == 'enabled' and item['python_type'] == 'ssh_server'))

@lru_cache(maxsize=4)
def _execute_returning_actions(src_rel: str, source_sha256: str='') -> frozenset:
    src = mirror_src(src_rel)
    if not src:
        return frozenset()
    mapping = dict(re.findall("'([^']+)':\\s*self\\.(func_\\d+)", src))
    returning: set[str] = set()
    for m in re.finditer('\\n    def (func_\\d+)\\(self[^)]*\\):(.*?)(?=\\n    def |\\Z)', src, re.S):
        if re.search('\\n\\s+return\\s+\\S', m.group(2)):
            returning.add(m.group(1))
    return frozenset((name for name, fn in mapping.items() if fn in returning))

@dataclass
class StructuralViolation:
    code: str
    detail: str
    step_index: int = -1

@dataclass
class StructuralResult:
    ok: bool = True
    violations: list[StructuralViolation] = field(default_factory=list)
    disabled: list[StructuralViolation] = field(default_factory=list)
    advisories: list[StructuralViolation] = field(default_factory=list)

    def add(self, code: str, detail: str, step_index: int=-1) -> None:
        self.ok = False
        self.violations.append(StructuralViolation(code, detail, step_index))

    def disable(self, code: str, detail: str, step_index: int=-1) -> None:
        from cex_core.engine.case_compiler.gate_advisories import STRUCTURAL_DISABLED_CODES
        if code not in STRUCTURAL_DISABLED_CODES:
            raise ValueError(f'unregistered structural disabled code: {code}')
        self.disabled.append(StructuralViolation(code, detail, step_index))

    def advise(self, code: str, detail: str, step_index: int=-1) -> None:
        from cex_core.engine.case_compiler.gate_advisories import STRUCTURAL_ADVISORY_CODES
        if code not in STRUCTURAL_ADVISORY_CODES:
            raise ValueError(f'unregistered structural advisory: {code}')
        self.advisories.append(StructuralViolation(code, detail, step_index))

    def render(self, autoid: str) -> str:
        lines = [f'case {autoid} violates structural constraints (correct-by-construction gate, independent of grade):']
        for v in self.violations:
            loc = f'step[{v.step_index}] ' if v.step_index >= 0 else ''
            lines.append(f'  - [{v.code}] {loc}{v.detail}')
        lines.append('\nThese are intent-independent **structural** errors (command legality / whether an assertion is dangling) — deterministically decidable and they must be fixed; they are not skeleton-choice issues. Fix them and emit again.')
        return '\n'.join(lines)
_NO_PATH_LESSON_TAIL = 'On the real bed such a query returns \'network unreachable\' or an empty reply and every assertion consuming it is bound to fail, so the case is falsified at compile time instead of burning a device round. Move the query to an executor that has a declared leg in the target\'s segment (the topology supply lists the pairing) — the exemption for an address the Author wrote covers the value, never the choice of executor, so this code still fires while another executor shares that segment. When the Author\'s own text names this target and no bed executor has a declared leg in its segment, the blocker is a bed prerequisite rather than a case defect: do not force-emit, report it with compile_report_underdetermined(reason_code="environment_prerequisite_gap") and state the missing prerequisite — intent is not mechanically visible on this surface.'
GATE_LESSON_TEXT: dict[str, str] = {'found_times_invalid_shape': 'found_times takes its three arguments from the row itself: G contains the static expected regex, H stays blank, I is a positive integer occurrence count, and the previous unregistered result is the actual text. Fix G/H/I instead of replacing the assertion with a weaker operator.', 'excel_contract_unavailable': 'the generated Excel function contract is missing, stale, or incomplete, so no E/F/G row in any case can be certified. This is NOT an authoring defect and nothing in the steps can fix it — rewriting E/F/G, switching dispatchers, or retrying the same case all hit the same wall. The contract is a generated projection; it has to be regenerated engine-side from the mirror before compilation can continue. Report the blocked state instead of reshaping the case around the error.', 'disabled_dispatch_method': 'the E/F pair exists in the generated contract but is disabled for this runtime; do not emit it until its implementation and runtime version are enabled together.', 'invalid_function_arguments': 'G must bind to the exact generated Python signature. Fix missing or extra arguments, duplicate keywords, unbalanced quotes, or an invalid zero-argument call.', 'dangling_assertion': "an assertion with I empty reads the framework's last result buffer — that buffer is only set by the nearest preceding step WITHOUT H that is itself an observation (dig/show), not a config step and not a step that saved its output with H. Put a plain (no-H) observation step immediately before the assertion.", 'manual_ip_cleanup': "do not add/delete IPs or routes on test_env or host-slot steps — the framework books every add and restores by delete at the next case's start; manual del or repeated add across cases collides with that restore and crashes or pollutes later cases' echoes.", 'empty_command_payload': 'a command step\'s G must not be None/empty-ish — the framework sends it verbatim and the device rejects a literal "None" with ^.', 'cmd_config_multiline': "cmd_config's G must be a single command — multi-line G gets flattened into one glued line and the device rejects it. Use cmds_config for multiple commands (sent line by line).", 'literal_backslash_n': 'a literal two-character backslash-n in G is a mis-escaped newline, not an intended token — use a real newline to separate multiple commands.', 'unknown_dispatch_target': "E must be a name in the framework's devices table — an unrecognized E makes the framework silently skip the whole step (no error, no log), and later assertions read the wrong buffer.", 'unknown_dispatch_method': 'F must be a valid method of its E object — an unrecognized method raises AttributeError and crashes the whole file.', 'execute_action_not_in_registry': "an execute step's action name must be an exact member of the framework's action registry — an inexact name falls back to fuzzy matching and can silently dispatch to a different (even opposite) action, invisibly on the device side.", 'execute_payload_required': 'this execute action requires a non-empty payload after the full-width separator; sending only the action name reaches its implementation with missing input and crashes the workbook run.', 'line_anchor_never_matches': 'the static regex contains a string-boundary anchor incompatible with a mandatory consuming prefix or suffix. No actual string can match that form. Boundary anchors themselves are legal: their meaning follows the selected actual channel and regex flags, including a register selected by I.', 'assertion_matches_command_echo': "an assertion pattern that also matches the command text of the step that produced the window it's checking verifies nothing (always-true fake pass for found/abs_found, always-fail for not_found) — match a data line, not the command echo.", 'no_assertion_in_case': 'a case needs at least one check_point step — the framework judges zero assertions as FAIL on settlement, so a config-only/observation-only sheet always fails.', 'empty_assertion_pattern': 'a check_point needs a non-empty pattern in G, or a captured register reference — leaving G/H/I all empty gives the framework nothing to compare (it falls back to searching with the command text itself).', 'mutation_control_after_config': 'an engine-issued IDE-MUTATION-BEFORE/CONTROL pair must sit immediately before its state-changing configuration anchor. A cmd_config/cmds_config dispatcher label alone is insufficient because read-only show/get windows use the same dispatcher. If the pair lands after configuration or before a read-only window, its control can observe state created by the same case; regenerate the compiler-owned rows rather than editing the assertion.', 'undefined_capture_ref': "a register referenced by H/I must be captured by an earlier step's H= first — an undefined reference reads back None (assertion distorted) or raises NameError (crashes the file, for non-assertion steps).", 'injection_format_crash': "a step's G that receives an I-injected value must be a plain single-slot format string — a named or multi-position placeholder raises KeyError/IndexError/ValueError on injection. A static G with blank I is not formatted; named braces in JSON or a downstream tool's literal syntax remain ordinary data.", 'injection_without_placeholder': 'a step carrying I needs a {} placeholder in G for the captured value to land — without one the injection silently does nothing and G runs as literal text.', 'injection_placeholder_scope': 'G is parsed into positional and keyword arguments before I injection; placeholders are allowed only in the first positional argument, and every placeholder requires a non-empty I reference. The two legal shapes are: (1) static value — G has no placeholder and I is blank; (2) injected value — an earlier step captures H=<name>, the current G first positional argument contains exactly one {} or {0}, and I=<name> (or a declared runtime-object attribute). Other positional/keyword arguments contain no placeholder.', 'register_shadows_framework_name': 'an H register name must not collide with a v2 runner **runtime object** name (a contract device slot, or a runner built-in). Registers live in their own dict in v2, so ordinary local-looking names are fine; only runtime objects are refused, and the refusal is hard: preflight raises NameError before any step runs and the whole pytest file crashes. Rename the register, do not rename the device.', 'destructive_command': 'the case sheet carries a line from the banned destructive-command set (grammar `destructive_commands.patterns`, the single source this rule and the emit rules both read). Two kinds live in that set and a case may carry neither: whole-device config wipes, which take the management IP down with everything else (two beds were killed this way on 2026-07-13 and needed console recovery), and device lifecycle commands, which really restart or power down a bed other runs are sharing. Framework capability is not the limit — the framework does ship a reboot helper that waits the device back up — but that helper is declared disabled in the Excel contract as a device-wide lifecycle helper, so no step can reach it: the ban is bed-sharing discipline, not a missing capability. Clean up only what the case itself created, with an object-scoped clear or the paired `no` form of what was configured. If the case objective genuinely needs the device to restart, it cannot be compiled for a shared bed — report it as not compilable instead of force-emitting.', 'command_not_in_tree': 'every E=APV_* cmd_config/cmds_config G line must be a product CLI command in the build-bound XML command tree. Shell/backend text and an ordinary absent command head are both terminal misses; use an equivalent CLI observation or report that the case cannot be compiled.', 'command_parameter_contract': 'the command head exists in the build-bound XML tree, but its argument count or argument shape does not match that exact XML signature. Re-read the projected parameter contract and correct the current command before sealing; do not infer an alternate command from this lesson or weaken the assertion.', 'governing_spec_identity_invalid': "an emit check reported a typed-expectation, compiled-form receipt, or governing-spec identity rejection. The gate name alone does not establish which condition failed or who caused it. Read the bound raw detail for the expectation_id and specific missing or mismatched identity. Keep the author's expected values and source claims unchanged; if the detail cannot establish a valid repair, disclose that gap instead of inventing an identity or substituting another source.", 'expectation_direction_rewrite': 'an expectation-direction seal check failed. The cause may be unreadable or malformed seal data, an identity mismatch, or a direction token set difference. The gate name alone establishes neither an expected rewrite nor an Author/model fault. Read the bound raw detail for the specific failed check; do not rewrite or replace an immutable seal. A token set difference is a representation comparison; it does not prove that the expected meaning changed or became its opposite. Preserve the signed source statement and author postconditions. If equivalence between the two representations is unverified, disclose that gap; do not infer a semantic flip, rewrite expected values, or split/merge expectation groups to bypass the seal.', 'driver_no_declared_path': "a query step executed on a test-driver host (router/client) may only target an IP the driver has a declared path to: sharing a declared segment with one of the driver's interfaces (IPv4 subnet / IPv6 prefix), or covered by an explicit route declaration in the topology. This target sits outside every declared path. " + _NO_PATH_LESSON_TAIL, 'executor_networks_undeclared': 'the same gate as driver_no_declared_path, hit from the other side: the executor is enabled in the contract, but the topology fact source declares NO network segment for it at all, so no path to the target can be verified either way. Treated as no-path. ' + _NO_PATH_LESSON_TAIL, 'cmd_not_in_allowlist': "the command's first-level module name is in no module of the verified-command footprint index. That index records command heads already observed to work, so a module absent from all of it is normally an out-of-scope or invented command. It is a heuristic: only the module root is judged (unknown sub-commands are logged, not refused), and the index is not the authority on whether a command exists — the build-bound XML command tree is (see command_not_in_tree). Check the module name against that tree; if the command really is in this build, say so rather than silently substituting a different command.", 'dead_capture': 'a step captured an echo into a register with H, but no check_point ever reads that register — neither as the H holding the expected value nor as the I naming the text under test. A register written and never read means the observation it captured was never verified, and the assertions in the case are checking some other buffer. Either consume the register with a check_point (capture-compare is the three-step form: observe with H, observe again without H, then assert with H), or drop the capture.', 'assertion_regex_invalid': 'the assertion pattern does not compile as a Python regex. The framework compiles it before matching, so the exception surfaces at re.compile and the whole pytest file crashes — every case after this one is skipped. Fix the pattern syntax; the usual cause is an unclosed character class.', 'short_mode_status_assertion': "the observation this assertion consumes was run in dig's short mode, whose output carries record values only — no header, status or section lines. Matching that kind of text against a short-mode window is always-fail; it never reaches the device to be decided. Either drop short mode from that observation so the full response is available, or rewrite the assertion against a record value.", 'comma_splits_parameters': "G for a test-env or host-slot step contains an unquoted comma, and the framework splits G on unquoted commas into positional parameters. The second segment onward lands on the host method's prompt/timeout parameters, so the step waits out the full timeout against a prompt that never appears and the captured output is silently truncated — no error is raised. Quote the segment if the comma belongs to the command, or use the framework's named timeout parameter form.", 'autoid_malformed': "the sheet's autoid is not an 18-digit number. A hand-copied or truncated id creates a junk output directory and can slip into the merged sheet as an extra case. Take the id verbatim from the machine-readable ledger (manifest / last_run), never by retyping it.", 'autoid_row_not_runnable': "an autoid row has an empty column E, so the deployed runner cannot start a case from it: the row is skipped before the new-case boundary, and the steps that follow execute inside the previous case — verdicts get attributed to the wrong case and device state carries across. The autoid must share its row with the case's first step; it must not sit on a row of its own.", 'xlsx_unreadable': "the workbook cannot be opened, or its execution sheet fails the identity check (marker / contract SHA / defined name). Nothing about the case's steps causes this and no step edit fixes it — the file is corrupt, or it was produced against a different runtime template than the one in force. Re-emit the case through the compiler rather than editing the workbook by hand."}

def _gate_lesson_sentence(code: str) -> str:
    """把批级黑板上那条静态教训接到案级事实后面。

    同一条规则有两张 LLM 可见面：批级黑板（`briefs._run_lessons` 读
    `GATE_LESSON_TEXT`，结构上够不到 case 数据）和案级拒绝（`StructuralResult
    .render` 的 detail）。两边各写一份时改一处漏一处，同一条规则会对同一个写手
    说两种话。教训只留黑板那一份，案级面只补这一次的事实。黑板面按 JSON 字段
    独立呈现、首字母小写，接进句子时补大写。
    """
    lesson = GATE_LESSON_TEXT.get(code, '')
    return lesson[:1].upper() + lesson[1:]

def _command_head_tokens(cmd: str) -> list[str]:
    toks = cmd.strip().split()
    while toks and toks[0].lower() in _OP_PREFIXES:
        toks = toks[1:]
    out: list[str] = []
    for t in toks:
        tl = t.strip().lower()
        if not tl:
            continue
        if not re.match('^[a-z][a-z0-9_-]*$', tl) or _VALUE_TOK_RE.match(tl):
            break
        out.append(tl)
    return out

def _load_allowlist_prefixes() -> tuple[set[str], set[str]]:
    full: set[str] = set()
    roots: set[str] = set()
    try:
        from cex_core.engine.ist_core.memory.footprint import get_footprint_index
        idx = get_footprint_index()
        for fid in idx.list_nodes():
            data = idx._nodes.get(fid) or {}
            behaviors = data.get('behaviors', [])
            if not data.get('cli', {}).get('commands') and (not data.get('decision_rules')) and behaviors and all((b.get('validity') == 'uncertain' for b in behaviors)):
                continue
            full.add(fid)
            roots.add(fid.split('.')[0])
    except Exception as exc:
        logger.debug('structural_gate 读 footprint 失败: %s', exc)
    return (full, roots)

def _check_command_allowlist(steps: list, init: str, result: StructuralResult) -> None:
    full, roots = _load_allowlist_prefixes()
    if not roots:
        result.disable('command_allowlist_footprint_unavailable', 'footprint 命令前缀集为空,本次未做命令头越界检查(判据缺供给,非案的问题)')
        return

    def _check_one(cmd: str, idx: int) -> None:
        head = _command_head_tokens(cmd)
        if not head:
            return
        module = head[0]
        if module not in roots:
            result.add('cmd_not_in_allowlist', f"command {cmd!r}: its first-level module {module!r} is in no module of the verified-command footprint index (modules on record: {', '.join(sorted(roots))}) — normally an out-of-scope or invented command. This gate judges the module root only (unknown sub-commands are logged, not refused) and the footprint index is NOT the authority on whether a command exists — the build-bound XML command tree is. Check the module name against that tree; if this command really does exist in this build, report that instead of quietly swapping in a different command.", idx)
    from cex_core.engine.ist_core.tools.device.emit_xlsx_tool import _ordered_apv_command_refs
    for ref in _ordered_apv_command_refs(steps, init):
        _check_one(str(ref.get('command') or ''), int(ref.get('step_index', -1)))

def _check_command_existence_same_source(steps: list, init: str, result: StructuralResult, *, device_build: str='') -> None:
    from cex_core.engine.ist_core.compile_engine.engine_errors import CommandTreeUnavailable
    from cex_core.engine.ist_core.tools.device.emit_xlsx_tool import COMMAND_PARAMETER_CONTRACT_CODE, COMMAND_NOT_IN_TREE_CODE, _command_tree_context, _looks_like_shell, _ordered_apv_command_refs, command_existence_verdict
    command_refs = _ordered_apv_command_refs(steps, init)
    if not command_refs:
        return
    loaded_context = _command_tree_context(device_build=device_build)
    for command_ref in command_refs:
        command = command_ref['command']
        verdict = command_existence_verdict(command, device_build=device_build, _loaded_context=loaded_context)
        if verdict['kind'] == 'tree_unavailable':
            raise CommandTreeUnavailable(verdict['detail'], device_build=verdict['device_build'] or device_build)
        if verdict['kind'] == 'hit' and (not verdict['parameters_valid']):
            result.add(COMMAND_PARAMETER_CONTRACT_CODE, f"command {command!r}: its head exists in build {verdict['device_build'] or '?'} but the supplied arguments violate the build-bound XML parameter contract; head={verdict['head']!r}, parameter_error={verdict['parameter_error']!r}. Re-read the exact XML-backed signature and correct this command before sealing.", command_ref['step_index'])
            continue
        if verdict['kind'] != 'missing':
            continue
        if _looks_like_shell(command):
            detail = f'command {command!r} is not a CLI command in this build. E=APV_* G may contain only build-bound product CLI commands (DESIGN §26.7); use an equivalent CLI observation or report that the case cannot be compiled.'
        else:
            missing_head = str(verdict.get('head') or command)
            detail = f"command head {missing_head!r} is absent from build {verdict['device_build'] or '?'} XML command tree" + (f' (full command {command!r})' if missing_head != command else '') + '. E=APV_* G may contain only commands from that tree (DESIGN §2.2/§26.7).'
        result.add(COMMAND_NOT_IN_TREE_CODE, detail, command_ref['step_index'])

def _is_observation_step(step: dict) -> bool:
    e = str(step.get('E', '')).strip()
    f = str(step.get('F', '')).strip()
    try:
        contract = load_excel_contract()
    except ExcelContractError:
        return False
    entry = contract_entry(e, f.lower() if e == 'test_env' else f, contract)
    if entry is None or entry['status'] != 'enabled':
        return False
    if e == 'test_env':
        return True
    if f in {'cmd', 'cmd_config'}:
        return True
    if f != 'execute':
        return False
    g = str(step.get('G', '')).strip()
    try:
        action_contract = validate_execute_action(e, g, contract)
    except ExcelContractError:
        return False
    src = CLIENT_ACTION_SRC if action_contract['dispatcher'] == 'client' else APV_ACTION_SRC
    return action_contract['name'] in _execute_returning_actions(src, str(contract['source_hashes'].get(src) or ''))

def _check_dangling_assertions(steps: list, result: StructuralResult) -> None:
    result_is_observe = False
    observe_bound_to: tuple[int, str] | None = None
    consecutive_observes: list[tuple[int, str]] = []
    for i, s in enumerate(steps):
        if not isinstance(s, dict):
            continue
        e = str(s.get('E', '')).strip()
        f = str(s.get('F', '')).strip()
        h = str(s.get('H', '') or '').strip()
        i_col = str(s.get('I', '') or '').strip()
        if e == 'check_point':
            reads_result = f == 'found_times' or not i_col
            if reads_result and result_is_observe and (len(consecutive_observes) >= 2):
                earlier = '、'.join((f'step[{n}] {c[:30]!r}' for n, c in consecutive_observes[:-1]))
                bound_n, bound_c = consecutive_observes[-1]
                result.advise('ambiguous_observation_binding', f'this assertion consumes only the echo of step[{bound_n}] ({bound_c[:40]!r}); the earlier consecutive observation(s) {earlier} produced echoes nothing reads — either dead observations, or the assertion was meant to read one of them (mis-binding). Put each assertion immediately after the observation it verifies, or capture earlier echoes with H and reference them explicitly.', i)
            if reads_result:
                consecutive_observes = []
            if reads_result and (not result_is_observe):
                result.add('dangling_assertion', 'this assertion reads the framework result but no valid observation echo is held there → framework result=None, found(None) raises TypeError and **crashes the entire file** (no case after this one runs). Cause: the nearest preceding step without H is a config step (cmds_config returns None), or every earlier observation step carries H (save_as does not update result). Fix: put an observation step **without H** (dig/show) **immediately before** the assertion so its echo becomes result; for capture-compare use the three-step form dig(H=v1) → dig(no H) → check_point(H=v1).', i)
        elif not h:
            if _is_observation_step(s):
                result_is_observe = True
                observe_bound_to = (i, str(s.get('G', '') or ''))
                consecutive_observes.append(observe_bound_to)
            else:
                result_is_observe = False
                observe_bound_to = None
                consecutive_observes = []

def _check_dead_capture(steps: list, result: StructuralResult) -> None:
    captured: dict[str, int] = {}
    referenced: set[str] = set()
    for i, s in enumerate(steps):
        if not isinstance(s, dict):
            continue
        e = str(s.get('E', '')).strip()
        h = str(s.get('H', '') or '').strip()
        i_col = str(s.get('I', '') or '').strip()
        if i_col and (not (e == 'check_point' and str(s.get('F', '')).strip() == 'found_times')):
            referenced.add(i_col)
        if e == 'check_point':
            if h:
                referenced.add(h)
        elif h:
            captured[h] = i
    for reg, idx in captured.items():
        if reg not in referenced:
            result.add('dead_capture', f"register '{reg}' is captured via save_as but is **never referenced by any check_point** (neither as the H for expect nor as the I for the text under test) = a register written but never read — a dead action / incomplete assertion. Typical: dig captures into '{reg}' with H yet the check_point reads result via found (the immediately preceding step is often a show config echo → asserting against the wrong buffer), with no abs_found(H='{reg}') consuming it → the dig behavior is never verified and the assertion reads the wrong buffer. Fix: reference '{reg}' with check_point abs_found (three-step form dig(H={reg})→dig(no H)→check_point(H={reg})), or delete the useless capture.", idx)

def check_structural_constraints(autoid: str, steps: list, init: str='') -> StructuralResult:
    result = StructuralResult()
    if not isinstance(steps, list):
        return result
    _check_command_allowlist(steps, init, result)
    _check_dangling_assertions(steps, result)
    _check_found_times_shape(steps, result)
    _check_dispatch_targets(steps, result)
    _check_contract_g_syntax(steps, result)
    _check_dead_capture(steps, result)
    return result

def _check_found_times_shape(steps: list, result: StructuralResult) -> None:
    for i, s in enumerate(steps):
        if not isinstance(s, dict):
            continue
        if str(s.get('E', '')).strip() == 'check_point' and str(s.get('F', '')).strip() == 'found_times':
            try:
                parse_found_times_cells(s.get('G'), s.get('H'), s.get('I'))
            except ValueError as exc:
                result.add('found_times_invalid_shape', f'{exc}; G supplies the static expected regex, H stays blank, and the previous unregistered result is the actual text under test.', i)

def _destructive_patterns() -> tuple[list, str | None]:
    try:
        from cex_core.engine.case_compiler.domain_grammar import load_grammar
        import re as _re
        pats = (load_grammar().get('destructive_commands') or {}).get('patterns') or []
    except Exception as exc:
        import logging
        logging.getLogger(__name__).warning('destructive_commands 文法读取失败——自毁命令规则本次未检查', exc_info=True)
        return ([], f'grammar unreadable ({exc.__class__.__name__}: {exc})')
    if not pats:
        return ([], 'grammar carries no destructive_commands.patterns entries')
    try:
        return ([_re.compile(str(p), _re.IGNORECASE) for p in pats], None)
    except Exception as exc:
        import logging
        logging.getLogger(__name__).warning('destructive_commands 正则编译失败——自毁命令规则本次未检查', exc_info=True)
        return ([], f'pattern compile failed ({exc.__class__.__name__}: {exc})')

def _check_no_destructive_commands(steps: list, result: StructuralResult) -> None:
    from cex_core.engine.ist_core.tools.device.emit_xlsx_tool import _apv_command_lines_for_step, is_apv_command_step
    pats, unavailable = _destructive_patterns()
    if unavailable:
        result.disable('destructive_command', f'destructive-command gate did not run: {unavailable}. Commands that wipe the whole device configuration (management IP included) were NOT screened in this volume — the bed-killing class this rule exists for is unguarded here.')
        return
    for i, s in enumerate(steps or []):
        if not isinstance(s, dict):
            continue
        if not is_apv_command_step(s, methods=None):
            continue
        for line in _apv_command_lines_for_step(steps, i):
            line = line.strip()
            if not line:
                continue
            matched = next((p for p in pats if p.search(line)), None)
            if matched is not None:
                result.add('destructive_command', f'step {i + 1}: {line!r} matches banned destructive-command rule {matched.pattern!r} (grammar `destructive_commands.patterns`, the single source this rule and the emit rules both read). Two kinds live in that set and a case sheet may carry neither: whole-device config wipes, which take the management IP down with everything else (two beds were killed this way on 2026-07-13 and needed console recovery), and device lifecycle commands, which really restart or power down a bed other runs are sharing. Framework capability is not the limit — the framework does ship a reboot helper that waits the device back up — but that helper is declared disabled in the Excel contract as a device-wide lifecycle helper, so no step can reach it: the ban is bed-sharing discipline, not a missing capability. Clean up only what this case created, with an object-scoped clear or the paired `no` form of what it configured. If the case objective genuinely needs the device to restart, it cannot be compiled for a shared bed — report it as not compilable instead of force-emitting.')

def _check_no_manual_ip_cleanup(steps: list, result: StructuralResult) -> None:
    contract = _load_contract_for_gate(result)
    if contract is None:
        return
    host_slots = _contract_host_slot_es(contract)
    for i, s in enumerate(steps):
        if not isinstance(s, dict):
            continue
        e = str(s.get('E', '')).strip()
        g = ' '.join(str(s.get('G', '') or '').split())
        if (e == 'test_env' or e in host_slots) and ('ip addr ' in g or 'ip address ' in g or 'ip route ' in g or ('ip -6 route ' in g)) and (' add' in g or ' del' in g):
            result.add('manual_ip_cleanup', "this step changes test-env host IPs/routes (add/del via ip addr / ip route) — the framework auto-books every add (without dedup) and restores by delete at the start of the next case: deleting on your own, or repeating the same add across cases, makes that restore fail, **crashing the entire file or polluting later cases' echo with RTNETLINK residue (batches of fake fails)**. Delete this step — host network state is managed by the framework; for multiple sources use the topology's existing trigger hosts, and for a new path check the topology's existing reachability first.", i)

def _check_dispatch_targets(steps: list, result: StructuralResult) -> None:
    contract = _load_contract_for_gate(result)
    if contract is None:
        return
    object_names = {str(item['e']) for item in contract['objects']}
    for i, s in enumerate(steps):
        if not isinstance(s, dict):
            continue
        e = str(s.get('E', '')).strip()
        f = str(s.get('F', '')).strip()
        if not e:
            continue
        if e not in object_names:
            result.add('unknown_dispatch_target', f"E={e!r} is not in the framework devices table (valid: {', '.join(sorted(object_names))}) — for an unknown E the framework **silently skips the whole step** (not executed, no log, no exception), and later assertions run against the wrong buffer. Fix it against excel_contract.json.", i)
            continue
        f_norm = f.lower() if e == 'test_env' else f
        entry = contract_entry(e, f_norm, contract)
        if entry is None:
            allowed = {str(item['f']) for item in contract['entries'] if item['e'] == e and item['status'] == 'enabled'}
            ranked = nearest_candidates(f_norm, allowed, top_n=5)
            hinted = [_candidate_signature_hint(e, c, contract) or c for c in ranked]
            result.add('unknown_dispatch_method', f"F={f!r} is not a valid method of the E={e} object — the framework getattr has no default, so a misspelled method name raises **AttributeError and crashes the entire file**. Nearest candidates (ranked by similarity, for reference only — confirm deliberately, this rule never auto-rewrites; signature shown inline where known): {'、'.join(hinted)}. Full enabled set ({len(allowed)}): {', '.join(sorted(allowed))}. Fix it against the E/F entry in excel_contract.json.", i)
            continue
        if entry['status'] != 'enabled':
            result.add('disabled_dispatch_method', f"E={e!r}, F={f_norm!r} is declared but {entry['status']}: {entry['reason']} (minimum runtime {entry['minimum_runtime']}).", i)

def _check_contract_g_syntax(steps: list, result: StructuralResult) -> None:
    contract = _load_contract_for_gate(result)
    if contract is None:
        return
    for i, step in enumerate(steps):
        if not isinstance(step, dict):
            continue
        e = str(step.get('E') or '').strip()
        f = str(step.get('F') or '').strip()
        f_norm = f.lower() if e == 'test_env' else f
        entry = contract_entry(e, f_norm, contract)
        if entry is None or entry['status'] != 'enabled':
            continue
        if entry.get('dispatch') == 'execute_registry':
            continue
        g = str(step.get('G') or '')
        if f == 'cmd_config' and ('\n' in g or '\r' in g):
            continue
        try:
            validate_g_for_entry(entry, g, contract)
        except ExcelContractError as exc:
            result.add('invalid_function_arguments', f"E={e!r}, F={f_norm!r} cannot bind G to {entry['signature'].get('text')}: {exc}", i)

def _check_command_payload_sanity(steps: list, result: StructuralResult) -> None:
    contract = _load_contract_for_gate(result)
    if contract is None:
        return
    for i, s in enumerate(steps):
        if not isinstance(s, dict):
            continue
        e = str(s.get('E', '')).strip()
        f = str(s.get('F', '')).strip()
        entry = contract_entry(e, f.lower() if e == 'test_env' else f, contract)
        if entry is None or entry.get('dispatch') not in {'cmd_primitive', 'direct_method_call', 'execute_registry'}:
            continue
        g_raw = s.get('G')
        g = str(g_raw) if g_raw is not None else ''
        if g_raw is None or g.strip().lower() == 'none':
            result.add('empty_command_payload', 'command step column G is None/literal "None" — the framework str()-ifies it and sends it as-is; the device receives "None" and always rejects it with ^. Fill in a real command or delete the step (a pure empty-string placeholder step is harmless and not covered here).', i)
        elif f == 'cmd_config' and ('\n' in g.strip() or '\r' in g.strip()):
            result.add('cmd_config_multiline', "cmd_config's G contains newlines — the framework first replace()-strips all newlines for cmd_config （框架执行器出处已脱敏）, so multiple commands get **concatenated with no separator** and sent as one line; the device always rejects it with ^. Use cmds_config for multiple commands (sent line by line).", i)
        elif '\\n' in g:
            result.add('literal_backslash_n', 'command column G contains a **literal** \\\\n (backslash + n, two characters, not a newline) — multiple commands get joined into one line and sent; the device always rejects at the second command with ^. Separate multiple commands with **real newlines** (the \\n escape written in JSON is restored to a newline by parsing; if you wrote \\\\\\\\n in the string it became a literal backslash — fix it).', i)

def _check_execute_action_registry(steps: list, result: StructuralResult) -> None:
    contract = _load_contract_for_gate(result)
    if contract is None:
        return
    for i, s in enumerate(steps):
        if not isinstance(s, dict):
            continue
        if str(s.get('F', '')).strip() != 'execute':
            continue
        e_value = str(s.get('E', '')).strip()
        raw = str(s.get('G', '') or '')
        try:
            validate_execute_action(e_value, raw, contract)
        except ExcelContractError as exc:
            detail = str(exc)
            code = 'execute_payload_required' if 'requires a non-empty payload' in detail else 'execute_action_not_in_registry'
            result.add(code, f'{detail}. execute uses the whole G cell as one action argument and must match the E-bound generated contract exactly; this rule covers only F=execute action routing and never rewrites it. Direct Python methods belong in F and are declared by their E/F pair in excel_contract.json. Distribution and membership assertions use the compiler combinators OBSERVE_DIST and OBSERVE_MEMBER; neither is an execute action.', i)
_DEST_ANCHOR_RE = re.compile('(?:@|://)\\[?([0-9A-Fa-f:.]+)\\]?')

def _extract_query_destinations(g: str) -> list[str]:
    import ipaddress as _ip
    out: list[str] = []
    for m in _DEST_ANCHOR_RE.finditer(g or ''):
        tok = m.group(1).strip().strip('.')
        if tok.count(':') == 1 and '.' in tok:
            host, port = tok.rsplit(':', 1)
            if port.isdecimal():
                tok = host
        try:
            _ip.ip_address(tok.split('%')[0])
        except ValueError:
            continue
        if tok not in out:
            out.append(tok)
    return out

def _check_driver_no_declared_path(steps: list, result: StructuralResult, *, author_ip_literals: Sequence[str] | None=None, author_hits: list[dict] | None=None) -> None:
    from cex_core.engine.ist_core.tools._shared.env_facts import normalize_ip_literal, require_env_facts
    facts = require_env_facts()
    author_set = {normalize_ip_literal(value) for value in author_ip_literals or () if normalize_ip_literal(value)}
    contract = _load_contract_for_gate(None)
    te_hosts: frozenset[str] = frozenset()
    slot_hosts: frozenset[str] = frozenset()
    if contract is not None:
        te_hosts = frozenset((str(item.get('f', '')).strip().lower() for item in contract.get('entries', []) if item.get('e') == 'test_env' and item.get('status') == 'enabled'))
        slot_hosts = frozenset((s.lower() for s in _contract_host_slot_es(contract)))
    driver_rows: dict[int, tuple[str, list[str]]] = {}
    save_var: dict[int, str] = {}
    for i, s in enumerate(steps):
        if not isinstance(s, dict):
            continue
        e = str(s.get('E', '')).strip()
        f = str(s.get('F', '')).strip()
        if e == 'check_point':
            continue
        executor = ''
        if e == 'test_env' and f:
            fl = f.lower()
            if contract is not None and fl in te_hosts or (contract is None and facts.is_driver_device(fl)):
                executor = fl
        elif contract is not None and e.lower() in slot_hosts or (contract is None and facts.is_driver_device(e)):
            executor = e.lower()
        if not executor:
            continue
        dests = _extract_query_destinations(str(s.get('G', '') or ''))
        if not dests:
            continue
        driver_rows[i] = (executor, dests)
        h = str(s.get('H', '') or '').strip()
        if h:
            save_var[i] = h
    if not driver_rows:
        return
    consuming: dict[int, int] = {}
    last_no_h_idx = -1
    for i, s in enumerate(steps):
        if not isinstance(s, dict):
            continue
        e = str(s.get('E', '')).strip()
        f = str(s.get('F', '')).strip()
        h = str(s.get('H', '') or '').strip()
        i_col = str(s.get('I', '') or '').strip()
        if e == 'check_point':
            reads_result = f == 'found_times' or not i_col
            if reads_result and last_no_h_idx in driver_rows:
                consuming.setdefault(last_no_h_idx, i)
            for ref in (i_col, h):
                if not ref:
                    continue
                cands = [j for j, var in save_var.items() if var == ref and j < i]
                if cands:
                    consuming.setdefault(max(cands), i)
        elif not h:
            last_no_h_idx = i
    for j in sorted(set(consuming) & set(driver_rows)):
        executor, dests = driver_rows[j]
        for dest in dests:
            verdict = facts.executor_path_verdict(executor, dest)
            if verdict is None:
                continue
            segs = ', '.join(verdict['dest_segments']) or 'none — this IP is in no declared segment of the bed'
            peers = verdict['reachable_drivers']
            peer_note = "Test-drivers that do share the target's declared segment: " + ', '.join(peers) if peers else "No test-driver shares the target's declared segment — no driver can reach it at all"
            if author_set and (not peers) and (normalize_ip_literal(dest) in author_set):
                result.advise('driver_no_declared_path_author_sourced', f"this step queries {dest} from executor '{verdict['executor']}', and step[{consuming[j]}]'s assertion consumes that reply; no test-driver on this bed has a declared path to that target (target's declared segment(s): [{segs}]). The address is written verbatim in the sealed Author case, so this is an execution-environment disclosure, not a case rejection: the case proceeds and the limitation is reported; expected is not rewritten.", j)
                if author_hits is not None:
                    author_hits.append({'step_index': j, 'consuming_step_index': consuming[j], 'executor': str(verdict['executor']), 'target': str(dest), 'dest_segments': list(verdict['dest_segments']), 'executor_networks': list(verdict.get('executor_networks') or []), 'networks_undeclared': bool(verdict.get('networks_undeclared'))})
                continue
            if verdict.get('networks_undeclared'):
                result.add('executor_networks_undeclared', f"this step queries {dest} from executor '{verdict['executor']}', and step[{consuming[j]}]'s assertion consumes that reply. Target's declared segment(s) on this bed: [{segs}]. {peer_note} (derived from the topology fact source). " + _gate_lesson_sentence('executor_networks_undeclared'), j)
                continue
            nets = ', '.join(verdict['executor_networks']) or 'none declared'
            result.add('driver_no_declared_path', f"this step queries {dest} from executor '{verdict['executor']}', and step[{consuming[j]}]'s assertion consumes that reply. Executor's declared segments: [{nets}]; target's declared segment(s) on this bed: [{segs}]; no explicit route declaration covers it either. {peer_note} (derived from the topology fact source). " + _gate_lesson_sentence('driver_no_declared_path'), j)

def _check_mutation_control_order(steps: list, result: StructuralResult) -> None:
    from cex_core.engine.case_compiler.mutation_testing import mutation_control_order_failure
    failure = mutation_control_order_failure(steps)
    if failure:
        result.add('mutation_control_after_config', failure)

def check_crash_gates_mandatory(steps: list) -> StructuralResult:
    result = StructuralResult()
    if isinstance(steps, list):
        _check_found_times_shape(steps, result)
        _check_dangling_assertions(steps, result)
        _check_no_manual_ip_cleanup(steps, result)
        _check_command_payload_sanity(steps, result)
        _check_dispatch_targets(steps, result)
        _check_contract_g_syntax(steps, result)
        _check_execute_action_registry(steps, result)
        _check_assertion_regex_compiles(steps, result)
        _check_line_anchor_assertions(steps, result)
        _check_assertion_matches_command_echo(steps, result)
        _check_has_assertion(steps, result)
        _check_empty_assertion_pattern(steps, result)
        _check_capture_refs_defined(steps, result)
        _check_no_destructive_commands(steps, result)
        _check_mutation_control_order(steps, result)
    return result

def lint_draft(steps: list, *, init: str='', final: bool=False, device_build: str='', author_ip_literals: Sequence[str] | None=None, author_driver_hits: list[dict] | None=None) -> StructuralResult:
    result = check_crash_gates_mandatory(steps)
    if not final:
        kept: list[StructuralViolation] = []
        for v in result.violations:
            if v.code == 'no_assertion_in_case':
                result.advise(v.code, 'draft not finished yet — no check_point step written so far in this prefix; this is informational, not a violation (you may simply not have reached the assertion step). It becomes a real block once you call compile_lint(final=True) before emit, or reach compile_emit with the case still missing an assertion.', v.step_index)
            else:
                kept.append(v)
        result.violations = kept
        result.ok = not result.violations
    dc_probe = StructuralResult()
    _check_dead_capture(steps, dc_probe)
    for v in dc_probe.violations:
        result.advise(v.code, v.detail, v.step_index)
    _check_command_allowlist(steps, init, result)
    _check_command_existence_same_source(steps, init, result, device_build=device_build)
    _check_driver_no_declared_path(steps, result, author_ip_literals=author_ip_literals, author_hits=author_driver_hits)
    _check_config_existence_advisories(steps, result)
    return result
_AUTOID_RE = re.compile('^\\d{18}$')
_SHORT_INCOMPATIBLE_RE = re.compile('status:|->>HEADER<<-|ANSWER SECTION|QUESTION SECTION')
_WORKBOOK_MODEL_LOCK = threading.Lock()
_WORKBOOK_MODEL_CACHE: OrderedDict[tuple, tuple[str, tuple[dict, ...], tuple[str, ...]]] = OrderedDict()
_WORKBOOK_MODEL_CAP = 1536

def _parse_workbook_model(xlsx_path, *, allow_legacy: bool) -> tuple[str, tuple[dict, ...], tuple[str, ...]]:
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path)
    ws, layout = resolve_execution_sheet(wb, allow_legacy=allow_legacy)
    autoid = ''
    steps: list[dict] = []
    empty_e_autoids: list[str] = []
    for row in ws.iter_rows(min_row=layout.data_start, max_col=layout.n_cols):
        a = str(row[0].value or '').strip()
        e = str(row[4].value or '').strip()
        is_autoid_row = a.isdigit() and len(a) >= 12 and (a != '999999999999999')
        if is_autoid_row:
            autoid = autoid or a
            if not e:
                empty_e_autoids.append(a)
        if not e:
            continue
        steps.append({'D': str(row[3].value or ''), 'E': e, 'F': str(row[5].value or ''), 'G': str(row[6].value or ''), 'H': str(row[7].value or ''), 'I': str(row[8].value or '') if len(row) > 8 else ''})
    return (autoid, tuple(steps), tuple(empty_e_autoids))

def _workbook_model(xlsx_path, *, allow_legacy: bool=True) -> tuple[str, list[dict], list[str]]:
    from cex_core.engine.common.file_identity import file_identity
    try:
        identity = file_identity(Path(xlsx_path))
    except (TypeError, OSError):
        autoid, steps, empty_e_autoids = _parse_workbook_model(xlsx_path, allow_legacy=allow_legacy)
        return (autoid, [dict(step) for step in steps], list(empty_e_autoids))
    resolved, size, mtime_ns, digest = identity
    key = (resolved, size, mtime_ns, digest, bool(allow_legacy))
    with _WORKBOOK_MODEL_LOCK:
        hit = _WORKBOOK_MODEL_CACHE.get(key)
        if hit is not None:
            _WORKBOOK_MODEL_CACHE.move_to_end(key)
    if hit is None:
        model = _parse_workbook_model(xlsx_path, allow_legacy=allow_legacy)
        try:
            after = file_identity(Path(xlsx_path))
        except Exception:
            after = None
        if after == (resolved, size, mtime_ns, digest):
            with _WORKBOOK_MODEL_LOCK:
                _WORKBOOK_MODEL_CACHE[key] = model
                while len(_WORKBOOK_MODEL_CACHE) > _WORKBOOK_MODEL_CAP:
                    _WORKBOOK_MODEL_CACHE.popitem(last=False)
        hit = model
    autoid, steps, empty_e_autoids = hit
    return (autoid, [dict(step) for step in steps], list(empty_e_autoids))

def clear_workbook_model_cache() -> None:
    with _WORKBOOK_MODEL_LOCK:
        _WORKBOOK_MODEL_CACHE.clear()

def steps_from_xlsx(xlsx_path, *, allow_legacy: bool=True) -> tuple[str, list[dict]]:
    autoid, steps, _empty_e = _workbook_model(xlsx_path, allow_legacy=allow_legacy)
    return (autoid, steps)

def _check_line_anchor_assertions(steps: list, result: StructuralResult) -> None:
    from cex_core.engine.case_compiler.regex_anchor_proof import RegexAnchorAnalysisUnavailable, analyze_regex_anchors
    for i, s in enumerate(steps):
        if not isinstance(s, dict):
            continue
        if str(s.get('E', '')).strip() != 'check_point':
            continue
        f = str(s.get('F', '')).strip()
        g = str(s.get('G', '') or '')
        if f not in ('found', 'not_found') or not g or str(s.get('H') or '').strip():
            continue
        try:
            analysis = analyze_regex_anchors(g)
        except RegexAnchorAnalysisUnavailable as exc:
            result.disable('regex_anchor_analysis_unavailable', str(exc), i)
            continue
        if analysis.contradiction:
            result.add('line_anchor_never_matches', f'static assertion regex cannot match any actual: {analysis.contradiction}. ' + ('The not_found assertion therefore always passes.' if f == 'not_found' else 'The found assertion therefore always fails.'), i)
        elif analysis.window_boundary_dependent and (not str(s.get('I') or '').strip()):
            result.advise('line_anchor_window_unverified', 'this pattern anchors on the start/end boundary of the whole output window; whether the expected data sits on that boundary depends on the live echo layout, so no pass/fail was decided from it this time', i)

def _echo_src_cmd(e: str, f: str, host_slots: frozenset[str], contract=None) -> bool:
    if e == 'test_env':
        return True
    if e in host_slots and f == 'cmd':
        return True
    if contract is not None:
        entry = contract_entry(e, f, contract)
        if isinstance(entry, Mapping) and entry.get('dispatch') == 'cmd_primitive':
            required = (entry.get('signature') or {}).get('required') if isinstance(entry.get('signature'), Mapping) else None
            first = str((required or [''])[0] if required else '')
            return first == 'cmd'
    return e.startswith('APV') and f == 'cmd_config'

def _check_assertion_matches_command_echo(steps: list, result: StructuralResult) -> None:
    contract = _load_contract_for_gate(result)
    if contract is None:
        return
    host_slots = _contract_host_slot_es(contract)
    last_src_g: str | None = None
    reg_src: dict[str, str] = {}
    for i, s in enumerate(steps):
        if not isinstance(s, dict):
            continue
        e = str(s.get('E', '')).strip()
        f = str(s.get('F', '')).strip()
        g = str(s.get('G', '') or '')
        h = str(s.get('H', '') or '').strip()
        i_col = str(s.get('I', '') or '').strip()
        if e == 'check_point':
            if h or not g.strip() or f not in ('found', 'not_found', 'abs_found'):
                continue
            src = reg_src.get(i_col) if i_col else last_src_g
            if not src:
                continue
            try:
                hit = src.find(g) >= 0 if f == 'abs_found' else re.compile(g, re.DOTALL).search(src) is not None
            except re.error:
                continue
            if hit:
                result.add('assertion_matches_command_echo', f"assertion {f}'s pattern already matches the **command text itself** of the step sourcing its window ({src[:60]!r}…) — the window always starts with the command echo, so this assertion is " + ('**always-fail** (it tries to verify "output contains no X" but X is in the command).' if f == 'not_found' else '**always-true fake PASS** (it passes no matter what the device outputs, verifying nothing).') + ' Rewrite it to match a data line, distinguishable from the command text.', i)
            continue
        if h:
            if _echo_src_cmd(e, f, host_slots, contract):
                reg_src[h] = g
            else:
                reg_src.pop(h, None)
        else:
            last_src_g = g if _echo_src_cmd(e, f, host_slots, contract) else None

def _check_has_assertion(steps: list, result: StructuralResult) -> None:
    for s in steps:
        if isinstance(s, dict) and str(s.get('E', '')).strip() == 'check_point':
            return
    if steps:
        result.add('no_assertion_in_case', 'this case has no check_point step at all — framework settlement judges success==0 as FAIL (check_point.py:126), so a config-only/observation-only sheet is always-fail on the device. Add at least one assertion; steps that only execute without verifying do not constitute a test case.', 0)

def _check_empty_assertion_pattern(steps: list, result: StructuralResult) -> None:
    for i, s in enumerate(steps):
        if not isinstance(s, dict) or str(s.get('E', '')).strip() != 'check_point':
            continue
        if str(s.get('G', '') or '').strip() or str(s.get('H', '') or '').strip() or str(s.get('I', '') or '').strip():
            continue
        result.add('empty_assertion_pattern', 'check_point has G/H/I all empty — no pattern, no register reference, the framework has nothing to compare (it falls back to searching with the observation command text; 044605 evidence: one on-device round wasted). Write a pattern in G, or reference an already-captured H register. If expected is not independently known, write <RUNTIME> only as an underdetermined marker; device actuals may be registered as digests but never backfilled into expected. An identity-bound Author, Spec, DefectSpec, Manual, ConfigBinding, or CapabilityXml claim must supply expected before delivery.', i)

def _check_assertion_regex_compiles(steps: list, result: StructuralResult) -> None:
    for i, s in enumerate(steps):
        if not isinstance(s, dict) or str(s.get('E', '')).strip() != 'check_point':
            continue
        f = str(s.get('F', '')).strip()
        g = str(s.get('G', '') or '')
        h = str(s.get('H', '') or '').strip()
        if f not in ('found', 'not_found', 'found_times') or not g:
            continue
        if h and f != 'found_times':
            continue
        try:
            re.compile(g)
        except re.error as exc:
            result.add('assertion_regex_invalid', f'assertion regex fails to compile ({exc}): {g[:80]!r} — the framework raises at re.compile and the entire file crashes. Fix the regex syntax (common: an unclosed character class, [^ should be [^\\n]).', i)

def _check_short_mode_assertions(steps: list, result: StructuralResult) -> None:
    last_obs_short = False
    for i, s in enumerate(steps):
        if not isinstance(s, dict):
            continue
        e = str(s.get('E', '')).strip()
        if e == 'check_point':
            g = str(s.get('G', '') or '')
            if last_obs_short and _SHORT_INCOMPATIBLE_RE.search(g):
                result.add('short_mode_status_assertion', 'the assertion matches dig status/HEADER/SECTION text, but the observation step it consumes used +short (output has only record values, none of those sections) — always-fail. Remove +short from that dig, or rewrite the assertion in record-value form.', i)
            continue
        h = str(s.get('H', '') or '').strip()
        if not h:
            g = str(s.get('G', '') or '')
            last_obs_short = 'dig' in g and '+short' in g

def _check_capture_refs_defined(steps: list, result: StructuralResult) -> None:
    contract = _load_contract_for_gate(result)
    object_names = {str(item['e']) for item in contract['objects']} if contract is not None else set()
    defined: set[str] = set()
    for i, s in enumerate(steps):
        if not isinstance(s, dict):
            continue
        e = str(s.get('E', '')).strip()
        f = str(s.get('F', '') or '').strip()
        h = str(s.get('H', '') or '').strip()
        i_col = str(s.get('I', '') or '').strip()
        if e == 'check_point':
            refs = () if f == 'found_times' else (h, i_col)
            for ref in refs:
                if ref and ref not in defined:
                    result.add('undefined_capture_ref', f'the assertion references register {ref!r}, but no earlier step captured it with H={ref} — the framework reads None and the assertion is distorted. Add a capture step or fix the reference name.', i)
        else:
            if i_col:
                base = i_col.split('.', 1)[0]
                if base not in defined and base not in object_names:
                    result.add('undefined_capture_ref', f'step I={i_col!r} references a variable that no earlier step captured with H — for an undefined I on a non-assertion step the framework **raises NameError and crashes the entire file**. Capture it with H first, or fix the reference name.', i)
            g = str(s.get('G', '') or '')
            if not i_col and _curl_write_out_has_doubled_variable(g):
                result.add('injection_format_crash', "curl -w/--write-out contains a doubled-brace variable such as '%{{name}}'. With column I blank the framework validates G but never calls .format(), so the doubled braces reach curl unchanged; curl then treats the variable name as invalid. Use curl's exact single-brace syntax only if the F/G carrier can represent it, otherwise use an equivalent observation without write-out variables or report the carrier-language gap.", i)
            try:
                validate_i_injection_syntax(g, i_col, f)
            except ExcelContractError:
                pass
            except IInjectionSyntaxError as exc:
                code = {'missing': 'injection_without_placeholder', 'scope': 'injection_placeholder_scope', 'format': 'injection_format_crash'}[exc.code]
                if exc.code == 'scope':
                    repair = " Accepted shapes: static — remove every placeholder from G and leave I blank; injected — capture an earlier H register (or use a declared runtime object attribute), put exactly one {} or {0} in G's first parsed positional argument, and set I to that register/attribute. No later positional argument or keyword argument may contain a placeholder."
                elif exc.code == 'format':
                    repair = " The currently bound runner still validates static G through Python's Formatter grammar even with I blank; this literal-brace feature is not certified on that runtime. Use an equivalent observation with no named braces, or complete the runner certification/deployment first." if not i_col else " Python placeholder grammar applies because I is present and the runtime will call .format() on G's first positional argument. Keep downstream literal braces out of that formatted argument, or use a static row with blank I when no injection is required."
                else:
                    repair = ''
                result.add(code, f'E={e!r}, F={f!r}, I={i_col!r}: {exc}. The runtime parses G first and formats only positional argument 1.{repair}', i)
            if h:
                if contract is not None and h in _framework_reserved_names(contract):
                    result.add('register_shadows_framework_name', f"H={h!r} collides with a v2 runner **runtime object name** (a contract device slot, or a runner built-in parsed from the mirror source). The v2 preflight rejects it before any step executes (NameError, verbatim runner message: 'H 不能覆盖运行时对象') and _store_register raises the same at runtime — either way the whole pytest file crashes. Registers are a separate dict in v2, so ordinary frame-local names are legal register names; pick one that is not a runtime object (like v1/ip1).", i)
                defined.add(h)

def _curl_write_out_has_doubled_variable(g: str) -> bool:
    import shlex
    try:
        tokens = shlex.split(str(g or ''), posix=True)
    except ValueError:
        return False
    curl_indexes = [index for index, token in enumerate(tokens) if token.rsplit('/', 1)[-1] == 'curl']
    if not curl_indexes:
        return False
    start = curl_indexes[-1] + 1
    doubled = re.compile('%\\{\\{[^{}\\s]+\\}\\}')
    args = tokens[start:]
    for index, token in enumerate(args):
        if token in {'-w', '--write-out'}:
            if index + 1 < len(args) and doubled.search(args[index + 1]):
                return True
            continue
        if token.startswith('--write-out=') and doubled.search(token.split('=', 1)[1]):
            return True
        if token.startswith('-w') and token != '-w' and doubled.search(token[2:]):
            return True
    return False

def _framework_reserved_names(contract: dict | None=None) -> frozenset:
    contract = contract if contract is not None else load_excel_contract()
    names: set[str] = {str(item['e']) for item in contract['objects']}
    for builtin in _v2_runner_builtin_names():
        names.add(builtin)
    return frozenset((n for n in names if n))

def _v2_runner_builtin_names() -> frozenset:
    src = mirror_src('lib/test_xlsx.py')
    if not src:
        return frozenset()
    import ast as _ast
    try:
        tree = _ast.parse(src)
    except SyntaxError:
        return frozenset()
    names: set[str] = set()
    for node in _ast.walk(tree):
        if isinstance(node, _ast.Assign) and any((isinstance(t, _ast.Name) and t.id == 'runtime_objects' for t in node.targets)) and isinstance(node.value, _ast.Dict):
            for key in node.value.keys:
                if isinstance(key, _ast.Constant) and isinstance(key.value, str):
                    names.add(key.value)
        elif isinstance(node, _ast.Call) and isinstance(node.func, _ast.Attribute) and (node.func.attr == 'update') and isinstance(node.func.value, _ast.Name) and (node.func.value.id in ('runtime_names', 'runtime_objects')):
            for arg in node.args:
                if isinstance(arg, (_ast.Tuple, _ast.List, _ast.Set)):
                    for el in arg.elts:
                        if isinstance(el, _ast.Constant) and isinstance(el.value, str):
                            names.add(el.value)
    return frozenset(names)
_KWARG_SEG_RE = re.compile('^\\s*(timeout|prompt)\\s*=')

def _check_parameter_splitting(steps: list, result: StructuralResult) -> None:
    contract = _load_contract_for_gate(result)
    if contract is None:
        return
    host_slots = _contract_host_slot_es(contract)
    for i, s in enumerate(steps):
        if not isinstance(s, dict):
            continue
        e = str(s.get('E', '')).strip()
        if not (e == 'test_env' or e in host_slots):
            continue
        f = str(s.get('F', '')).strip()
        if f == 'execute':
            continue
        g = str(s.get('G', '') or '')
        if '\n' in g or ',' not in g:
            continue
        segs = [p.strip() for p in _PARAM_SPLIT_RE.findall(g) if p.strip()]
        stray = [p for p in segs[1:] if not _KWARG_SEG_RE.match(p)]
        if stray:
            result.add('comma_splits_parameters', f"G contains an unquoted comma, so the framework splits it into {len(segs)} positional parameters — {stray[0]!r} gets mispassed to the host method's prompt/timeout parameter; the step always waits out the full timeout and the output is distorted. If the comma is part of the command, quote that segment, or rewrite the command to avoid commas; to pass a timeout use the timeout=N form (framework named parameter).", i)

def _check_config_existence_advisories(steps: list, result: StructuralResult) -> None:
    config_context: list[str] = []
    last_observe = ''
    for i, step in enumerate(steps if isinstance(steps, list) else []):
        if not isinstance(step, dict):
            continue
        e = str(step.get('E', '') or '').strip()
        f = str(step.get('F', '') or '').strip()
        g = str(step.get('G', '') or '').strip()
        if e == 'check_point':
            is_echo, matched = config_existence_check(last_observe, g, config_context, f or 'found')
            if is_echo:
                result.advise('config_existence_only', 'the assertion only confirms that a configuration written earlier is present in a read-only configuration query; it does not demonstrate the target runtime behavior. Keep it only when configuration existence is the stated test objective, otherwise add an independently sourced behavior observation.', i)
            continue
        if f in {'cmd_config', 'cmds_config'}:
            config_context.extend((line.strip() for line in g.splitlines() if line.strip()))
        if observe_kind(g):
            last_observe = g

def lint_xlsx_case(xlsx_path) -> StructuralResult:
    result = StructuralResult()
    try:
        autoid, steps = steps_from_xlsx(xlsx_path, allow_legacy=True)
    except Exception as exc:
        result.add('xlsx_unreadable', f'sheet is unreadable: {exc}')
        return result
    if autoid and (not _AUTOID_RE.match(autoid)):
        result.add('autoid_malformed', f'sheet autoid {autoid!r} is not an 18-digit number — a hand-copied truncated id silently creates a junk directory and sneaks into the final sheet (evidence: once produced a 35-case final sheet). Use the machine-readable full id from last_run.json/manifest.')
    mand = check_crash_gates_mandatory(steps)
    if not mand.ok:
        result.ok = False
        result.violations.extend(mand.violations)
    result.disabled.extend(mand.disabled)
    _check_assertion_regex_compiles(steps, result)
    _check_short_mode_assertions(steps, result)
    _check_parameter_splitting(steps, result)
    _check_autoid_rows_runnable(xlsx_path, result)
    _check_config_existence_advisories(steps, result)
    return result

def _check_autoid_rows_runnable(xlsx_path, result: StructuralResult) -> None:
    try:
        _autoid, _steps, empty_e_autoids = _workbook_model(xlsx_path, allow_legacy=True)
    except Exception:
        return
    for a in empty_e_autoids:
        result.add('autoid_row_not_runnable', f"autoid row ({a}) has an empty column E — the deployed runner cannot start a case from it: the v1 framework silently skipped the whole case (not executed, no fail counted, no log), and the v2 scheduler skips the row before the new-case boundary, so this case's steps execute inside the previous case (verdicts attributed to the wrong case, device state carried across). The autoid must share its row with the first step (column E non-empty).")
