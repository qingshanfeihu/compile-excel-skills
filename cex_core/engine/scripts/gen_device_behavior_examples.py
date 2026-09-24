# 生成：tools/extract_engine.py ← InfoTest scripts/gen_device_behavior_examples.py（sha256 1a0335a11bd8e642）。不在这里手改。
from __future__ import annotations
from cex_core.engine._root import _cex_data_path
import argparse
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile
import yaml
from pathlib import Path
from typing import Any, Callable
from openpyxl import load_workbook
ROOT = _cex_data_path('')
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
DEFAULT_OUTPUT = ROOT / 'knowledge/data/compile_ref/device_behavior_examples.json'
CORPUS_ROOT_REL = Path('knowledge/data/device_behavior_corpus')
CORPUS_SEAL_NAME = 'corpus.seal.json'
CORPUS_SEAL_SCHEMA = 'ist.device-behavior-corpus-seal.v1'
CORPUS_ROLES = frozenset({'verified_workbook', 'attribution_receipt', 'harvest_workbook', 'user_adjudication', 'domain_grammar'})
TRANSLATION_SHAPE_AXES: tuple[tuple[str, str], ...] = (('certificate_method_surface', 'm_cert'), ('positive_text_assertion', 'op_found'), ('negative_text_assertion', 'op_not_found'), ('stable_wait', 'sleep'), ('curl_trigger', 't_curl'), ('dig_trigger', 't_dig'))

class DeviceBehaviorGenerationError(RuntimeError):
    pass

class DeviceBehaviorSourceUnavailableError(DeviceBehaviorGenerationError):
    pass

def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()

def _text(value: Any) -> str:
    return '' if value is None else str(value).strip()

def _rel_posix(path: Path, project_root: Path) -> str:
    try:
        return path.relative_to(project_root).as_posix()
    except ValueError:
        return path.as_posix()

def _receipt_row_valid(row: Any) -> bool:
    return isinstance(row, dict) and bool(_text(row.get('autoid'))) and (len(_text(row.get('raw_echo_sha256'))) == 64) and bool(_text(row.get('verdict') or row.get('framework_result')))

def _load_corpus(corpus_root: Path, project_root: Path) -> dict[str, Any]:
    if not corpus_root.is_dir():
        raise DeviceBehaviorSourceUnavailableError(f'corpus root missing: {corpus_root}')
    seal_path = corpus_root / CORPUS_SEAL_NAME
    if seal_path.is_symlink() or not seal_path.is_file():
        raise DeviceBehaviorSourceUnavailableError(f'corpus seal missing: {seal_path}')
    try:
        seal = json.loads(seal_path.read_text(encoding='utf-8'))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise DeviceBehaviorSourceUnavailableError(f'corpus seal unreadable: {seal_path}: {exc}') from exc
    if not isinstance(seal, dict) or seal.get('schema') != CORPUS_SEAL_SCHEMA:
        raise DeviceBehaviorSourceUnavailableError(f'corpus seal schema mismatch: {seal_path}')
    raw_files = seal.get('files')
    if not isinstance(raw_files, list) or not raw_files:
        raise DeviceBehaviorSourceUnavailableError(f'corpus seal lists no evidence files: {seal_path}')
    resolved_root = corpus_root.resolve()
    entries: list[dict[str, Any]] = []
    seen: set[str] = set()
    for index, raw in enumerate(raw_files):
        if not isinstance(raw, dict):
            raise DeviceBehaviorSourceUnavailableError(f'corpus seal entry #{index} is not an object')
        rel = _text(raw.get('path'))
        role = _text(raw.get('role'))
        sha256 = _text(raw.get('sha256'))
        autoid = _text(raw.get('autoid'))
        if role not in CORPUS_ROLES:
            raise DeviceBehaviorSourceUnavailableError(f'corpus seal entry {rel or index} has unknown role: {role!r}')
        rel_path = Path(rel)
        if not rel or rel_path.is_absolute() or '..' in rel_path.parts:
            raise DeviceBehaviorSourceUnavailableError(f'corpus seal entry escapes corpus root: {rel!r}')
        if not re.fullmatch('[0-9a-f]{64}', sha256):
            raise DeviceBehaviorSourceUnavailableError(f'corpus seal entry {rel} has malformed sha256')
        if autoid and (not re.fullmatch('\\d{18}', autoid)):
            raise DeviceBehaviorSourceUnavailableError(f'corpus seal entry {rel} has malformed autoid: {autoid!r}')
        candidate = corpus_root / rel_path
        if candidate.is_symlink() or not candidate.is_file():
            raise DeviceBehaviorSourceUnavailableError(f'corpus evidence file missing: {rel}')
        if not candidate.resolve().is_relative_to(resolved_root):
            raise DeviceBehaviorSourceUnavailableError(f'corpus evidence file resolves outside corpus root: {rel}')
        if _sha256(candidate) != sha256:
            raise DeviceBehaviorSourceUnavailableError(f'corpus evidence file sha256 mismatch: {rel}')
        if role == 'verified_workbook' and (not re.fullmatch('verified_\\d+\\.xlsx', candidate.name)):
            raise DeviceBehaviorSourceUnavailableError(f'verified workbook entry has unexpected name: {rel}')
        if role == 'harvest_workbook' and (not autoid):
            raise DeviceBehaviorSourceUnavailableError(f'harvest workbook entry lacks autoid identity: {rel}')
        if rel in seen:
            raise DeviceBehaviorSourceUnavailableError(f'corpus seal lists {rel} twice')
        seen.add(rel)
        entries.append({'path': candidate, 'relpath': _rel_posix(candidate, project_root), 'corpus_relpath': rel, 'sha256': sha256, 'role': role, 'autoid': autoid})
    by_role: dict[str, list[dict[str, Any]]] = {role: [] for role in CORPUS_ROLES}
    for entry in entries:
        by_role[entry['role']].append(entry)
    if not by_role['verified_workbook']:
        raise DeviceBehaviorSourceUnavailableError(f'sealed corpus contains no verified workbook: {corpus_root}')
    if len(by_role['domain_grammar']) != 1:
        raise DeviceBehaviorSourceUnavailableError(f'sealed corpus must carry exactly one domain_grammar entry: {corpus_root}')
    for path in sorted(corpus_root.rglob('*')):
        if not path.is_file() or path.name.startswith('.'):
            continue
        rel = path.relative_to(corpus_root).as_posix()
        if rel == CORPUS_SEAL_NAME:
            continue
        if rel not in seen:
            raise DeviceBehaviorSourceUnavailableError(f'unlisted file inside sealed corpus: {rel}')
    return {'root': corpus_root, 'seal_path': seal_path, 'seal_sha256': _sha256(seal_path), 'files': entries, 'by_role': by_role}

def _corpus_grammar(corpus: dict[str, Any]) -> dict[str, Any]:
    entry = corpus['by_role']['domain_grammar'][0]
    try:
        payload = json.loads(entry['path'].read_text(encoding='utf-8'))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise DeviceBehaviorSourceUnavailableError(f"sealed domain_grammar unreadable: {entry['corpus_relpath']}") from exc
    if not isinstance(payload, dict):
        raise DeviceBehaviorSourceUnavailableError(f"sealed domain_grammar is not an object: {entry['corpus_relpath']}")
    return payload

def _workbook_rows(path: Path, *, autoid: str='') -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    rows: list[dict[str, Any]] = []
    try:
        for sheet in workbook.worksheets:
            current_autoid = ''
            target_seen = False
            for number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                cells = list(values)
                row_autoid = _text(cells[0]) if cells else ''
                if re.fullmatch('\\d{18}', row_autoid):
                    if autoid and target_seen and (row_autoid != autoid):
                        break
                    current_autoid = row_autoid
                    target_seen = target_seen or row_autoid == autoid
                if autoid and current_autoid != autoid:
                    continue
                if len(cells) < 7:
                    continue
                e, f, g = (_text(cells[index]) for index in (4, 5, 6))
                if not (e or f or g):
                    continue
                rows.append({'sheet': sheet.title, 'row': number, 'desc': _text(cells[3]), 'E': e, 'F': f, 'G': g})
    finally:
        workbook.close()
    return rows

def _verified_translation_features(path: Path, autoid: str) -> tuple[dict[str, bool], list[dict[str, Any]]]:
    axis = {source_key: False for _name, source_key in TRANSLATION_SHAPE_AXES}
    semantic_rows: list[dict[str, Any]] = []
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        for sheet in workbook.worksheets:
            in_body = False
            in_case = False
            for number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                cells = [_text(value) for value in values]
                first = cells[0] if cells else ''
                if first == '自动化ID':
                    in_body = True
                    continue
                if not in_body:
                    continue
                if any((re.search('(?:cert|key|csr|rootca|crlca|interca)', cell, re.IGNORECASE) for cell in cells)):
                    axis['m_cert'] = True
                if first == autoid:
                    in_case = True
                if first == '999999999999999':
                    e = cells[4] if len(cells) > 4 else ''
                    f = cells[5] if len(cells) > 5 else ''
                    g = cells[6] if len(cells) > 6 else ''
                    axis['op_found'] |= e == 'check_point' and f == 'found'
                    axis['op_not_found'] |= e == 'check_point' and f == 'not_found'
                    axis['sleep'] |= e == 'time' and f == 'sleep'
                    axis['t_curl'] |= bool(re.search('(?<!\\w)curl(?!\\w)', g, re.I))
                    axis['t_dig'] |= bool(re.search('(?<!\\w)dig(?!\\w)', g, re.I))
                    in_case = False
                    continue
                if len(cells) < 7:
                    continue
                e, f, g = (cells[4], cells[5], cells[6])
                if in_case and (e or f or g):
                    semantic_rows.append({'sheet': sheet.title, 'row': number, 'desc': cells[3] if len(cells) > 3 else '', 'E': e, 'F': f, 'G': g, 'H': cells[7] if len(cells) > 7 else '', 'I': cells[8] if len(cells) > 8 else ''})
                if not in_case:
                    continue
                axis['op_found'] |= e == 'check_point' and f == 'found'
                axis['op_not_found'] |= e == 'check_point' and f == 'not_found'
                axis['sleep'] |= e == 'time' and f == 'sleep'
                axis['t_curl'] |= bool(re.search('(?<!\\w)curl(?!\\w)', g, re.I))
                axis['t_dig'] |= bool(re.search('(?<!\\w)dig(?!\\w)', g, re.I))
    finally:
        workbook.close()
    if not semantic_rows:
        raise DeviceBehaviorGenerationError(f'verified workbook contains no bound case rows: {path.name}')
    return (axis, semantic_rows)

def _translation_signature(axis: dict[str, bool]) -> str:
    return '-'.join((f'{name}={int(bool(axis[source_key]))}' for name, source_key in TRANSLATION_SHAPE_AXES))

def _row_naming_styles(rows: list[dict[str, Any]]) -> list[str]:
    commands = [command for row in rows if row.get('E') != 'check_point' for command in _commands(row)]
    quoted = [match.group(2) for command in commands for match in re.finditer('([\\"\'])([^\\"\'\\r\\n]{1,96})\\1', command)]
    tokens = [token for command in commands for token in re.findall('[A-Za-z][A-Za-z0-9_.-]*', command)]
    tags: list[str] = []
    if quoted:
        tags.append('quoted_identifiers')
    if any(('_' in token for token in tokens)):
        tags.append('underscore_separated_identifiers')
    if any(('-' in token for token in tokens)):
        tags.append('hyphen_separated_identifiers')
    if any(('.' in token for token in tokens)):
        tags.append('dotted_identifiers')
    if any((re.search('[A-Za-z][0-9]+\\Z', token) for token in tokens)):
        tags.append('numeric_suffix_identifiers')
    if len(set(quoted)) < len(quoted):
        tags.append('quoted_identifier_reuse')
    return tags

def _translation_shape_summary(rows: list[dict[str, Any]]) -> dict[str, Any]:
    product_roles = sorted({_text(row.get('E')) for row in rows if _text(row.get('E')).startswith('APV_')})
    trigger_roles = sorted({_text(row.get('F')) for row in rows if row.get('E') == 'test_env' and _text(row.get('F'))})
    assertion_ops = sorted({_text(row.get('F')) for row in rows if row.get('E') == 'check_point' and _text(row.get('F'))})
    trigger_kinds = sorted({'curl' if re.search('(?<!\\w)curl(?!\\w)', _text(row.get('G')), re.I) else 'dig' if re.search('(?<!\\w)dig(?!\\w)', _text(row.get('G')), re.I) else 'environment_command' for row in rows if row.get('E') == 'test_env'})
    first_trigger = next((index for index, row in enumerate(rows) if row.get('E') == 'test_env'), len(rows))
    wait_before_trigger = any((index < first_trigger and row.get('E') == 'time' and (row.get('F') == 'sleep') for index, row in enumerate(rows)))
    witness_rows = sorted({int(row['row']) for row in rows if row.get('E') in {'check_point', 'test_env', 'time'} or _text(row.get('E')).startswith('APV_')})[:20]
    return {'topology_roles': {'product_devices': product_roles, 'trigger_roles': trigger_roles, 'multi_product_device': len(product_roles) > 1}, 'trigger_channels': trigger_kinds, 'assertion_evidence_shapes': assertion_ops, 'naming_style': _row_naming_styles(rows), 'lifecycle_shape': {'config_row_count': sum((_text(row.get('E')).startswith('APV_') and row.get('F') in {'cmd_config', 'cmds_config', 'cmd_enable'} for row in rows)), 'stable_wait_before_trigger': wait_before_trigger, 'captured_value_relation': any((row.get('E') == 'check_point' and _text(row.get('H')) for row in rows))}, 'witness_rows': witness_rows}

def _shape_statement(summary: dict[str, Any]) -> str:
    topology = summary['topology_roles']
    product_count = len(topology['product_devices'])
    trigger_count = len(topology['trigger_roles'])
    assertions = summary['assertion_evidence_shapes'] or ['none']
    triggers = summary['trigger_channels'] or ['none']
    wait = summary['lifecycle_shape']['stable_wait_before_trigger']
    return f"Epoch-revalidated translation shape: configure {product_count} product role(s), drive {trigger_count} trigger role(s) through {', '.join(triggers)}, and carry assertion evidence with {', '.join(assertions)}. Stable wait before the first trigger is {('present' if wait else 'absent')}. Reuse only this topology/trigger/evidence/naming structure; derive every expected value from the current case's identity-bound Author, Spec, DefectSpec, Manual, ConfigBinding or CapabilityXml claim."

def _commands(row: dict[str, Any]) -> list[str]:
    return [line.strip() for line in _text(row.get('G')).splitlines() if line.strip()]

def _receipt_text(row: dict[str, Any]) -> str:
    values: list[str] = []
    for key in ('detail_full', 'device_context', 'causality_full', 'causality'):
        value = row.get(key)
        if isinstance(value, str):
            values.append(value)
    for key in ('anomaly_lines', '_fail_signatures'):
        value = row.get(key)
        if isinstance(value, list):
            values.extend((str(item) for item in value))
    return '\n'.join(values)

def _translation_adjudications(entries: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    by_entry: dict[str, dict[str, Any]] = {}
    for entry in entries:
        path = entry['path']
        try:
            text = path.read_text(encoding='utf-8')
        except (OSError, UnicodeError):
            continue
        match = re.match('^---\\n(.*?)\\n---\\n', text, re.DOTALL)
        if match is None:
            continue
        try:
            frontmatter = yaml.safe_load(match.group(1)) or {}
        except yaml.YAMLError:
            continue
        entry_id = _text(frontmatter.get('projection_entry_id'))
        if not entry_id:
            continue
        channel = _text(frontmatter.get('channel'))
        if channel != 'ordered_user_adjudications':
            continue
        ruling_body = text[match.end():].strip()
        heading = '# 裁决'
        if not ruling_body.startswith(heading):
            continue
        ruling = ruling_body[len(heading):].strip()
        if not ruling:
            continue
        if entry_id in by_entry:
            raise DeviceBehaviorGenerationError(f'multiple ordered adjudications target projection entry {entry_id}')
        by_entry[entry_id] = {'path': entry['relpath'], 'sha256': entry['sha256'], 'intent_signature': _text(frontmatter.get('intent_signature')), 'conflict_shape': _text(frontmatter.get('conflict_shape')), 'version_family': _text(frontmatter.get('version_family')), 'case_family': _text(frontmatter.get('case_family')), 'channel': channel, 'lineage': _text((frontmatter.get('anchor') or {}).get('lineage')), 'ruling_sha256': hashlib.sha256(ruling.encode('utf-8')).hexdigest()}
    return by_entry

def _load_receipts(entries: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    by_autoid: dict[str, list[dict[str, Any]]] = {}
    for entry in entries:
        path = entry['path']
        try:
            payload = json.loads(path.read_text(encoding='utf-8'))
        except (OSError, UnicodeError, json.JSONDecodeError):
            continue
        rows = payload if isinstance(payload, list) else [payload]
        for row in rows:
            if not _receipt_row_valid(row):
                continue
            autoid = _text(row.get('autoid'))
            by_autoid.setdefault(autoid, []).append({'path': entry['relpath'], 'sha256': entry['sha256'], 'raw_echo_sha256': _text(row.get('raw_echo_sha256')), 'verdict': _text(row.get('verdict') or row.get('framework_result')), 'receipt_schema': _text((row.get('framework_result_receipt') or {}).get('schema') if isinstance(row.get('framework_result_receipt'), dict) else ''), 'run_ts': float(row.get('_run_ts') or 0), '_text': _receipt_text(row)})
    for receipts in by_autoid.values():
        receipts.sort(key=lambda item: (item['run_ts'], item['path']))
    return by_autoid

def _pick_receipt(receipts: list[dict[str, Any]], predicate: Callable[[str], bool] | None=None) -> dict[str, Any] | None:
    for receipt in receipts:
        if predicate is None or predicate(str(receipt.get('_text') or '')):
            return receipt
    return None

def _source_receipt(receipt: dict[str, Any]) -> dict[str, Any]:
    return {key: receipt[key] for key in ('path', 'sha256', 'raw_echo_sha256', 'verdict', 'receipt_schema')}

def _harvest_workbooks(entries: list[dict[str, Any]]) -> tuple[dict[str, dict[str, Any]], list[str]]:
    found: dict[str, dict[str, Any]] = {}
    unreadable: list[str] = []
    for entry in entries:
        autoid = entry['autoid']
        if autoid in found:
            continue
        path = entry['path']
        try:
            rows = _workbook_rows(path, autoid=autoid)
        except Exception as exc:
            unreadable.append(f"{entry['relpath']}:{type(exc).__name__}")
            continue
        found[autoid] = {'rows': rows, 'source': {'path': entry['relpath'], 'sha256': entry['sha256']}}
    return (found, unreadable)

def _persistence_family(command: str, patterns: list[re.Pattern[str]]) -> tuple[str, bool] | None:
    for pattern in patterns:
        match = pattern.search(command)
        if match is None:
            continue
        family = next((str(value).casefold() for value in match.groups() if value), ' '.join(command[:match.end()].casefold().split()[:2]))
        return (family, bool(command[match.end():].strip()))
    return None

def _argumented_save_family_equivalent_observation(verified_rows: list[dict[str, Any]], harvest_rows: list[dict[str, Any]], grammar: dict[str, Any]) -> dict[str, Any] | None:
    local_disk = (grammar.get('persistence_channels') or {}).get('local_disk') or {}
    raw_persistence = local_disk.get('patterns') or []
    destructive = (grammar.get('destructive_commands') or {}).get('patterns') or []
    observe_verbs = set(((grammar.get('verb_classes') or {}).get('observe_leading') or {}).get('verbs') or [])
    try:
        persistence = [re.compile(pattern, re.IGNORECASE) for pattern in raw_persistence]
        destructive_res = [re.compile(pattern, re.IGNORECASE) for pattern in destructive]
    except re.error:
        return None
    if not persistence or not destructive_res or (not observe_verbs):
        return None
    harvest_by_family: dict[str, tuple[int, bool]] = {}
    harvest_forbidden_rows: list[int] = []
    for row in harvest_rows:
        for command in _commands(row):
            classified = _persistence_family(command, persistence)
            if classified is not None:
                family, argumented = classified
                harvest_by_family.setdefault(family, (row['row'], argumented))
            if any((pattern.search(command) for pattern in destructive_res)):
                harvest_forbidden_rows.append(row['row'])
    if not harvest_forbidden_rows:
        return None
    for index, row in enumerate(verified_rows):
        for command in _commands(row):
            classified = _persistence_family(command, persistence)
            if classified is None:
                continue
            family, argumented = classified
            harvest = harvest_by_family.get(family)
            if not argumented or harvest is None or harvest[1] is True:
                continue
            observe_at = next((at for at in range(index + 1, min(index + 5, len(verified_rows))) if any((candidate.split(None, 1)[0].casefold() in observe_verbs for candidate in _commands(verified_rows[at]) if candidate.split(None, 1)))), None)
            if observe_at is None:
                continue
            assertion_at = next((at for at in range(observe_at + 1, min(observe_at + 3, len(verified_rows))) if verified_rows[at].get('E') == 'check_point'), None)
            if assertion_at is None:
                continue
            if any((pattern.search(candidate) for verified in verified_rows for candidate in _commands(verified) for pattern in destructive_res)):
                continue
            return {'detector': 'argumented_persistence_family_with_non_destructive_equivalent_observation', 'persistence_family': family, 'verified_rows': [row['row'], verified_rows[observe_at]['row'], verified_rows[assertion_at]['row']], 'harvest_rows': [harvest[0], *harvest_forbidden_rows[:4]], 'argumented_form': True, 'self_authored_target_lane': 'W24_value_grounding', 'harvest_bare_form': True, 'harvest_forbidden_mechanism': True, 'equivalent_observation_without_forbidden_mechanism': True}
    return None

def _rebuild_enable_wait(rows: list[dict[str, Any]], receipts: list[dict[str, Any]]) -> tuple[dict[str, Any], dict[str, Any]] | None:
    for delete_at, delete_row in enumerate(rows):
        if not any((command.casefold().startswith('no ') for command in _commands(delete_row))):
            continue
        for rebuild_at in range(delete_at + 1, min(delete_at + 7, len(rows))):
            commands = _commands(rows[rebuild_at])
            has_create = any((not command.casefold().startswith('no ') and ' virtual ' in command.casefold() for command in commands))
            has_enable = any((' virtual enable ' in f' {command.casefold()} ' for command in commands))
            if not (has_create and has_enable):
                continue
            wait_at = next((index for index in range(rebuild_at + 1, min(rebuild_at + 4, len(rows))) if rows[index]['E'] == 'time' and rows[index]['F'] == 'sleep'), None)
            if wait_at is None:
                continue
            probe_at = next((index for index in range(wait_at + 1, min(wait_at + 5, len(rows))) if rows[index]['E'] == 'test_env' or any((command.casefold().startswith('show ') for command in _commands(rows[index])))), None)
            if probe_at is None:
                continue
            receipt = _pick_receipt(receipts, lambda actual: ' virtual enable ' not in f' {actual.casefold()} ') or _pick_receipt(receipts)
            if receipt is None:
                continue
            return ({'detector': 'rebuild_then_enable_then_stable_wait_then_probe', 'rows': [delete_row['row'], rows[rebuild_at]['row'], rows[wait_at]['row'], rows[probe_at]['row']], 'stable_wait_seconds': _text(rows[wait_at]['G']), 'failed_route_contrast_omitted_enable': ' virtual enable ' not in f" {receipt['_text'].casefold()} "}, receipt)
    return None

def _negative_actual_feedback(rows: list[dict[str, Any]], receipts: list[dict[str, Any]]) -> tuple[dict[str, Any], dict[str, Any]] | None:
    literal_feedback = re.compile('^[A-Za-z][A-Za-z0-9 .,:;\'\\"_-]{11,}$')
    for index, row in enumerate(rows):
        pattern = _text(row.get('G'))
        if row.get('E') != 'check_point' or row.get('F') != 'found':
            continue
        if not literal_feedback.fullmatch(pattern):
            continue
        receipt = _pick_receipt(receipts, lambda actual: pattern in actual)
        if receipt is None:
            continue
        prior = rows[index - 1] if index else {}
        return ({'detector': 'literal_assertion_reappears_in_device_or_framework_actual', 'rows': [row['row']], 'preceded_by_config_attempt': prior.get('F') in {'cmd_config', 'cmds_config'}, 'observed_actual_sample': pattern}, receipt)
    return None

def _typed_list_existence(rows: list[dict[str, Any]], receipts: list[dict[str, Any]]) -> tuple[dict[str, Any], dict[str, Any]] | None:
    pairs: list[tuple[int, dict[str, Any], dict[str, Any], list[str]]] = []
    for index in range(len(rows) - 1):
        observe, assertion = (rows[index], rows[index + 1])
        commands = _commands(observe)
        if len(commands) == 1 and commands[0].casefold().startswith('show ') and (assertion.get('E') == 'check_point') and (assertion.get('F') in {'found', 'not_found'}):
            pairs.append((index, observe, assertion, commands[0].split()))
    for left_at in range(len(pairs)):
        _, left_show, left_assert, left_tokens = pairs[left_at]
        for right_at in range(left_at + 1, len(pairs)):
            _, right_show, right_assert, right_tokens = pairs[right_at]
            if _text(left_assert.get('G')) != _text(right_assert.get('G')):
                continue
            if {left_assert.get('F'), right_assert.get('F')} != {'found', 'not_found'}:
                continue
            if len(left_tokens) < 3 or len(right_tokens) != len(left_tokens):
                continue
            if left_tokens[:-1] != right_tokens[:-1] or left_tokens[-1] == right_tokens[-1]:
                continue
            first_show_index = min(pairs[left_at][0], pairs[right_at][0])
            preceding = rows[max(0, first_show_index - 5):first_show_index]
            flat = [command for row in preceding for command in _commands(row)]
            if not any((command.casefold().startswith('no ') for command in flat)):
                continue
            if not any((not command.casefold().startswith('no ') and ' virtual ' in command.casefold() for command in flat)):
                continue
            receipt = _pick_receipt(receipts)
            if receipt is None:
                continue
            return ({'detector': 'same_identity_found_in_one_typed_list_and_absent_in_another', 'rows': [left_show['row'], left_assert['row'], right_show['row'], right_assert['row']], 'typed_show_list_count': 2, 'uses_traffic_as_type_identity_proof': False}, receipt)
    return None

def _per_object_associations(rows: list[dict[str, Any]], receipts: list[dict[str, Any]]) -> tuple[dict[str, Any], dict[str, Any]] | None:
    pairs: list[tuple[dict[str, Any], dict[str, Any], str, str]] = []
    for index in range(len(rows) - 1):
        observe, assertion = (rows[index], rows[index + 1])
        commands = _commands(observe)
        pattern = _text(assertion.get('G'))
        if len(commands) == 1 and commands[0].casefold().startswith('show ') and (assertion.get('E') == 'check_point') and (assertion.get('F') == 'found') and (pattern.count('.*') == 1):
            left, right = (part.strip() for part in pattern.split('.*', 1))
            if left and right:
                pairs.append((observe, assertion, left, right))
    unique = {(left, right) for _, _, left, right in pairs}
    if len(unique) < 3:
        return None
    receipt = _pick_receipt(receipts, lambda actual: all((left in actual and right in actual for left, right in unique)))
    if receipt is None:
        return None
    return ({'detector': 'distinct_per_object_association_assertions', 'rows': [item[1]['row'] for item in pairs], 'association_assertion_count': len(unique), 'single_aggregate_existence_assertion': False}, receipt)
_IPV4_TOKEN = re.compile('(?<![0-9])(?:25[0-5]|2[0-4][0-9]|1?[0-9]{1,2})(?:\\.(?:25[0-5]|2[0-4][0-9]|1?[0-9]{1,2})){3}(?![0-9])')

def _ipv4_tokens(text: str) -> set[str]:
    return set(_IPV4_TOKEN.findall(text))

def _multi_device_forwarding_business_response(rows: list[dict[str, Any]], receipts: list[dict[str, Any]]) -> tuple[dict[str, Any], dict[str, Any]] | None:
    config_methods = {'cmd_config', 'cmds_config', 'cmd_enable'}
    for observe_at, observe in enumerate(rows):
        if observe.get('E') != 'test_env' or not _commands(observe):
            continue
        assertion_at = next((index for index in range(observe_at + 1, min(observe_at + 3, len(rows))) if rows[index].get('E') == 'check_point' and rows[index].get('F') in {'found', 'abs_found'} and _text(rows[index].get('G'))), None)
        if assertion_at is None:
            continue
        assertion = rows[assertion_at]
        pattern = _text(assertion.get('G'))
        if 'IST_EXIT_STATUS' in pattern:
            continue
        normalized_pattern = pattern.replace('\\.', '.')
        probe_ips = _ipv4_tokens('\n'.join(_commands(observe)))
        if not probe_ips:
            continue
        configs: dict[str, list[tuple[dict[str, Any], str]]] = {}
        for row in rows[:observe_at]:
            host = _text(row.get('E'))
            if not host.startswith('APV_') or row.get('F') not in config_methods:
                continue
            for command in _commands(row):
                if not command.casefold().startswith('no '):
                    configs.setdefault(host, []).append((row, command))
        if len(configs) < 2:
            continue
        for front_host, front_rows in configs.items():
            front_text = '\n'.join((command for _, command in front_rows))
            if not probe_ips.intersection(_ipv4_tokens(front_text)):
                continue
            for producer_host, producer_rows in configs.items():
                if producer_host == front_host:
                    continue
                if len(producer_rows) < 5 or len(front_rows) < 2:
                    continue
                producer_text = '\n'.join((command for _, command in producer_rows))
                producer_ips = _ipv4_tokens(producer_text)
                linked_ips = producer_ips.intersection(_ipv4_tokens(front_text))
                response_ips = {value for value in producer_ips if value in normalized_pattern and value not in linked_ips}
                if not linked_ips or not response_ips:
                    continue
                wait_rows = [row for row in rows[:observe_at] if row.get('E') == 'time' and row.get('F') == 'sleep' and (row['row'] > max((item[0]['row'] for item in producer_rows + front_rows)))]
                if not wait_rows:
                    continue
                receipt = _pick_receipt(receipts, lambda actual: 'IST_EXIT_STATUS' in actual and (not any((value in actual for value in response_ips)))) or _pick_receipt(receipts)
                if receipt is None:
                    continue
                config_rows = sorted({item[0]['row'] for item in producer_rows + front_rows})
                return ({'detector': 'two_device_forwarding_chain_with_business_response', 'rows': [*config_rows, wait_rows[-1]['row'], observe['row'], assertion['row']], 'device_role_count': 2, 'producer_config_step_count': len(producer_rows), 'forwarding_config_step_count': len(front_rows), 'stable_wait_before_probe': True, 'probe_targets_forwarding_listener': True, 'assertion_uses_business_response_not_process_exit': True, 'failed_route_contrast_omits_business_response': not any((value in str(receipt.get('_text') or '') for value in response_ips))}, receipt)
    return None
RULES: tuple[dict[str, Any], ...] = ({'id': 'rebuild_requires_enable_and_stable_wait', 'kind': 'device_behavior_fact', 'statement': 'After an object is rebuilt, enable it and allow a stable wait before probing it again; object creation alone does not establish readiness.', 'detector': _rebuild_enable_wait}, {'id': 'negative_outcome_uses_actual_feedback_or_structure', 'kind': 'translation_example', 'statement': "Bind a negative outcome to feedback actually emitted by the device or framework (for example, 'Failed to execute the command') or to an observable structural consequence; never treat the author's sentence as device bytes.", 'detector': _negative_actual_feedback}, {'id': 'protocol_change_proved_by_typed_object_existence', 'kind': 'device_behavior_fact', 'statement': 'Prove a protocol or type change through object presence and absence in the corresponding typed show lists, not through traffic behavior.', 'detector': _typed_list_existence}, {'id': 'multi_instance_success_asserts_each_association', 'kind': 'translation_example', 'statement': "For a multi-instance 'configuration succeeded' criterion, when the observation surface distinguishes objects, assert each object's own association fields as one dimension per object. When the surface cannot distinguish objects, an aggregate assertion is allowed only with an explicit coverage disclosure.", 'detector': _per_object_associations, 'adjudication_entry_id': 'multi_instance_success_asserts_each_association'}, {'id': 'save_family_uses_argumented_form_and_equivalent_observation', 'kind': 'translation_example', 'statement': 'Treat an authored persistence/save family word as an intent family: preserve that family while grounding a tree-legal argumented form in the manual, bind its required target as self-authored setup data, and replace a bed-forbidden mechanism with an equivalent non-destructive observation. If bare-form semantics are the test point, keep the claim underdetermined instead of silently rewriting it.', 'detector': _argumented_save_family_equivalent_observation, 'source_mode': 'harvest_verified'}, {'id': 'multi_device_forwarding_requires_complete_topology', 'kind': 'translation_example', 'statement': 'For cross-device forwarding validation, construct the complete producer-side object chain and the front-side listener/forwarding relation before probing the front listener. Assert the protocol or business response, such as returned answer data, rather than only the client process exit status.', 'detector': _multi_device_forwarding_business_response})

def _ledger_receipt_rows(path: Path) -> list[dict[str, Any]]:
    from cex_core.engine.ist_core.security_scrub import scrub_text
    try:
        text = path.read_text(encoding='utf-8', errors='replace')
    except OSError:
        return []
    rows: list[dict[str, Any]] = []
    for line in text.splitlines():
        try:
            fact = json.loads(line)
        except json.JSONDecodeError:
            continue
        if not isinstance(fact, dict) or fact.get('ev') != 'verdict':
            continue
        inline = fact.get('evidence_inline')
        if not isinstance(inline, dict):
            continue
        row = {key: scrub_text(value, scrub_paths=False) if isinstance(value, str) else value for key, value in inline.items()}
        row['autoid'] = _text(fact.get('aid') or fact.get('autoid'))
        row['verdict'] = _text(fact.get('result') or fact.get('verdict'))
        if _receipt_row_valid(row):
            rows.append(row)
    return rows

def _rule_required_role(rule: dict[str, Any]) -> str:
    return 'harvest_workbook' if rule.get('source_mode') == 'harvest_verified' else 'attribution_receipt'

def _seal_entry(relpath: str, role: str, sha256: str, *, autoid: str='', origin: str='') -> dict[str, Any]:
    entry: dict[str, Any] = {'path': relpath, 'role': role, 'sha256': sha256}
    if autoid:
        entry['autoid'] = autoid
    if origin:
        entry['origin'] = origin
    return entry

def _qualifying_adjudication(path: Path) -> bool:
    try:
        text = path.read_text(encoding='utf-8')
    except (OSError, UnicodeError):
        return False
    match = re.match('^---\\n(.*?)\\n---\\n', text, re.DOTALL)
    if match is None:
        return False
    try:
        frontmatter = yaml.safe_load(match.group(1)) or {}
    except yaml.YAMLError:
        return False
    return bool(_text(frontmatter.get('projection_entry_id'))) and _text(frontmatter.get('channel')) == 'ordered_user_adjudications'

def seal_corpus(root: Path=ROOT, *, corpus_root: Path | None=None) -> dict[str, Any]:
    corpus = corpus_root or root / CORPUS_ROOT_REL
    staged: list[tuple[Path, dict[str, Any]]] = []
    verified_root = root / 'knowledge/framework/verified'
    for path in sorted(verified_root.glob('verified_*.xlsx')):
        match = re.fullmatch('verified_(\\d+)\\.xlsx', path.name)
        if match is None:
            continue
        staged.append((path, _seal_entry(f'verified/{path.name}', 'verified_workbook', _sha256(path), autoid=match.group(1), origin=_rel_posix(path, root))))
    if not any((entry['role'] == 'verified_workbook' for _, entry in staged)):
        raise DeviceBehaviorSourceUnavailableError(f'refuse to seal an empty verified corpus: {verified_root} has no verified_*.xlsx')
    outputs_root = root / 'workspace/outputs'
    for path in sorted(outputs_root.glob('*/attribution_evidence_holds/*.json')):
        try:
            payload = json.loads(path.read_text(encoding='utf-8'))
        except (OSError, UnicodeError, json.JSONDecodeError) as exc:
            raise DeviceBehaviorSourceUnavailableError(f'refuse to seal unreadable attribution receipt: {path}: {exc}') from exc
        rows = payload if isinstance(payload, list) else [payload]
        if not any((_receipt_row_valid(row) for row in rows)):
            raise DeviceBehaviorSourceUnavailableError(f'refuse to seal attribution receipt without identity-bound rows: {path}')
        staged.append((path, _seal_entry(f'receipts/{path.parent.parent.name}__{path.name}', 'attribution_receipt', _sha256(path), origin=_rel_posix(path, root))))
    for path in sorted(outputs_root.glob('*/facts.jsonl')):
        ledger_rows = _ledger_receipt_rows(path)
        if not ledger_rows:
            continue
        minted = json.dumps(ledger_rows, ensure_ascii=False, indent=1) + '\n'
        staged.append((minted, _seal_entry(f'receipts/ledger__{path.parent.name}.json', 'attribution_receipt', hashlib.sha256(minted.encode('utf-8')).hexdigest(), origin=_rel_posix(path, root))))
    verified_autoids = {entry['autoid'] for _, entry in staged if entry['role'] == 'verified_workbook' and entry.get('autoid')}
    mirror = root / 'knowledge/framework/mirror/smoke_test'
    sealed_harvest: set[str] = set()
    for path in sorted(mirror.glob('**/case.xlsx')):
        autoid = path.parent.name
        if not re.fullmatch('\\d{18}', autoid) or autoid in sealed_harvest:
            continue
        if autoid not in verified_autoids:
            continue
        sealed_harvest.add(autoid)
        staged.append((path, _seal_entry(f'harvest/{autoid}.xlsx', 'harvest_workbook', _sha256(path), autoid=autoid, origin=_rel_posix(path, root))))
    adjudication_root = root / 'knowledge/adjudications'
    skipped_adjudications = 0
    if adjudication_root.is_dir():
        for path in sorted(adjudication_root.glob('*.md')):
            if not _qualifying_adjudication(path):
                skipped_adjudications += 1
                continue
            staged.append((path, _seal_entry(f'adjudications/{path.name}', 'user_adjudication', _sha256(path), origin=_rel_posix(path, root))))
    grammar_path = root / 'knowledge/data/compile_ref/domain_grammar.json'
    try:
        grammar = json.loads(grammar_path.read_text(encoding='utf-8'))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise DeviceBehaviorSourceUnavailableError(f'refuse to seal without a readable domain_grammar: {grammar_path}: {exc}') from exc
    if not isinstance(grammar, dict):
        raise DeviceBehaviorSourceUnavailableError(f'refuse to seal a non-object domain_grammar: {grammar_path}')
    staged.append((grammar_path, _seal_entry('grammar/domain_grammar.json', 'domain_grammar', _sha256(grammar_path), origin=_rel_posix(grammar_path, root))))
    corpus.parent.mkdir(parents=True, exist_ok=True)
    staging = corpus.parent / f'.{corpus.name}.seal-staging'
    legacy = corpus.parent / f'.{corpus.name}.pre-seal'
    shutil.rmtree(staging, ignore_errors=True)
    shutil.rmtree(legacy, ignore_errors=True)
    files = sorted((entry for _, entry in staged), key=lambda item: item['path'])
    try:
        for source, entry in staged:
            target = staging / entry['path']
            target.parent.mkdir(parents=True, exist_ok=True)
            if isinstance(source, Path):
                shutil.copyfile(source, target)
            else:
                target.write_text(source, encoding='utf-8')
        seal_payload = {'schema': CORPUS_SEAL_SCHEMA, 'files': files}
        (staging / CORPUS_SEAL_NAME).write_text(json.dumps(seal_payload, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        if corpus.exists():
            os.replace(corpus, legacy)
        os.replace(staging, corpus)
    finally:
        shutil.rmtree(staging, ignore_errors=True)
    shutil.rmtree(legacy, ignore_errors=True)
    counts: dict[str, int] = {}
    for entry in files:
        counts[entry['role']] = counts.get(entry['role'], 0) + 1
    return {'corpus_root': corpus.as_posix(), 'seal': (corpus / CORPUS_SEAL_NAME).as_posix(), 'schema': CORPUS_SEAL_SCHEMA, 'file_count': len(files), 'counts': counts, 'skipped_adjudication_files': skipped_adjudications}

def build(root: Path=ROOT, *, corpus_root: Path | None=None) -> dict[str, Any]:
    corpus = _load_corpus(corpus_root or root / CORPUS_ROOT_REL, root)
    by_role = corpus['by_role']
    receipts_by_autoid = _load_receipts(by_role['attribution_receipt'])
    adjudications = _translation_adjudications(by_role['user_adjudication'])
    harvest_by_autoid, harvest_unreadable = _harvest_workbooks(by_role['harvest_workbook'])
    grammar = _corpus_grammar(corpus)
    candidates: dict[str, list[dict[str, Any]]] = {rule['id']: [] for rule in RULES}
    shape_candidates: dict[str, list[dict[str, Any]]] = {}
    unreadable: list[str] = list(harvest_unreadable)
    verified_entries = by_role['verified_workbook']
    for entry in verified_entries:
        workbook_path = entry['path']
        match = re.fullmatch('verified_(\\d+)\\.xlsx', workbook_path.name)
        if match is None:
            continue
        autoid = match.group(1)
        try:
            shape_axis, shape_rows = _verified_translation_features(workbook_path, autoid)
            shape_signature = _translation_signature(shape_axis)
            shape_summary = _translation_shape_summary(shape_rows)
            shape_candidates.setdefault(shape_signature, []).append({'workbook': {'path': entry['relpath'], 'sha256': entry['sha256']}, 'axis': shape_axis, 'summary': shape_summary, 'score': (len(shape_rows), int(shape_summary['lifecycle_shape']['config_row_count']), len(shape_summary['topology_roles']['product_devices']))})
            rows = _workbook_rows(workbook_path)
            rows_scoped = _workbook_rows(workbook_path, autoid=autoid)
        except Exception as exc:
            unreadable.append(f"{entry['relpath']}:{type(exc).__name__}")
            continue
        workbook_source = {'path': entry['relpath'], 'sha256': entry['sha256']}
        for rule in RULES:
            if rule.get('source_mode') == 'harvest_verified':
                harvest = harvest_by_autoid.get(autoid)
                if not harvest:
                    continue
                witness = rule['detector'](rows_scoped, harvest['rows'], grammar)
                if witness is None:
                    continue
                candidates[rule['id']].append({'workbook': workbook_source, 'witness': witness, 'harvest': harvest['source']})
                continue
            receipts = receipts_by_autoid.get(autoid) or []
            if not receipts:
                continue
            detected = rule['detector'](rows, receipts)
            if detected is None:
                continue
            witness, receipt = detected
            candidates[rule['id']].append({'workbook': workbook_source, 'witness': witness, 'receipt': _source_receipt(receipt)})
    entries: list[dict[str, Any]] = []
    missing: list[str] = []
    for rule in RULES:
        rows = candidates[rule['id']]
        if not rows:
            missing.append(rule['id'])
            continue
        adjudication_id = _text(rule.get('adjudication_entry_id'))
        adjudication = adjudications.get(adjudication_id) if adjudication_id else None
        if adjudication_id and adjudication is None:
            missing.append(f"{rule['id']}:ordered_user_adjudication")
            continue
        chosen = sorted(rows, key=lambda item: (item['workbook']['path'], tuple(item['witness'].get('rows') or []), (item.get('receipt') or item.get('harvest') or {}).get('path', '')))[0]
        entries.append({'id': rule['id'], 'kind': rule['kind'], 'statement': rule['statement'], 'mechanical_enforcement': False, 'authority': {'class': 'translation_shape_only' if adjudication else 'actual_only', 'knowledge_layer': 'K_ought' if adjudication else 'K_is', 'can_sign_expected': False, 'note': 'This entry guides translation shape and is not an expected-value signer. Device actuals remain diagnostic; the ordered user ruling does not turn them into expected bytes.' if adjudication else 'This entry is a device-behavior or translation example. It is not Author, Spec, DefectSpec, Manual, ConfigBinding, or CapabilityXml.'}, 'derivation': chosen['witness'], 'sources': {'verified_workbook': chosen['workbook'], **({'attribution_receipt': chosen['receipt']} if 'receipt' in chosen else {'harvest_workbook': chosen['harvest']}), **({'user_adjudication': adjudication} if adjudication else {})}})
    corpus_role_counts = {role: len(by_role.get(role) or []) for role in sorted(CORPUS_ROLES)}
    rule_by_id = {rule['id']: rule for rule in RULES}
    unavailable_causes: dict[str, str] = {}
    for key in missing:
        if key.endswith(':ordered_user_adjudication'):
            role = 'user_adjudication'
        else:
            role = _rule_required_role(rule_by_id[key])
        unavailable_causes[key] = 'no_qualifying_candidate' if corpus_role_counts.get(role) else 'evidence_class_absent'
    absent_classes = sorted({'user_adjudication' if key.endswith(':ordered_user_adjudication') else _rule_required_role(rule_by_id[key]) for key, cause in unavailable_causes.items() if cause == 'evidence_class_absent'})
    status = 'ready' if not missing else 'unavailable'
    published = f"{len(entries)} entry/entries are published and usable on their own ({', '.join((row['id'] for row in entries))}); " if entries else ''
    if not missing:
        unavailable_reason = ''
    elif absent_classes:
        unavailable_reason = published + 'The sealed corpus carries no evidence at all of these classes: ' + ', '.join(absent_classes) + '. These entries cannot close until a corpus is sealed while that evidence still exists on the live surface (attribution receipts are short-lived: closing deletes the holds). This is a corpus that was sealed too late, not a device that never showed the behavior — reseal with `python scripts/gen_device_behavior_examples.py seal`.'
    else:
        unavailable_reason = published + 'A sealed, sha256-verified corpus was searched and no identity-bound candidate closed these required entries. This is a true absence of qualifying evidence in a verified corpus, not an unavailable evidence surface: an unavailable corpus fails the generator and emits no projection at all.'
    translation_shapes: list[dict[str, Any]] = []
    for signature, rows in sorted(shape_candidates.items()):
        chosen = sorted(rows, key=lambda item: (tuple((-int(value) for value in item['score'])), item['workbook']['path']))[0]
        axis = {public_name: bool(chosen['axis'][source_key]) for public_name, source_key in TRANSLATION_SHAPE_AXES}
        translation_shapes.append({'id': 'verified_shape_' + hashlib.sha256(signature.encode('utf-8')).hexdigest()[:16], 'kind': 'translation_example', 'statement': _shape_statement(chosen['summary']), 'mechanical_enforcement': False, 'authority': {'class': 'translation_shape_only', 'knowledge_layer': 'K_is', 'can_sign_expected': False, 'note': 'This entry is an epoch-revalidated translation shape only. It carries no command recipe or expected value and never authorizes observe-then-assert.'}, 'derivation': {'detector': 'verified_six_axis_translation_signature', 'signature': signature, 'axes': axis, 'axis_note': 'These six booleans reproduce the disclosure-study workbook-surface signature, including legacy side columns and the sentinel row. They are opaque clustering identity, not claims that the case executed those actions. The executable topology/trigger/assertion/naming summary is translation_shape below.', 'translation_shape': chosen['summary']}, 'sources': {'verified_workbook': chosen['workbook']}})
    return {'schema': 'ist.device-behavior-examples.v1', 'status': status, 'coverage': {'required': [rule['id'] for rule in RULES], 'present': [row['id'] for row in entries], 'missing': missing, 'note': 'Entries listed in `entries` are closed and usable on their own. `status` reports whether the required set is complete; a partial projection still supplies every entry it did close.'}, 'unavailable_causes': unavailable_causes, 'absent_evidence_classes': absent_classes, 'unavailable_reason': unavailable_reason, 'missing_required_entries': missing, 'verified_corpus': {'path': _rel_posix(corpus['root'], root), 'seal': _rel_posix(corpus['seal_path'], root), 'seal_schema': CORPUS_SEAL_SCHEMA, 'file_count': len(verified_entries), 'corpus_file_count': len(corpus['files']), 'role_counts': corpus_role_counts, 'source_manifest_sha256': corpus['seal_sha256'], 'reseal': 'python scripts/gen_device_behavior_examples.py seal'}, 'authority': {'class': 'actual_only', 'can_sign_expected': False, 'expected_authorities': ['Author', 'Spec', 'DefectSpec', 'Manual', 'ConfigBinding', 'CapabilityXml']}, 'regenerate': 'python scripts/gen_device_behavior_examples.py', 'source_policy': 'All evidence is read exclusively from the sealed corpus pinned by corpus.seal.json (ist.device-behavior-corpus-seal.v1); live runtime surfaces (framework mirror, workspace outputs) are never consulted at build time, and a missing, unreadable, or identity-mismatched corpus fails the generator instead of emitting an empty projection. Actual device-behavior entries are structurally derived only when a case identity has both a verified workbook and an attribution receipt; the persistence-family translation exemplar instead requires the same identity in the framework harvest and verified corpus. Source payload text is not copied, except for a literal proven to reappear in device/framework actuals. A translation-shape entry may additionally bind an identity-marked K_ought record from ordered_user_adjudications; no entry mechanically enforces an assertion count or signs expected. The translation_shape_vocabulary is recomputed directly from the sealed verified workbooks and does not require a failure-attribution receipt; it copies only topology/trigger/assertion/naming shape and never source commands or expected literals.', 'unreadable_sources': unreadable, 'entries': entries, 'translation_shape_vocabulary': {'schema': 'ist.translation-shape-vocabulary', 'axes': [name for name, _source_key in TRANSLATION_SHAPE_AXES], 'signature_count': len(translation_shapes), 'entries': translation_shapes}}

def render(payload: dict[str, Any]) -> str:
    return json.dumps(payload, ensure_ascii=False, indent=2) + '\n'

def write_atomic(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.is_symlink():
        raise DeviceBehaviorGenerationError('output path is a symbolic link')
    fd, temporary = tempfile.mkstemp(prefix=f'.{path.name}.', dir=path.parent)
    try:
        with os.fdopen(fd, 'w', encoding='utf-8') as stream:
            stream.write(render(payload))
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    except Exception:
        Path(temporary).unlink(missing_ok=True)
        raise

def main(argv: list[str] | None=None, *, root: Path=ROOT) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument('--check', action='store_true')
    parser.add_argument('--output', type=Path, default=None)
    parser.add_argument('--corpus-root', type=Path, default=None)
    subparsers = parser.add_subparsers(dest='command')
    seal_parser = subparsers.add_parser('seal', help='把现役活面证据固化成密封语料（构建只读语料）')
    seal_parser.add_argument('--corpus-root', type=Path, default=None)
    args = parser.parse_args(argv)
    try:
        if args.command == 'seal':
            summary = seal_corpus(root, corpus_root=args.corpus_root)
            print(json.dumps(summary, ensure_ascii=False, indent=2))
            return 0
        payload = build(root, corpus_root=args.corpus_root)
    except DeviceBehaviorSourceUnavailableError as exc:
        print(f'device-behavior corpus unavailable: {exc}', file=sys.stderr)
        return 2
    output = args.output or root / 'knowledge/data/compile_ref/device_behavior_examples.json'
    generated = render(payload)
    if args.check:
        try:
            current = output.read_text(encoding='utf-8')
        except OSError:
            return 1
        return 0 if current == generated else 1
    write_atomic(output, payload)
    return 0
if __name__ == '__main__':
    raise SystemExit(main())
