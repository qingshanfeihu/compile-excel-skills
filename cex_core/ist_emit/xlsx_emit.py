"""xlsx 出件（skill 剪切版）。

上游来源：InfoTest_Engine main/case_compiler/xlsx_emit.py。
保留：fd 级原子写盘（staging inode + 硬链备份 + 回滚 + 读回复核）、
目录绑定校验（O_NOFOLLOW/O_DIRECTORY 全链）、凭据 redact。
砍掉：run_governance.run_owner_commit_scope（不随包分发）；
excel_release.py 晋升链（模板 = 内置冻结快照 + SHA 钉死，选择逻辑收进本文件）。
行为差异声明见 reference/excel-contract.md。
"""

from __future__ import annotations

import io
import os
import re
import secrets
import stat
from dataclasses import dataclass
from pathlib import Path
from typing import Any, BinaryIO, Optional

from openpyxl import load_workbook

from ._sealed_io import (
    read_regular_nofollow,
    sha256_bytes,
    validate_xlsx_zip_budget,
)
from .case_ir import FileIR, Row
from .config import get_config
from .credential_literals import mirror_credential_literals
from .excel_contract import (
    PINNED_CONTRACT_SHA256,
    TEMPLATE_SHA256,
    ExcelContractError,
    resolve_execution_sheet,
)

_TARGET_NAME = "case.xlsx"
_TMP_PREFIX = ".case.xlsx.tmp."
_BACKUP_PREFIX = ".case.xlsx.previous."
_REDACTED_TEMPLATE_CREDENTIAL = "<已移除凭据>"

TEMPLATE_PATH = Path(__file__).resolve().parents[1] / "templates" / "case_template.xlsx"


@dataclass(frozen=True)
class RuntimeTemplateSelection:

    path: Path
    content: bytes
    release_state: str
    contract_sha256: str | None = None


def select_runtime_template() -> RuntimeTemplateSelection:
    """内置冻结快照模板 + SHA 钉死。

    skill 版不存在 candidate/legacy 通道：模板就是 templates/case_template.xlsx，
    SHA 必须等于 excel_contract.TEMPLATE_SHA256。模板身份与晋升回执见
    reference/excel-contract.md（冻结快照，非晋升链）。
    """
    content = read_regular_nofollow(
        TEMPLATE_PATH,
        error_type=ExcelContractError,
        invalid_message="runtime template path is invalid",
        directory_message="runtime template directory is unavailable",
        open_message="runtime template is unavailable",
        bounds_message="runtime template is not a bounded regular file",
        changed_message="runtime template changed while reading",
        max_bytes=8 * 1024 * 1024,
    )
    validate_xlsx_zip_budget(
        content,
        error_type=ExcelContractError,
        message="runtime template is not a bounded xlsx archive",
    )
    if sha256_bytes(content) != TEMPLATE_SHA256:
        raise ExcelContractError(
            "runtime template sha256 does not match the pinned identity "
            f"(expected {TEMPLATE_SHA256[:12]}…)"
        )
    return RuntimeTemplateSelection(
        path=TEMPLATE_PATH,
        content=content,
        release_state="promoted",
        contract_sha256=PINNED_CONTRACT_SHA256,
    )


def _inode_id(st: os.stat_result) -> tuple[int, int]:
    return int(st.st_dev), int(st.st_ino)


def _directory_flags() -> int:
    if not hasattr(os, "O_NOFOLLOW") or not hasattr(os, "O_DIRECTORY"):
        raise RuntimeError("secure xlsx emit requires O_NOFOLLOW and O_DIRECTORY")
    return (
        os.O_RDONLY
        | os.O_DIRECTORY
        | os.O_NOFOLLOW
        | getattr(os, "O_CLOEXEC", 0)
    )


def _open_directory_chain(path: Path, *, create: bool = False) -> int:
    absolute = Path(os.path.abspath(os.fspath(path)))
    if not absolute.is_absolute():
        raise ValueError("trusted outputs root must be absolute")
    flags = _directory_flags()
    fd = os.open(absolute.anchor, flags)
    try:
        for part in absolute.parts[1:]:
            try:
                next_fd = os.open(part, flags, dir_fd=fd)
            except FileNotFoundError:
                if not create:
                    raise
                try:
                    os.mkdir(part, mode=0o700, dir_fd=fd)
                except FileExistsError:
                    pass
                next_fd = os.open(part, flags, dir_fd=fd)
            os.close(fd)
            fd = next_fd
        return fd
    except BaseException:
        os.close(fd)
        raise


def _open_or_create_case_directory(root_fd: int, component: str) -> int:
    flags = _directory_flags()
    try:
        return os.open(component, flags, dir_fd=root_fd)
    except FileNotFoundError:
        try:
            os.mkdir(component, mode=0o700, dir_fd=root_fd)
        except FileExistsError:
            pass
        return os.open(component, flags, dir_fd=root_fd)


def _assert_directory_binding(
    trusted_root: Path,
    root_fd: int,
    component: str,
    case_dir_fd: int,
) -> None:
    check_root_fd = _open_directory_chain(trusted_root)
    check_case_fd = -1
    try:
        if _inode_id(os.fstat(check_root_fd)) != _inode_id(os.fstat(root_fd)):
            raise RuntimeError("trusted outputs root was exchanged during xlsx emit")
        check_case_fd = os.open(
            component, _directory_flags(), dir_fd=check_root_fd)
        if _inode_id(os.fstat(check_case_fd)) != _inode_id(os.fstat(case_dir_fd)):
            raise RuntimeError("case output directory was exchanged during xlsx emit")
    except (FileNotFoundError, NotADirectoryError, OSError) as exc:
        raise RuntimeError(
            "case output directory binding changed during xlsx emit") from exc
    finally:
        if check_case_fd >= 0:
            os.close(check_case_fd)
        os.close(check_root_fd)


def _validate_existing_target(case_dir_fd: int) -> os.stat_result | None:
    try:
        st = os.stat(
            _TARGET_NAME, dir_fd=case_dir_fd, follow_symlinks=False)
    except FileNotFoundError:
        return None
    if stat.S_ISLNK(st.st_mode):
        raise ValueError("case.xlsx target must not be a symbolic link")
    if not stat.S_ISREG(st.st_mode):
        raise ValueError("case.xlsx target must be a regular file")
    if int(st.st_nlink) != 1:
        raise ValueError("case.xlsx hard-link target is rejected")
    return st


def _unique_name(prefix: str) -> str:
    return prefix + secrets.token_hex(16)


def _create_new_file(case_dir_fd: int) -> tuple[str, int]:
    flags = (
        os.O_WRONLY
        | os.O_CREAT
        | os.O_EXCL
        | os.O_NOFOLLOW
        | getattr(os, "O_CLOEXEC", 0)
    )
    for _ in range(32):
        name = _unique_name(_TMP_PREFIX)
        try:
            return name, os.open(name, flags, 0o600, dir_fd=case_dir_fd)
        except FileExistsError:
            continue
    raise RuntimeError("could not allocate a unique xlsx staging inode")


def _make_target_backup(
    case_dir_fd: int,
) -> tuple[str, os.stat_result] | None:
    try:
        old_fd = os.open(
            _TARGET_NAME,
            os.O_RDONLY | os.O_NOFOLLOW | getattr(os, "O_CLOEXEC", 0),
            dir_fd=case_dir_fd,
        )
    except FileNotFoundError:
        return None
    try:
        old_st = os.fstat(old_fd)
        if not stat.S_ISREG(old_st.st_mode):
            raise ValueError("case.xlsx target must be a regular file")
        if int(old_st.st_nlink) != 1:
            raise ValueError("case.xlsx hard-link target is rejected")
        for _ in range(32):
            backup = _unique_name(_BACKUP_PREFIX)
            try:
                os.link(
                    _TARGET_NAME,
                    backup,
                    src_dir_fd=case_dir_fd,
                    dst_dir_fd=case_dir_fd,
                    follow_symlinks=False,
                )
                break
            except FileExistsError:
                continue
        else:
            raise RuntimeError("could not allocate a unique xlsx rollback link")

        try:
            backup_st = os.stat(
                backup, dir_fd=case_dir_fd, follow_symlinks=False)
            current_st = os.stat(
                _TARGET_NAME, dir_fd=case_dir_fd, follow_symlinks=False)
            expected = _inode_id(old_st)
            if (
                _inode_id(backup_st) != expected
                or _inode_id(current_st) != expected
                or not stat.S_ISREG(current_st.st_mode)
                or int(backup_st.st_nlink) != 2
                or int(current_st.st_nlink) != 2
            ):
                raise RuntimeError("case.xlsx target changed while preparing atomic replace")
            return backup, old_st
        except BaseException:
            try:
                os.unlink(backup, dir_fd=case_dir_fd)
            except OSError:
                pass
            raise
    finally:
        os.close(old_fd)


def _same_target_inode(case_dir_fd: int, expected: os.stat_result) -> bool:
    try:
        current = os.stat(
            _TARGET_NAME, dir_fd=case_dir_fd, follow_symlinks=False)
    except OSError:
        return False
    return (
        stat.S_ISREG(current.st_mode)
        and _inode_id(current) == _inode_id(expected)
    )


def _rollback_replaced_target(
    case_dir_fd: int,
    *,
    backup_name: str | None,
    new_inode: os.stat_result,
) -> None:
    if backup_name is not None:
        os.replace(
            backup_name,
            _TARGET_NAME,
            src_dir_fd=case_dir_fd,
            dst_dir_fd=case_dir_fd,
        )
        return
    if _same_target_inode(case_dir_fd, new_inode):
        os.unlink(_TARGET_NAME, dir_fd=case_dir_fd)


def _validate_destination(
    out_path: Path,
    trusted_outputs_root: Path,
) -> tuple[Path, str]:
    root = Path(os.path.abspath(os.fspath(trusted_outputs_root)))
    target = Path(os.path.abspath(os.fspath(out_path)))
    try:
        rel = target.relative_to(root)
    except ValueError as exc:
        raise ValueError("case.xlsx destination is outside trusted outputs root") from exc
    if (
        len(rel.parts) != 2
        or rel.parts[0] in {"", ".", ".."}
        or rel.parts[1] != _TARGET_NAME
    ):
        raise ValueError(
            "case.xlsx destination must be outputs/<single-component>/case.xlsx")
    return root, rel.parts[0]


def _set_row(sheet, r: int, cols: dict[int, Any]) -> None:
    for c in range(1, 10):
        cell = sheet.cell(r, c)
        value = cols.get(c)
        cell.value = value
        if c == 1 and value not in (None, ""):
            cell.value = str(value)
            cell.number_format = "@"


def _sanitize_inherited_template_credentials(workbook) -> None:
    literals = sorted(
        (value for value in mirror_credential_literals() if value),
        key=len,
        reverse=True,
    )
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if not isinstance(value, str):
                    continue
                cleaned = value
                for literal in literals:
                    cleaned = re.sub(
                        re.escape(literal),
                        _REDACTED_TEMPLATE_CREDENTIAL,
                        cleaned,
                        flags=re.IGNORECASE,
                    )
                if cleaned != value:
                    cell.value = cleaned


def _row_cols(row: Row, *, stmt_type: Optional[int], description: Optional[str],
              autoid: Optional[str] = None, priority: Optional[str] = None) -> dict[int, Any]:
    cols: dict[int, Any] = {}
    if autoid is not None:
        cols[1] = autoid
    if priority is not None:
        cols[2] = priority
    if stmt_type is not None:
        cols[3] = stmt_type
    if description is not None:
        cols[4] = description
    cols[5] = row.test_object
    cols[6] = row.method
    cols[7] = row.data
    if row.save_as:
        cols[8] = row.save_as
    if row.input_var:
        cols[9] = row.input_var
    return cols


def _build_workbook(
    file_ir: FileIR,
    *,
    selection: RuntimeTemplateSelection | None = None,
):
    selection = selection or select_runtime_template()
    wb = load_workbook(io.BytesIO(selection.content))
    _sanitize_inherited_template_credentials(wb)
    sheet, layout = resolve_execution_sheet(
        wb,
        allow_legacy=False,
    )
    current_autoid_width = float(sheet.column_dimensions["A"].width or 0)
    sheet.column_dimensions["A"].width = max(current_autoid_width, 22.0)

    data_start = layout.data_start

    for r in range(data_start, sheet.max_row + 1):
        _set_row(sheet, r, {})

    r = data_start
    _set_row(sheet, r, {3: 0, 4: f"Author         : {file_ir.author}\n{file_ir.feature}"})
    r += 1

    for row in file_ir.init_rows:
        _set_row(sheet, r, _row_cols(row, stmt_type=1, description="初始化配置"))
        r += 1

    for case in file_ir.cases:
        first_case_row = True
        for st in case.steps:
            first_row_of_step = True
            for row in st.rows:
                if first_row_of_step:
                    cols = _row_cols(
                        row,
                        stmt_type=st.stmt_type,
                        description=st.description,
                        autoid=(case.autoid if first_case_row else None),
                        priority=(case.priority if first_case_row else None),
                    )
                    first_case_row = False
                    first_row_of_step = False
                else:
                    cols = _row_cols(row, stmt_type=None, description=None)
                _set_row(sheet, r, cols)
                r += 1
        r += 1
    return wb


def emit_xlsx(
    file_ir: FileIR,
    out_path: Path,
    *,
    trusted_outputs_root: Path,
) -> dict:
    out_path = Path(out_path)
    trusted_root, component = _validate_destination(
        out_path, Path(trusted_outputs_root))
    selection = select_runtime_template()
    root_fd = _open_directory_chain(trusted_root, create=True)
    case_dir_fd = -1
    temp_fd = -1
    temp_name: str | None = None
    backup_name: str | None = None
    new_inode: os.stat_result | None = None
    try:
        case_dir_fd = _open_or_create_case_directory(root_fd, component)
        _assert_directory_binding(
            trusted_root, root_fd, component, case_dir_fd)
        _validate_existing_target(case_dir_fd)

        wb = _build_workbook(file_ir, selection=selection)
        temp_name, temp_fd = _create_new_file(case_dir_fd)
        try:
            with os.fdopen(temp_fd, "wb", closefd=False) as stream:
                wb.save(stream)
                stream.flush()
            os.fsync(temp_fd)
            new_inode = os.fstat(temp_fd)
            if (
                not stat.S_ISREG(new_inode.st_mode)
                or int(new_inode.st_nlink) != 1
            ):
                raise RuntimeError("xlsx staging inode lost its regular-file identity")
        finally:
            wb.close()
            os.close(temp_fd)
            temp_fd = -1

        _assert_directory_binding(
            trusted_root, root_fd, component, case_dir_fd)
        _validate_existing_target(case_dir_fd)
        backup = _make_target_backup(case_dir_fd)
        if backup is not None:
            backup_name, old_inode = backup
            if not _same_target_inode(case_dir_fd, old_inode):
                raise RuntimeError("case.xlsx target changed before atomic replace")

        _assert_directory_binding(
            trusted_root, root_fd, component, case_dir_fd)
        did_replace = False
        try:
            os.replace(
                temp_name,
                _TARGET_NAME,
                src_dir_fd=case_dir_fd,
                dst_dir_fd=case_dir_fd,
            )
            did_replace = True
            temp_name = None
            os.fsync(case_dir_fd)

            _assert_directory_binding(
                trusted_root, root_fd, component, case_dir_fd)
            if new_inode is None or not _same_target_inode(case_dir_fd, new_inode):
                raise RuntimeError("atomic xlsx target identity check failed")
            current = os.stat(
                _TARGET_NAME, dir_fd=case_dir_fd, follow_symlinks=False)
            if int(current.st_nlink) != 1:
                raise RuntimeError("atomic xlsx target unexpectedly has multiple links")

            read_fd = os.open(
                _TARGET_NAME,
                os.O_RDONLY | os.O_NOFOLLOW | getattr(os, "O_CLOEXEC", 0),
                dir_fd=case_dir_fd,
            )
            try:
                if _inode_id(os.fstat(read_fd)) != _inode_id(new_inode):
                    raise RuntimeError("case.xlsx changed before round-trip readback")
                with os.fdopen(read_fd, "rb", closefd=False) as stream:
                    stats = _readback(stream, display_path=out_path)
                    stats["template_release_state"] = selection.release_state
                    stats["runtime_template_sha256"] = TEMPLATE_SHA256
                    if selection.contract_sha256 is not None:
                        stats["excel_contract_sha256"] = selection.contract_sha256
            finally:
                os.close(read_fd)

            _assert_directory_binding(
                trusted_root, root_fd, component, case_dir_fd)
            if not _same_target_inode(case_dir_fd, new_inode):
                raise RuntimeError("case.xlsx changed before atomic emit commit")
        except BaseException:
            if did_replace:
                _rollback_replaced_target(
                    case_dir_fd,
                    backup_name=backup_name,
                    new_inode=new_inode,
                )
                backup_name = None
                os.fsync(case_dir_fd)
            raise

        if backup_name is not None:
            os.unlink(backup_name, dir_fd=case_dir_fd)
            backup_name = None
            os.fsync(case_dir_fd)
        return stats
    finally:
        if temp_fd >= 0:
            os.close(temp_fd)
        if case_dir_fd >= 0:
            if temp_name is not None:
                try:
                    os.unlink(temp_name, dir_fd=case_dir_fd)
                except OSError:
                    pass
            if backup_name is not None:
                try:
                    os.unlink(backup_name, dir_fd=case_dir_fd)
                except OSError:
                    pass
            os.close(case_dir_fd)
        os.close(root_fd)


def _readback(
    source: Path | BinaryIO,
    *,
    display_path: Path | None = None,
) -> dict:
    wb = load_workbook(source, data_only=True)
    sheet, _layout = resolve_execution_sheet(wb, allow_legacy=False)
    grid = [list(row) for row in sheet.iter_rows(values_only=True)]
    anchor = get_config().xlsx.header_anchor
    autoids = []
    case_begin = False
    n_check = 0
    for row in grid:
        a = row[0] if len(row) > 0 else None
        c = row[2] if len(row) > 2 else None
        e = row[4] if len(row) > 4 else None
        if a is not None and str(a).strip() == anchor:
            case_begin = True
            continue
        if not case_begin:
            continue
        if a is not None and str(c) not in ("1", "0", "None"):
            autoids.append(str(a))
        if e == "check_point":
            n_check += 1
    wb.close()
    return {
        "path": str(display_path if display_path is not None else source),
        "rows": len(grid),
        "case_count": len(autoids),
        "autoids": autoids,
        "check_point_count": n_check,
    }


__all__ = [
    "RuntimeTemplateSelection", "select_runtime_template", "emit_xlsx",
]
