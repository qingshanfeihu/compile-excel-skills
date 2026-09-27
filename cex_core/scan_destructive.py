"""自毁命令扫描：整机清配置、恢复出厂、重启/关机这类命令不许出现在用例里。

规则只有一个来源：数据包 projections 里的 `domain_grammar.json` 的
`destructive_commands.patterns`（与 InfoTest structural_gate 读同一份）。这里不写任何命令。

与 InfoTest 的对应：
- 辖区同 InfoTest `is_apv_command_step(methods=None)`：E 列以 APV 开头的每一步都查，
  不按 F 列分辖区（按 F 分等于留一个“换个方法名就能绕过”的口子）；
- 执行页逐行查到最后一行：文件末尾 999999999999999 伪案的行框架照样执行，不在那里停；
- G 列按框架的读法查（knowledge/framework/mirror/lib/test_xlsx.py 的 get_parameter：
  逗号切参、`键=值` 取值、剥一层成对引号）：原文与框架会交给设备的每个参数串都拿规则比，
  `"clear config all"`、`cmd=clear config all` 这类写法绕不过去；
- 工作簿按框架的方式整本载入、取缓存值（read_excel_with_openpyxl 的 data_only=True）；
  不用只读模式——只读模式按表里声明的 dimension 截行，一条伪造的 dimension 就能把后面的行藏起来；
- 执行页里的公式格一律报出：框架执行的是缓存值，公式本身看不出会跑什么；
- 与 InfoTest 不同：规则读不到、为空或编译失败时**拒绝**（抛 DestructiveRulesUnavailable），
  而不是记一笔“本次未检查”后放行——这里是上机前的最后一道闸。
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any, Iterable

# 公式格的命中：不是命令规则，框架按缓存值执行，所以整格拒收
FORMULA_RULE = "formula cell: the framework runs the cached value, not the formula"
_KEYWORD = re.compile(r"^[A-Za-z_]\w*$")


class DestructiveRulesUnavailable(RuntimeError):
    pass


def load_patterns(grammar: dict[str, Any] | str | Path) -> list[re.Pattern[str]]:
    if not isinstance(grammar, dict):
        try:
            grammar = json.loads(Path(grammar).read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            raise DestructiveRulesUnavailable(f"domain grammar unreadable: {exc}") from None
    raw = ((grammar or {}).get("destructive_commands") or {}).get("patterns") or []
    if not isinstance(raw, list) or not raw:
        raise DestructiveRulesUnavailable("domain grammar carries no destructive_commands.patterns")
    try:
        return [re.compile(str(pattern), re.IGNORECASE) for pattern in raw]
    except re.error as exc:
        raise DestructiveRulesUnavailable(f"destructive pattern does not compile: {exc}") from None


# ── G 列在框架里变成什么（test_xlsx.py 的 _split_parameter_parts / _keyword_split /
#    _unquote_parameter，逐条同一判定）──────────────────────────────────────────


def _split_parameter_parts(text: str) -> list[str] | None:
    """按引号外的逗号切参；引号不闭合时框架拒绝整格，这里返回 None（只查原文）。"""
    parts: list[str] = []
    current: list[str] = []
    quote = None
    escaped = False
    for char in text:
        if escaped:
            current.append(char)
            escaped = False
            continue
        if char == "\\":
            current.append(char)
            escaped = True
            continue
        if quote is not None:
            current.append(char)
            if char == quote:
                quote = None
            continue
        if char in ("\"", "'"):
            current.append(char)
            quote = char
            continue
        if char == ",":
            part = "".join(current).strip()
            if part:
                parts.append(part)
            current = []
            continue
        current.append(char)
    if quote is not None:
        return None
    tail = "".join(current).strip()
    if tail:
        parts.append(tail)
    return parts


def _unquote_parameter(value: str) -> str:
    text = value.strip()
    if len(text) >= 2 and text[0] in ("\"", "'") and text[-1] == text[0]:
        return text[1:-1].strip()
    return text


def _keyword_value(part: str) -> str | None:
    """`键=值` 形态的参数返回值（未剥引号）；不是关键字参数返回 None。"""
    quote = None
    escaped = False
    for index, char in enumerate(part):
        if escaped:
            escaped = False
            continue
        if char == "\\":
            escaped = True
            continue
        if quote is not None:
            if char == quote:
                quote = None
            continue
        if char in ("\"", "'"):
            quote = char
            continue
        if char == "=":
            if _KEYWORD.match(part[:index].strip()):
                return part[index + 1:].strip()
            return None
    return None


def command_forms(line: str) -> list[str]:
    """一行 G 的原文，加上框架 get_parameter 会交给设备的每个参数串（位置参数、关键字值，
    都剥掉一层成对引号）。多行格按行拆开后逐行调用：cmds_config/execute 整格下发，设备那边
    仍是一行一条命令，按单行解析只会多查、不会少查。"""
    text = (line or "").strip()
    forms = [text] if text else []
    for part in _split_parameter_parts(text) or []:
        value = _keyword_value(part)
        form = _unquote_parameter(part if value is None else value)
        if form and form not in forms:
            forms.append(form)
    return forms


def scan_lines(lines: Iterable[tuple[str, str]],
               patterns: list[re.Pattern[str]]) -> list[dict[str, str]]:
    """lines 是 (位置, 命令行)；返回每个命中：位置、命令（原文）、命中的规则。

    命中的是框架解析后的参数串（而不是原文）时，另带 ``executed``：设备实际会收到的那一串。"""
    findings = []
    for where, line in lines:
        forms = command_forms(line)
        hit = next(((form, p) for form in forms for p in patterns if p.search(form)), None)
        if hit is None:
            continue
        form, pattern = hit
        finding = {"where": where, "command": forms[0], "rule": pattern.pattern}
        if form != forms[0]:
            finding["executed"] = form
        findings.append(finding)
    return findings


# ── 工作簿 ─────────────────────────────────────────────────────────────────


def _load(xlsx: str | Path, *, data_only: bool):
    from openpyxl import load_workbook

    return load_workbook(Path(xlsx), data_only=data_only)


def _execution_sheet(wb: Any):
    from .ist_emit.excel_contract import resolve_execution_sheet

    ws, _ = resolve_execution_sheet(wb, allow_legacy=False)
    return ws


def _command_lines(ws: Any) -> list[tuple[str, str]]:
    out = []
    for row_no, row in enumerate(ws.iter_rows(values_only=True), start=1):
        device = str(row[4] or "").strip() if len(row) > 4 else ""
        command = str(row[6]) if len(row) > 6 and row[6] is not None else ""
        if not device.startswith("APV") or not command.strip():
            continue
        for line in command.splitlines():
            if line.strip():
                out.append((f"{ws.title}!G{row_no}", line))
    return out


def _formula_cells(cached: Any, formulas: Any) -> list[dict[str, str]]:
    findings = []
    for row in formulas.iter_rows():
        for cell in row:
            if cell.data_type != "f":
                continue
            value = cached[cell.coordinate].value
            findings.append({"where": f"{cached.title}!{cell.coordinate}",
                             "command": "" if value is None else str(value),
                             "rule": FORMULA_RULE,
                             "formula": str(getattr(cell.value, "text", cell.value) or "")})
    return findings


def workbook_command_lines(xlsx: str | Path) -> list[tuple[str, str]]:
    """执行页里每个 APV 步的 G 列逐行拆开（框架读到的缓存值，逐行到表尾）。"""
    wb = _load(xlsx, data_only=True)
    try:
        return _command_lines(_execution_sheet(wb))
    finally:
        wb.close()


def workbook_formula_cells(xlsx: str | Path) -> list[dict[str, str]]:
    """执行页里的每个公式格：位置、框架会执行的缓存值（command）、公式原文（formula）。"""
    cached = _load(xlsx, data_only=True)
    formulas = _load(xlsx, data_only=False)
    try:
        ws = _execution_sheet(cached)
        return _formula_cells(ws, formulas[ws.title])
    finally:
        cached.close()
        formulas.close()


def scan_workbook(xlsx: str | Path, grammar: dict[str, Any] | str | Path) -> list[dict[str, str]]:
    """公式格与自毁命令的全部命中；空列表才说明这本工作簿可以上共享床。"""
    patterns = load_patterns(grammar)
    cached = _load(xlsx, data_only=True)
    formulas = _load(xlsx, data_only=False)
    try:
        ws = _execution_sheet(cached)
        return (_formula_cells(ws, formulas[ws.title])
                + scan_lines(_command_lines(ws), patterns))
    finally:
        cached.close()
        formulas.close()
