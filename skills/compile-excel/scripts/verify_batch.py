#!/usr/bin/env python3
"""verify_batch: case.xlsx 验收报告（pass/fail/totals）。

本地终验（工作流 D4 的本地部分）：
1. 结构检查（ist_emit 真语义：唯一 A-I 表头、defined-name、marker 身份钉死）；
2. 布局检查（Author 行/init 行/步骤行/空行分隔/哨兵）；
3. E/F 合法集检查——**从冻结模板 K-P 列自举**（模板第 2 行自述分组）；
4. 逐 case check_point≥1、found_times 的 G/H/I 硬契约、autoid 唯一；
5. 断言不得命中上一条命令原文（否则是恒真或恒假）；
6. 断言读得到观察回显（悬空断言在设备上让整卷崩掉，见 _dangling_assertions）。

本脚本就是验收。不要再调用 InfoTest 引擎。
用法：python verify_batch.py --xlsx <case.xlsx>
退出码：0 = 全过；1 = 有失败项。
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import load_workbook  # noqa: E402

import _cex_path  # noqa: F401,E402 — 发行根进 sys.path

from cex_core.ist_emit.case_ir import FileIR  # noqa: E402,F401 — 路径注册
from cex_core.ist_emit.excel_contract import (  # noqa: E402
    CONTRACT_MARKER,
    EXECUTION_HEADERS,
    PINNED_CONTRACT_SHA256,
    ExcelContractError,
    resolve_execution_sheet,
)
from cex_core.ist_emit.xlsx_emit import TEMPLATE_PATH  # noqa: E402

REPORT_SCHEMA = "ist.excel.verify-report"


class Report:
    def __init__(self) -> None:
        self.checks: list[dict] = []

    def add(self, name: str, ok: bool, detail: str = "") -> None:
        self.checks.append({"name": name, "ok": bool(ok), "detail": detail})

    def payload(self, path: str) -> dict:
        failures = [c for c in self.checks if not c["ok"]]
        return {
            "schema": REPORT_SCHEMA,
            "path": path,
            "totals": len(self.checks),
            "pass": len(self.checks) - len(failures),
            "fail": len(failures),
            "failures": failures,
            "checks": self.checks,
        }


def _ef_sets_from_template() -> tuple[set[str], dict[str, set[str]]]:
    """从冻结模板 K-P 列自举 E 集与 E→F 集合。

    模板第 2 行自述分组（如 L2="APV_0 / APV_1"），列下方即该组 F 值。
    """
    wb = load_workbook(TEMPLATE_PATH)
    ws, _layout = resolve_execution_sheet(wb, allow_legacy=False)
    e_values: set[str] = set()
    f_map: dict[str, set[str]] = {}
    for row in ws.iter_rows(min_row=3, max_row=12, min_col=11, max_col=16, values_only=True):
        for idx, value in enumerate(row):
            if value is None or not str(value).strip():
                continue
            text = str(value).strip()
            if idx == 0:
                e_values.add(text)
            else:
                group_header = ws.cell(row=2, column=11 + idx).value
                for e_name in str(group_header or "").split("/"):
                    e_name = e_name.strip()
                    if e_name:
                        f_map.setdefault(e_name, set()).add(text)
    wb.close()
    return e_values, f_map


def _command_echo_hits(data: list[list]) -> list[str]:
    """found/not_found/abs_found must not already match the command that feeds it."""
    hits: list[str] = []
    last_cmd = ""
    saved: dict[str, str] = {}
    for row in data:
        e = str(row[4] or "").strip()
        f = str(row[5] or "").strip()
        g = str(row[6] or "")
        h = str(row[7] or "").strip()
        i_col = str(row[8] or "").strip()
        if e == "check_point":
            if h or not g.strip() or f not in {"found", "not_found", "abs_found"}:
                continue
            src = saved.get(i_col) if i_col else last_cmd
            if not src:
                continue
            try:
                matched = (
                    src.find(g) >= 0 if f == "abs_found"
                    else re.compile(g, re.DOTALL).search(src) is not None
                )
            except re.error:
                continue
            if matched:
                hits.append(f"{f} {g!r} matches command {src[:48]!r}")
            continue
        if e.startswith("APV") and f == "cmd_config":
            if h:
                saved[h] = g
            else:
                last_cmd = g
    return hits


_PROMPT_LIKE = re.compile(
    r"^(apv|router[a-z]?|[a-z_]+\(config\)#?|[#>*%-]{1,6}|\s*[#>-]{1,4}\s*)$",
    re.I,
)


def _matches_empty(pattern: str) -> bool:
    """正则可匹配空串 ⇒ 可匹配任何回显（恒真）。"""
    try:
        return re.compile(pattern, re.DOTALL).search("") is not None
    except re.error:
        return False


def _tautology_family(data: list[list]) -> list[str]:
    """恒真/恒假断言族（静态可判子集，镜像引擎 emit 必崩规则的机械部分）。

    - 期望命中提示符形态（APV/#/>/(config)#）⇒ 每行回显都命中 ⇒ 恒真
      （not_found 同形 ⇒ 恒假）；
    - found 正则可匹配空串 ⇒ 恒真；
    - not_found 的期望词出现在喂给它的命令行里 ⇒ 回显必含命令行 ⇒ 恒假。
    """
    bad: list[str] = []
    last_cmd = ""
    saved: dict[str, str] = {}
    for row in data:
        e = str(row[4] or "").strip()
        f = str(row[5] or "").strip()
        g = str(row[6] or "")
        h = str(row[7] or "").strip()
        i_col = str(row[8] or "").strip()
        if e == "check_point":
            if h or not g.strip() or f not in {"found", "not_found", "abs_found"}:
                continue
            expected = g.strip()
            if _PROMPT_LIKE.match(expected):
                bad.append(f"{f} {expected!r} 是提示符形态（每行回显都命中）"
                           f" → 恒真{'假' if f == 'not_found' else ''}")
                continue
            if f == "found" and _matches_empty(g):
                bad.append(f"found /{g}/ 可匹配空串 → 恒真")
                continue
            if f == "not_found":
                src = saved.get(i_col) if i_col else last_cmd
                if src:
                    toks = {t.strip("\"'") for t in src.split()}
                    if expected.strip("\"'") in toks:
                        bad.append(f"not_found {expected!r} 出现在命令 "
                                   f"{src[:40]!r} 里（回显含命令行）→ 恒假")
            continue
        if e.startswith("APV") and f == "cmd_config":
            if h:
                saved[h] = g
            else:
                last_cmd = g
    return bad


def _provenance_problems(xlsx: Path, data: list[list]) -> list[str]:
    """provenance.json 边车：逐 case check_point 必须有非空来源（kind+ref）。"""
    prov_path = xlsx.parent / "provenance.json"
    if not prov_path.is_file():
        return [f"缺 {prov_path.name}（先跑 compile_excel，不要手写 xlsx）"]
    try:
        prov = json.loads(prov_path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError) as exc:
        return [f"provenance.json 不可读: {exc}"]
    if prov.get("schema") != "ist.excel.provenance":
        return [f"provenance schema 不符: {prov.get('schema')!r}"]
    cases = prov.get("cases") if isinstance(prov.get("cases"), dict) else {}
    problems: list[str] = []
    cp_counts: dict[str, int] = {}
    for row in data:
        a = str(row[0] or "").strip()
        e = str(row[4] or "").strip()
        if a.isdigit() and len(a) >= 12 and a != "999999999999999":
            cp_counts.setdefault(a, 0)
        if e == "check_point" and cp_counts:
            last = next(reversed(cp_counts))
            cp_counts[last] += 1
    for autoid, n in cp_counts.items():
        entries = cases.get(autoid)
        if not isinstance(entries, list) or not entries:
            problems.append(f"{autoid}: provenance 缺该 case 的来源记录")
            continue
        if len(entries) != n:
            problems.append(f"{autoid}: provenance {len(entries)} 条 ≠ xlsx "
                            f"check_point {n} 条")
        for ent in entries:
            src = (ent or {}).get("source") or {}
            if not str(src.get("kind") or "").strip() \
                    or not str(src.get("ref") or "").strip():
                problems.append(f"{autoid}: 存在 kind/ref 为空的来源记录")
    return problems


# 产生可供断言读取的回显的方法（引擎 structural_gate._is_observation_step 同一口径）：
# test_env 探针、CLI 的 cmd_config、设备 Linux root shell 的 cmd；execute 引擎只认有返回值的动作，
# 这里一律放行，由引擎的结构闸细判。
_OBSERVATION_METHODS = frozenset({"cmd", "cmd_config", "execute"})


def _dangling_assertions(data: list[list]) -> list[str]:
    """悬空断言：check_point 读框架 result，而 result 只由本案最近一条不带 H 的步骤给出。

    镜像引擎 structural_gate._check_dangling_assertions 与框架 test_xlsx 的 preflight：
    - 本案在它之前没有一条不带 H 的步骤（案首就是断言，或前面每步都带 H——H 只把输出存进寄存器，
      不更新 result）：框架 preflight 报「check_point 前没有可用的本案 observation result」，整卷碰设备前就拒跑；
    - 最近那条不带 H 的步骤不是观察：cmds_config、time::sleep 返回 None（check_point 不允许比较 None，
      运行到这里抛错，后面的案都不跑），cmd_enable 等引擎同样判 dangling_assertion。
    带 I 的断言读寄存器、不读 result，不在此列；found_times 永远读 result。
    """
    bad: list[str] = []
    autoid, in_case, has_result, observe, last = "", False, False, False, ""
    for row in data:
        a = str(row[0]).strip() if row[0] is not None else ""
        c = str(row[2]).strip() if row[2] is not None else ""
        if a and c not in ("0", "1", "None") and a != EXECUTION_HEADERS[0]:
            autoid, in_case, has_result, observe, last = a, True, False, False, ""
        if not in_case:
            continue
        e = str(row[4] or "").strip()
        f = str(row[5] or "").strip()
        h = str(row[7] or "").strip()
        i_col = str(row[8] or "").strip()
        if not e:
            continue
        if e == "check_point":
            if i_col and f != "found_times":
                continue
            what = f"{f} {str(row[6] or '')[:40]!r}"
            if not has_result:
                bad.append(f"{autoid}: {what} 之前本案没有不带 H 的观察步（框架 preflight 整卷拒跑）")
            elif not observe:
                bad.append(f"{autoid}: {what} 读的是 {last} 的返回，不是观察回显")
            continue
        if h:
            continue
        has_result = True
        observe = e == "test_env" or f in _OBSERVATION_METHODS
        last = f"{e}::{f}"
    return bad


_INIT_NEUTRAL = re.compile(
    r"^(clear|no|show)\s|^(config(ure)?\s+t(erminal)?|conf\s+t|enable|end|exit)$", re.IGNORECASE)


def _init_isolation_problems(data: list[list]) -> list[str]:
    """init 行（C=1）在每个案之前都会重放：只许清场（clear / no）、只读（show）和模式切换。

    在 init 里建对象，等于给每个案都塞了一条隐藏前置：某案重定义同名对象或与之冲突时，
    设备会拒绝那一步，而断言可能照样命中别的东西。案要用的对象写进该案自己的步骤。
    """
    problems: list[str] = []
    for row in data:
        if str(row[2]).strip() != "1":
            continue
        e = str(row[4] or "").strip()
        if not e.startswith("APV_"):
            continue
        for line in str(row[6] or "").splitlines():
            command = line.strip()
            if command and not _INIT_NEUTRAL.search(command):
                problems.append(command)
    return problems


def verify(path: Path) -> dict:
    report = Report()
    wb = load_workbook(path, data_only=True)

    # 1) 真语义结构解析（含 marker 身份钉死）
    try:
        sheet, layout = resolve_execution_sheet(wb, allow_legacy=False)
        report.add("execution-sheet resolve (pinned identity)", True,
                   f"header_row={layout.header_row} data_start={layout.data_start}")
    except ExcelContractError as exc:
        report.add("execution-sheet resolve (pinned identity)", False, str(exc))
        return report.payload(str(path))
    marker_cell = sheet.cell(row=1, column=1).value
    marker_sha = str(sheet.cell(row=1, column=3).value or "")
    report.add("contract marker present", marker_cell == CONTRACT_MARKER,
               f"A1={marker_cell!r}")
    report.add("marker contract sha pinned",
               marker_sha == PINNED_CONTRACT_SHA256, marker_sha[:12] + "…")

    grid = [list(r) for r in sheet.iter_rows(values_only=True)]
    data = [row + [None] * max(0, 9 - len(row)) for row in grid[layout.data_start - 1:]]
    data = [row for row in data if any(v not in (None, "") for v in row[:9])]

    # 2) 布局：首行 Author（C=0），init 行 C=1，case 行 C>=2
    first = data[0] if data else []
    report.add("author row first (C=0)",
               bool(data) and str(first[2]) == "0" and "Author" in str(first[3] or ""))
    stmt_bad = [i for i, row in enumerate(data)
                if row[2] is not None and str(row[2]).strip() not in {"0", "1"}
                and not str(row[2]).isdigit()]
    report.add("stmt_type column discipline (0/1/int>=2)", not stmt_bad,
               f"bad_rows={[data[i][0] or data[i][4] for i in stmt_bad][:5]}")

    # 3) E/F 合法集（模板自举）
    e_values, f_map = _ef_sets_from_template()
    ef_bad: list[str] = []
    for row in data:
        e = str(row[4] or "").strip()
        f = str(row[5] or "").strip()
        if not e:
            continue
        if e not in e_values:
            ef_bad.append(f"E={e}")
            continue
        allowed = f_map.get(e)
        if allowed and f and f not in allowed:
            ef_bad.append(f"E={e} F={f}")
    report.add("E/F membership (template-derived)", not ef_bad,
               f"violations={ef_bad[:5]} known_E={sorted(e_values)}")

    # 4) 逐 case：check_point≥1；found_times 契约；autoid 唯一
    autoids: list[str] = []
    current: dict | None = None
    case_results: list[tuple[str, bool, str]] = []
    for row in data:
        a = str(row[0]).strip() if row[0] is not None else ""
        c = str(row[2]).strip() if row[2] is not None else ""
        if a and c not in ("1", "0", "None") and a != EXECUTION_HEADERS[0]:
            if current is not None:
                case_results.append((current["autoid"], current["cp"] > 0, ""))
            current = {"autoid": a, "cp": 0}
            if a in autoids:
                case_results.append((a, False, "autoid 重复"))
            autoids.append(a)
        e = str(row[4] or "").strip()
        f = str(row[5] or "").strip()
        if current is not None and e == "check_point":
            current["cp"] += 1
            if f == "found_times":
                i_text = str(row[8] or "").strip()
                h_text = str(row[7] or "").strip()
                try:
                    count = int(i_text)
                    ok_ft = count > 0 and str(count) == i_text and not h_text
                except ValueError:
                    ok_ft = False
                if not ok_ft:
                    case_results.append(
                        (current["autoid"], False,
                         f"found_times 契约违规：I={i_text!r} H={h_text!r}"))
    if current is not None:
        case_results.append((current["autoid"], current["cp"] > 0, ""))
    no_cp = [aid for aid, ok, _note in case_results if not ok and aid != "999999999999999"]
    report.add("every case has check_point", not no_cp, f"missing={no_cp}")
    report.add("autoid unique", len(autoids) == len(set(autoids)))
    # InfoTest structural_gate 按全数字 ≥12 位识别用例边界（哨兵除外）
    real_autoids = [a for a in autoids if a != "999999999999999"]
    bad_ids = [a for a in real_autoids if not (a.isdigit() and len(a) >= 12)]
    report.add("autoid >= 12 digits (framework boundary)", not bad_ids,
               f"bad={bad_ids}（生产惯例 18 位）")
    report.add("case count", True, f"cases={len(autoids)} autoids={autoids}")
    dangling = _dangling_assertions(data)
    report.add("every assertion reads an observation echo (no dangling check_point)", not dangling,
               f"dangling={dangling[:4]}")
    echo_bad = _command_echo_hits(data)
    report.add("assertion does not match the command text", not echo_bad,
               f"hits={echo_bad[:4]}")
    taut_bad = _tautology_family(data)
    report.add("tautology family (prompt-like / empty-match / not_found-in-command)",
               not taut_bad, f"hits={taut_bad[:4]}")
    prov_bad = _provenance_problems(path, data)
    report.add("provenance sidecar (expected-value sources)", not prov_bad,
               f"problems={prov_bad[:4]}")
    init_bad = _init_isolation_problems(data)
    report.add("init rows only reset state (they replay before every case)", not init_bad,
               f"state-creating init commands={init_bad[:4]}（放进用到它的那个案的步骤里）")

    wb.close()
    return report.payload(str(path))


def main() -> int:
    parser = argparse.ArgumentParser(description="case.xlsx 验收报告")
    parser.add_argument("--xlsx", required=True, help="待验 case.xlsx")
    args = parser.parse_args()

    result = verify(Path(args.xlsx))
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["fail"] == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
