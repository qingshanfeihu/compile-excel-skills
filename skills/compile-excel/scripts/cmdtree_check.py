#!/usr/bin/env python3
"""cmdtree_check: 编译期命令 grounding —— cases.json 里的每条 APV CLI 命令必须在命令树中存在。

把「设备不支持这个命令」从上机期（烧一轮床才发现 % invalid）提前到编译期。
对齐引擎口径：命令在不在这个 build 上，由命令树投影单方面说了算——查不到就拒绝，
不管规格书或用例怎么写。

判定来源：--projection 显式路径，否则工作区已同步数据包（cex_sync）里的命令树投影
（vendor_stdlib_*.json）；只有显式给了 --tree 才读原始 cmdtree XML（本地调试用）。
投影判定用 cex_core.vendor_cmd.resolve_vendor_command——与引擎同一个判定函数（逐字抽取），
既判命令头在不在，也判参数是否合契约。服务端只下发投影，不下发原始 XML
（原始 XML 带参数默认值，含凭据默认值）。

查哪些行（与引擎 emit_xlsx_tool._ordered_apv_command_refs / _apv_command_lines_for_step 同一解释）：
- 文件级 init_commands（每条一行）与每个案的步骤；步骤键名 e/f/g 与 E/F/G 都认
  （手写 cases.json 用小写，cex_author_emit 出件的用大写）；
- E=APV_* 且 F 是 CLI 方法（cmd_config / cmds_config / cmd_enable）的步骤；cmds_config 的多行 G
  逐行拆开，每行是一条命令；
- G 里执行器的关键字参数（`,timeout=60`、`,prompt=abort:`）不是命令的一部分，先剥掉；
  紧跟在带 prompt= 的 cmd_config 之后的 YES / NO 是交互应答，不是命令；
- F=cmd 在设备的 Linux root shell 里执行（框架 APV.cmd → APV_Root），G 是 shell 命令（ls、cat、
  /ca/bin/…），不在命令树里，不查；它若恰好是一条 CLI 命令，报警告（CLI 命令要用 cmd_config）。

命令按原文判，不剥前导 no/show/clear：它们在命令树里各有自己的路径和参数契约
（`no slb policy default <vs>` 只收一个参数，剥成正向命令去判就会误报参数不合）——与引擎、
cex_cmd_check 同。--tree 原始 XML 模式：在树中做贪心路径匹配；全路径命中 = ok；部分命中 =
head 只到 menu；零命中 = head 不在树中（附同前缀兄弟节点作建议）。
树的 build（文件名 _585 等）与工作区 device_build 不一致时警告（不阻断）。

用法：
  python3 scripts/cmdtree_check.py --cases cases.json [--projection vendor_stdlib_x.json]
  python3 scripts/cmdtree_check.py --xlsx compile_outputs/<batch>/case.xlsx
退出码：0 = 全部命中；1 = 存在未知命令或参数不合契约（应修 cases.json 再编译）；2 = 树缺失/不可解析。
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import _cex_path  # noqa: F401,E402 — 发行根进 sys.path

REPORT_SCHEMA = "ist.excel.cmdtree-check"

_BUILD_RE = re.compile(r"cmdtree[_\-]?(\w+)\.xml$", re.I)
_SENTINEL_AUTOID = "999999999999999"
_INIT = "init"

# 设备 CLI 方法（引擎存在性判据的辖区是 cmd_config / cmds_config；cmd_enable 同是 CLI，一并查）
_CLI_METHODS = frozenset({"cmd_config", "cmds_config", "cmd_enable"})
_ROOT_SHELL_METHOD = "cmd"
_PROMPT_ANSWERS = frozenset({"YES", "NO"})

# 剥掉后不查树的纯框架词（不出现在设备命令树里的执行器指令）
_FRAMEWORK_WORDS = {"page"}


def _field(step: dict, key: str):
    """小写 e/f/g/h/i 与大写 E/F/G/H/I 等价。"""
    if key in step:
        return step[key]
    return step.get(key.upper())


def _workspace(start: Path):
    from cex_client import workspace as wsmod

    return wsmod.find(start)


def _find_tree(explicit: str) -> tuple[Path | None, str]:
    """返回 (树路径, build 标签)。原始 XML 只认显式路径，不在任何目录里自动找。"""
    path = Path(explicit).expanduser()
    if not path.is_file():
        return None, ""
    m = _BUILD_RE.search(path.name)
    return path, (m.group(1) if m else "")


def _find_projection(explicit: str, workspace: Path) -> Path | None:
    if explicit:
        path = Path(explicit).expanduser()
        return path if path.is_file() else None
    ws = _workspace(workspace)
    if ws is None:
        return None
    from cex_client import bundle

    try:
        return bundle.entry_path(ws, "cmdtree", "vendor_stdlib_")
    except Exception:  # noqa: BLE001 — 工作区坏了按“没有投影”处理，由调用方报 exit 2
        return None


def _normalized(step: dict) -> dict:
    return {"E": str(_field(step, "e") or "").strip(), "F": str(_field(step, "f") or "").strip(),
            "G": "" if _field(step, "g") is None else str(_field(step, "g"))}


def _command_lines(groups: list[tuple[str, list[dict]]]) -> tuple[list[dict], list[dict]]:
    """[(autoid, 步骤)] → (要查树的 CLI 命令行, F=cmd 的 root shell 行)。每行 {autoid, e, f, command}。"""
    from cex_core.engine.case_compiler.excel_contract import (
        parse_g_arguments,
        strip_apv_command_kwargs,
    )

    cli: list[dict] = []
    shell: list[dict] = []
    for autoid, steps in groups:
        for index, step in enumerate(steps):
            e, f, g = step["E"], step["F"], step["G"]
            if not e.startswith("APV") or not g.strip():
                continue
            if f not in _CLI_METHODS and f != _ROOT_SHELL_METHOD:
                continue
            try:
                command = strip_apv_command_kwargs(g, f)
            except Exception:  # noqa: BLE001 — 与引擎同：G 解析不了就按原文查
                command = g
            lines = [line.strip() for line in command.splitlines() if line.strip()]
            if f == _ROOT_SHELL_METHOD:
                shell.extend({"autoid": autoid, "e": e, "f": f, "command": line} for line in lines)
                continue
            if (f == "cmd_config" and lines and lines[0].upper() in _PROMPT_ANSWERS and index > 0
                    and steps[index - 1]["E"] == e and steps[index - 1]["F"] == "cmd_config"):
                try:
                    _args, kwargs = parse_g_arguments(steps[index - 1]["G"], "cmd_config")
                except Exception:  # noqa: BLE001
                    kwargs = {}
                if str(kwargs.get("prompt") or "").strip():
                    continue  # 上一步 prompt= 等来的交互应答，不是产品命令
            cli.extend({"autoid": autoid, "e": e, "f": f, "command": line} for line in lines)
    return cli, shell


def _groups_from_cases(cases_path: Path) -> list[tuple[str, list[dict]]]:
    data = json.loads(cases_path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise TypeError("cases JSON 顶层必须是对象")
    groups: list[tuple[str, list[dict]]] = []
    init = [str(c) for c in data.get("init_commands") or [] if str(c).strip()]
    if init:
        # compile_excel 把 init_commands 合成一条 APV_0::cmds_config 共享前置块
        groups.append((_INIT, [{"E": "APV_0", "F": "cmds_config", "G": "\n".join(init)}]))
    for case in data.get("cases") or []:
        if not isinstance(case, dict):
            continue
        steps = [_normalized(s) for s in case.get("steps") or [] if isinstance(s, dict)]
        groups.append((str(case.get("autoid") or ""), steps))
    return groups


def _groups_from_xlsx(xlsx: Path) -> list[tuple[str, list[dict]]]:
    from openpyxl import load_workbook

    from cex_core.ist_emit.excel_contract import resolve_execution_sheet

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws, layout = resolve_execution_sheet(wb, allow_legacy=False)
    rows = list(ws.iter_rows(min_row=layout.data_start, values_only=True))
    wb.close()
    groups: list[tuple[str, list[dict]]] = []
    init: list[dict] = []
    current: list[dict] | None = None
    for row in rows:
        row = list(row) + [None] * max(0, 9 - len(row))
        a = str(row[0]).strip() if row[0] is not None else ""
        c = str(row[2]).strip() if row[2] is not None else ""
        step = {"E": str(row[4] or "").strip(), "F": str(row[5] or "").strip(),
                "G": "" if row[6] is None else str(row[6])}
        if c == "1":
            init.append(step)
            continue
        if a.isdigit() and len(a) >= 12:
            current = [] if a != _SENTINEL_AUTOID else None
            if current is not None:
                groups.append((a, current))
        if current is not None and step["E"]:
            current.append(step)
    return ([(_INIT, init)] if init else []) + groups


def _check_with_projection(lines: list[dict], shell: list[dict], projection_path: Path,
                           report: dict) -> None:
    from cex_core.vendor_cmd import load_projection, resolve_vendor_command

    projection = load_projection(projection_path)
    for item in lines:
        cmd = item["command"]
        report["checked"] += 1
        if cmd.lower() in _FRAMEWORK_WORDS:
            report["ok"] += 1
            continue
        verdict = resolve_vendor_command(cmd, projection)
        if verdict.get("hit"):
            report["ok"] += 1
            continue
        if not verdict.get("decided"):
            report["warnings"].append(f"{cmd!r}: 投影无法判定（不是以命令词开头的行）")
            report["ok"] += 1
            continue
        report["unknown"].append({
            "autoid": item["autoid"],
            "command": cmd,
            "matched_prefix": verdict.get("head", ""),
            "reason": ("参数不合投影记录的契约" if verdict.get("reason_code")
                       == "parameter_contract_violation" else "命令头不在投影里"),
            "reason_code": verdict.get("reason_code", ""),
            "parameter_error": verdict.get("parameter_error"),
        })
    for item in shell:
        verdict = resolve_vendor_command(item["command"], projection)
        if verdict.get("head"):
            report["warnings"].append(
                f"{item['autoid']}: {item['command']!r} 是一条 CLI 命令，但 F=cmd 在设备的 Linux root shell "
                "里执行它，拿不到 CLI 回显；CLI 命令用 cmd_config")


def _tree_paths(xml_path: Path) -> tuple[dict, dict]:
    """返回 (路径→is_item 叶命令, 每个前缀的直接子节点)。

    匹配语义：命令 head = 树中一条路径；head 之后的 token 全是运行期参数，不校验。
    head 必须落在 item（叶命令）上——只到 menu 不算完整命令（如 `slb virtual`）。
    """
    root = ET.parse(str(xml_path)).getroot()
    is_item: dict = {}
    children: dict = {}

    def walk(node, prefix):
        for child in node:
            name = (child.get("name") or "").strip().lower()
            if not name:
                continue
            p = prefix + (name,)
            leaf = child.tag == "item"
            if leaf:
                is_item[p] = True
            else:
                is_item.setdefault(p, False)
                walk(child, p)
            children.setdefault(prefix, set()).add(name)

    for scope in root:
        walk(scope, ())
    return is_item, children


def _check_with_tree(lines: list[dict], is_item: dict, children: dict, report: dict) -> None:
    for item in lines:
        cmd = item["command"]
        report["checked"] += 1
        if cmd.lower() in _FRAMEWORK_WORDS:
            report["ok"] += 1
            continue
        toks = cmd.lower().split()
        # 贪心最深路径；之后的 token 全是参数，不校验
        depth = 0
        for d in range(len(toks), 0, -1):
            if tuple(toks[:d]) in is_item:
                depth = d
                break
        if depth and is_item[tuple(toks[:depth])]:
            report["ok"] += 1
            continue
        prefix = tuple(toks[:depth]) if depth else ()
        sibs = sorted(children.get(prefix, set()))
        last = toks[depth] if depth < len(toks) else ""

        # 近似优先：前缀命中 / 单复数之差 / 包含关系排前，其余字典序截断
        def _rank(s: str, last: str = last) -> tuple:
            near = (
                s.startswith(last[:3])
                or s.rstrip("s") == last.rstrip("s")
                or last.rstrip("s") in s
            )
            return (0 if near else 1, s)
        sibs = [s for _, s in sorted((_rank(s), s) for s in sibs)][:8]
        report["unknown"].append({
            "autoid": item["autoid"],
            "command": cmd,
            "match_depth": depth,
            "tokens": len(toks),
            "matched_prefix": " ".join(prefix),
            "reason": "head 只到 menu 不是完整命令" if depth else "head 不在树中",
            "suggestions": sibs,
        })


def main() -> int:
    ap = argparse.ArgumentParser(description="编译期命令存在性校验（cmdtree grounding）")
    ap.add_argument("--cases")
    ap.add_argument("--xlsx")
    ap.add_argument("--projection", default="", help="命令树投影 JSON（缺省从工作区数据包找）")
    ap.add_argument("--tree", default="", help="原始 cmdtree XML（只在显式给出时使用）")
    args = ap.parse_args()
    if not args.cases and not args.xlsx:
        print(json.dumps({"ok": False, "error": "--cases 或 --xlsx 必填其一"},
                         ensure_ascii=False))
        return 2

    src = Path(args.cases or args.xlsx).expanduser().resolve()
    if not src.exists():
        print(json.dumps({"ok": False, "error": f"输入不存在: {src}"},
                         ensure_ascii=False))
        return 2
    try:
        groups = (_groups_from_cases if args.cases else _groups_from_xlsx)(src)
    except (ValueError, TypeError, OSError) as exc:  # JSONDecodeError 是 ValueError
        print(json.dumps({"ok": False, "error": f"输入不可读: {exc}"}, ensure_ascii=False))
        return 2
    lines, shell = _command_lines(groups)

    if not args.tree:
        projection = _find_projection(args.projection, src.parent)
        if projection is None:
            print(json.dumps({
                "ok": False,
                "error": "找不到命令树投影：先 cex_sync 同步数据包，或用 --projection 指定",
            }, ensure_ascii=False))
            return 2
        report = {"schema": REPORT_SCHEMA, "projection": str(projection), "source": str(src),
                  "checked": 0, "ok": 0, "unknown": [], "warnings": []}
        try:
            _check_with_projection(lines, shell, projection, report)
        except ValueError as exc:
            print(json.dumps({"ok": False, "error": f"投影不可用: {exc}"}, ensure_ascii=False))
            return 2
        report["ok_flag"] = not report["unknown"]
        print(json.dumps(report, ensure_ascii=False, indent=1))
        return 0 if report["ok_flag"] else 1

    tree_path, build = _find_tree(args.tree)
    if tree_path is None:
        print(json.dumps({"ok": False, "error": f"--tree 指定的文件不存在: {args.tree}"},
                         ensure_ascii=False))
        return 2

    try:
        is_item, children = _tree_paths(tree_path)
    except ET.ParseError as exc:
        print(json.dumps({"ok": False, "error": f"命令树解析失败: {exc}"},
                         ensure_ascii=False))
        return 2

    report = {"schema": REPORT_SCHEMA, "tree": str(tree_path), "tree_build": build,
              "source": str(src), "checked": 0, "ok": 0, "unknown": [], "warnings": []}

    ws = _workspace(src.parent)
    want_build = ws.device_build if ws is not None else ""
    if build and want_build and build.lower() not in want_build.lower():
        report["warnings"].append(
            f"树 build={build} 与工作区 device_build={want_build} 不一致——"
            "树判定可能不代表床固件，建议换对应 build 的树"
        )
    _check_with_tree(lines, is_item, children, report)
    report["ok_flag"] = not report["unknown"]
    print(json.dumps(report, ensure_ascii=False, indent=1))
    return 0 if report["ok_flag"] else 1


if __name__ == "__main__":
    sys.exit(main())
