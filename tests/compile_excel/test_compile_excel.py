#!/usr/bin/env python3
"""compile_excel 出件结构自检（stdlib unittest + openpyxl，无需 pytest）。

跑法：python3 tests/test_compile_excel.py
覆盖：执行页真语义解析（表头第 29 行）、marker/defined-name 校验、
行布局（Author/C=1 init/C=2+i 步骤/空行分隔/哨兵）、原子写盘产物、
错误路径（无 check_point、坏 batch、found_times 契约）。
"""

from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
SKILL_ROOT = REPO_ROOT / "skills" / "compile-excel"
SCRIPTS = SKILL_ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

from openpyxl import load_workbook  # noqa: E402

import _cex_path  # noqa: F401,E402
from cex_core.ist_emit.excel_contract import (  # noqa: E402
    EXECUTION_HEADERS,
    ExcelContractError,
    resolve_execution_sheet,
)
from compile_excel import CompileError, build_file_ir  # noqa: E402

SAMPLE = {
    "batch": "unit_batch",
    "init_commands": ["configure terminal"],
    "cases": [
        {
            "autoid": "202609239000010001",
            "steps": [
                {"e": "APV_0", "f": "cmd_config", "g": "show version", "h": "", "i": ""},
                {"e": "check_point", "f": "found", "g": "version", "h": "", "i": ""},
            ],
        },
        {
            "autoid": "202609239000020001",
            "steps": [
                {"e": "APV_0", "f": "cmd_config", "g": "show slb all", "h": "out1", "i": ""},
                {"e": "check_point", "f": "found", "g": "out1", "h": "", "i": ""},
            ],
        },
    ],
}


class BuildFileIRTests(unittest.TestCase):

    def test_sentinel_appended_and_init_block(self):
        fir = build_file_ir(SAMPLE)
        self.assertEqual(fir.feature, "unit_batch")
        self.assertEqual(len(fir.cases), 3)
        self.assertEqual(fir.cases[-1].autoid, "999999999999999")
        self.assertEqual(fir.cases[-1].priority, "P9")
        self.assertEqual(fir.init_rows[0].method, "cmds_config")
        self.assertEqual(fir.init_rows[0].test_object, "APV_0")
        self.assertEqual(fir.init_rows[0].data, "configure terminal")

    def test_no_sentinel_opt_out(self):
        fir = build_file_ir(SAMPLE, sentinel=False)
        self.assertEqual(len(fir.cases), 2)

    def test_stmt_type_and_capture_normalization(self):
        fir = build_file_ir(SAMPLE)
        case2 = fir.cases[1]
        # check_point 引用前序捕获变量 out1：G→H、found→abs_found
        cp = case2.steps[1].rows[0]
        self.assertEqual(cp.method, "abs_found")
        self.assertEqual(cp.save_as, "out1")
        self.assertEqual(cp.data, "")
        # stmt_type 从 2 递增
        self.assertEqual([s.stmt_type for s in fir.cases[0].steps], [2, 3])

    def test_missing_check_point_rejected(self):
        doc = {"batch": "x", "cases": [{"autoid": "202609239000980001", "steps": [
            {"e": "APV_0", "f": "cmd_config", "g": "show version"}]}]}
        with self.assertRaises(CompileError):
            build_file_ir(doc)

    def test_bad_batch_rejected(self):
        doc = {"batch": "a/b", "cases": SAMPLE["cases"]}
        with self.assertRaises(CompileError):
            build_file_ir(doc)

    def test_found_times_contract(self):
        doc = {"batch": "x", "cases": [{"autoid": "202609239000990001", "steps": [
            {"e": "check_point", "f": "found_times", "g": "a", "h": "", "i": "x"}]}]}
        with self.assertRaises(CompileError):
            build_file_ir(doc)

    def test_short_autoid_rejected(self):
        doc = {"batch": "x", "cases": [{"autoid": "668301", "steps": [
            {"e": "APV_0", "f": "cmd_config", "g": "show version"},
            {"e": "check_point", "f": "found", "g": "version"}]}]}
        with self.assertRaises(CompileError):
            build_file_ir(doc)

    def test_duplicate_autoid_rejected(self):
        doc = {"batch": "x", "cases": [SAMPLE["cases"][0], SAMPLE["cases"][0]]}
        with self.assertRaises(CompileError):
            build_file_ir(doc)

    def test_missing_batch_says_where_it_belongs(self):
        doc = {"cases": SAMPLE["cases"]}
        with self.assertRaises(CompileError) as ctx:
            build_file_ir(doc)
        self.assertIn("顶层", str(ctx.exception))
        self.assertIn("batch", str(ctx.exception))

    def test_init_commands_grouped_by_device_give_one_row_per_device(self):
        doc = {**SAMPLE, "init_commands": {"APV_1": ["clear sdns all"],
                                           "APV_0": ["clear sdns all", "clear slb all"]}}
        rows = build_file_ir(doc).init_rows
        self.assertEqual([(r.test_object, r.method, r.data) for r in rows],
                         [("APV_0", "cmds_config", "clear sdns all\nclear slb all"),
                          ("APV_1", "cmds_config", "clear sdns all")])

    def test_init_commands_for_an_unknown_device_are_refused(self):
        for bad in ({"APV_3": ["clear sdns all"]}, {"APV_0": "clear sdns all"}, "clear sdns all"):
            with self.assertRaises(CompileError):
                build_file_ir({**SAMPLE, "init_commands": bad})


class OutRootTests(unittest.TestCase):

    def test_out_that_already_names_the_batch_is_not_nested_again(self):
        tmp = Path(tempfile.mkdtemp(prefix="ist_out_"))
        cases = tmp / "cases.json"
        cases.write_text(json.dumps(SAMPLE, ensure_ascii=False), encoding="utf-8")
        for out in (tmp / "compile_outputs", tmp / "compile_outputs" / "unit_batch"):
            proc = subprocess.run(
                [sys.executable, str(SCRIPTS / "compile_excel.py"), "--cases", str(cases),
                 "--out", str(out)], capture_output=True, text=True)
            self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
            stats = json.loads(proc.stdout)
            self.assertEqual(stats["batch_dir"], str((tmp / "compile_outputs" / "unit_batch").resolve()))
        self.assertTrue((tmp / "compile_outputs" / "unit_batch" / "case.xlsx").is_file())
        self.assertFalse((tmp / "compile_outputs" / "unit_batch" / "unit_batch").exists())

    def test_cli_error_names_the_cases_file(self):
        tmp = Path(tempfile.mkdtemp(prefix="ist_out_"))
        cases = tmp / "cases.json"
        cases.write_text(json.dumps({"cases": SAMPLE["cases"]}), encoding="utf-8")
        proc = subprocess.run([sys.executable, str(SCRIPTS / "compile_excel.py"), "--cases",
                               str(cases), "--out", str(tmp)], capture_output=True, text=True)
        out = json.loads(proc.stdout)
        self.assertEqual(proc.returncode, 1)
        self.assertEqual(out["cases"], str(cases))
        self.assertIn("顶层缺少 batch", out["error"])


class EmitStructureTests(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="ist_emit_test_")
        cls.cases_path = Path(cls.tmp) / "cases.json"
        cls.cases_path.write_text(json.dumps(SAMPLE, ensure_ascii=False), encoding="utf-8")
        cls.out = Path(cls.tmp) / "out"
        proc = subprocess.run(
            [sys.executable, str(SCRIPTS / "compile_excel.py"),
             "--cases", str(cls.cases_path), "--out", str(cls.out)],
            capture_output=True, text=True,
        )
        assert proc.returncode == 0, proc.stdout + proc.stderr
        cls.stats = json.loads(proc.stdout)
        cls.target = cls.out / "unit_batch" / "case.xlsx"

    def test_stats_and_path(self):
        self.assertTrue(self.stats["ok"])
        self.assertEqual(self.stats["batch"], "unit_batch")
        self.assertEqual(self.stats["case_count"], 3)  # 2 真 case + 哨兵
        self.assertEqual(self.stats["autoids"],
                         ["202609239000010001", "202609239000020001", "999999999999999"])
        self.assertTrue(self.target.is_file())

    def test_execution_sheet_real_semantics(self):
        wb = load_workbook(self.target)
        sheet, layout = resolve_execution_sheet(wb)
        self.assertEqual(layout.header_row, 29)
        self.assertEqual(layout.data_start, 30)
        headers = tuple(sheet.cell(row=29, column=c).value for c in range(1, 10))
        self.assertEqual(headers, EXECUTION_HEADERS)
        wb.close()

    def test_row_layout(self):
        wb = load_workbook(self.target)
        ws, layout = resolve_execution_sheet(wb)
        rows = []
        for r in range(30, 42):
            rows.append([ws.cell(row=r, column=c).value for c in range(1, 10)])
        # r30: Author 行（C=0）
        self.assertEqual(rows[0][2], 0)
        self.assertIn("Author", str(rows[0][3]))
        # r31: init 行（C=1, E=APV_0, F=cmds_config）
        self.assertEqual(rows[1][2], 1)
        self.assertEqual(rows[1][4], "APV_0")
        self.assertEqual(rows[1][5], "cmds_config")
        # r32: case 900001 首步（A/B/C/D 齐, C=2）
        self.assertEqual(rows[2][0], "202609239000010001")
        self.assertEqual(rows[2][1], "P1")
        self.assertEqual(rows[2][2], 2)
        # r33: 第二步只有 E-I，A 列为空
        self.assertIsNone(rows[3][0])
        self.assertEqual(rows[3][4], "check_point")
        # r34: case 间空行
        self.assertTrue(all(v is None for v in rows[4]))
        # r35: case 900002 首步
        self.assertEqual(rows[5][0], "202609239000020001")
        # r38: 哨兵 case（P9, time sleep 1）；其后全空
        self.assertEqual(rows[8][0], "999999999999999")
        self.assertEqual(rows[8][1], "P9")
        self.assertEqual(rows[8][4], "time")
        self.assertEqual(rows[8][5], "sleep")
        self.assertTrue(all(v is None for v in rows[10]), rows[10])
        wb.close()

    def test_marker_and_defined_name_survive(self):
        wb = load_workbook(self.target)
        ws = wb["执行"]
        self.assertEqual(ws.cell(row=1, column=1).value, "IST_EXCEL_CONTRACT")
        self.assertEqual(ws.cell(row=1, column=2).value, "ist.excel.runtime")
        dn = wb.defined_names.get("IST_EXECUTION_SHEET")
        self.assertEqual(list(dn.destinations), [("执行", "$A$29:$I$29")])
        wb.close()

    def test_reject_tampered_contract_marker(self):
        wb = load_workbook(REPO_ROOT / "cex_core" / "templates" / "case_template.xlsx")
        wb["执行"].cell(row=1, column=3).value = "0" * 64
        with self.assertRaises(ExcelContractError):
            resolve_execution_sheet(wb)
        wb.close()

    def test_no_legacy_fallback(self):
        wb = load_workbook(REPO_ROOT / "cex_core" / "templates" / "case_template.xlsx")
        del wb.defined_names["IST_EXECUTION_SHEET"]
        wb["执行"].cell(row=1, column=1).value = None
        with self.assertRaises(ExcelContractError):
            resolve_execution_sheet(wb)
        wb.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
