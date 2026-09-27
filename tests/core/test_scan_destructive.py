"""自毁命令扫描：规则只从 grammar 读；读不到就拒绝；E 列 APV 的每一步都查，不按 F 列分辖区。

扫描要看见框架真正会执行的东西（knowledge/framework/mirror/lib/test_xlsx.py）：
- G 列按框架的参数解析取值：`"clear config all"`、`cmd=clear config all` 交给设备的都是那条命令；
- 逐行查到表尾：999999999999999 伪案的行框架照样执行；
- 公式格框架执行缓存值：一律报出；
- 框架整本载入工作簿：表里声明的 dimension 藏不住行。
"""

from __future__ import annotations

import ast
import io
import re
import shutil
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conftest import INFOTEST_ROOT, REPO_ROOT

from cex_core.scan_destructive import (
    FORMULA_RULE,
    DestructiveRulesUnavailable,
    command_forms,
    load_patterns,
    scan_lines,
    scan_workbook,
)

GRAMMAR = {"destructive_commands": {"patterns": [
    r"^clear\s+config\s+all\b", r"\breboot\b", r"\bshutdown\b"]}}
TEMPLATE = REPO_ROOT / "cex_core" / "templates" / "case_template.xlsx"
SENTINEL = "999999999999999"
AID1, AID2 = "202609260000000001", "202609260000000002"


def _workbook(tmp_path: Path, rows, *, formula=None, dimension: str | None = None) -> Path:
    """模板里追加行 (A, C, E, F, G)；formula=(行偏移, 列, 公式, 缓存值) 写成带缓存值的公式格；
    dimension 改写执行页 XML 里声明的 dimension。"""
    from cex_core.ist_emit.excel_contract import resolve_execution_sheet

    target = tmp_path / "case.xlsx"
    shutil.copy2(TEMPLATE, target)
    wb = load_workbook(target)
    ws, _ = resolve_execution_sheet(wb, allow_legacy=False)
    start = ws.max_row + 1
    for offset, row in enumerate(rows):
        for column, value in zip((1, 3, 5, 6, 7), row):
            if value is not None:
                ws.cell(row=start + offset, column=column, value=value)
    ref = None
    if formula:
        ref = f"{formula[1]}{start + formula[0]}"
        ws[ref] = formula[2]
    wb.save(target)
    if formula or dimension:
        source = zipfile.ZipFile(target)
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as out:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename.startswith("xl/worksheets/sheet"):
                    text = data.decode("utf-8")
                    if formula:
                        cell = re.compile(f'<c r="{ref}"' + r'[^>]*><f>(.*?)</f>(?:<v\s*/>|<v></v>)?</c>')
                        assert cell.search(text), "formula cell not written"
                        text = cell.sub(lambda m: f'<c r="{ref}" t="str"><f>{m.group(1)}</f>'
                                                  f"<v>{formula[3]}</v></c>", text, count=1)
                    if dimension:
                        text, count = re.subn(r'<dimension ref="[^"]*"\s*/>',
                                              f'<dimension ref="{dimension}"/>', text, count=1)
                        assert count == 1
                    data = text.encode("utf-8")
                out.writestr(info, data)
        source.close()
        target.write_bytes(buffer.getvalue())
    return target


def test_fails_closed_without_rules(tmp_path):
    for grammar in ({}, {"destructive_commands": {"patterns": []}},
                    {"destructive_commands": {"patterns": ["("]}}, tmp_path / "missing.json"):
        with pytest.raises(DestructiveRulesUnavailable):
            load_patterns(grammar)


def test_scan_lines_hits_and_misses():
    patterns = load_patterns(GRAMMAR)
    findings = scan_lines([("a", "clear config all"), ("b", "system reboot noninteractive"),
                           ("c", "clear slb real r1"), ("d", "no slb real http r1")], patterns)
    assert [f["where"] for f in findings] == ["a", "b"]


def test_scan_workbook_checks_every_apv_step_regardless_of_method(tmp_path):
    target = _workbook(tmp_path, [
        (None, None, "APV_1", "cmd", "system reboot noninteractive"),
        (None, None, "APV_1", "cmd_config", "slb real http r1 10.0.0.1 80\nclear config all"),
        (None, None, "check_point", "found", "reboot"),
        (None, None, "APV_1", "cmd_config", "clear slb real r1")])
    findings = scan_workbook(target, GRAMMAR)
    assert [f["command"] for f in findings] == ["system reboot noninteractive", "clear config all"]


def test_quoted_and_keyword_arguments_are_matched_as_the_framework_passes_them(tmp_path):
    """框架 get_parameter 剥一层成对引号、`键=值` 取值：设备收到的都是 clear config all。"""
    target = _workbook(tmp_path, [
        (AID1, 2, "APV_0", "cmd_config", '"clear config all"'),
        (None, 2, "APV_0", "cmd_config", "'clear config all'"),
        (None, 2, "APV_0", "cmd_config", "cmd=clear config all"),
        (None, 2, "APV_0", "cmd_config", 'cmd = "clear config all"'),
        (None, 2, "APV_0", "cmd", 'show version, prompt="#"'),
        (SENTINEL, 2, "time", "sleep", "1")])
    findings = scan_workbook(target, GRAMMAR)
    assert len(findings) == 4
    assert {f["executed"] for f in findings} == {"clear config all"}
    assert [f["command"] for f in findings] == [
        '"clear config all"', "'clear config all'", "cmd=clear config all",
        'cmd = "clear config all"']


def test_rows_of_the_trailing_pseudo_case_are_scanned(tmp_path):
    """999999999999999 伪案的行在文件末尾照样执行，扫描不在它那里停。"""
    target = _workbook(tmp_path, [
        (AID1, 2, "APV_0", "cmd", "show version"),
        (SENTINEL, 2, "time", "sleep", "1"),
        (None, None, "APV_0", "cmd_config", "clear config all"),
        (AID2, 2, "APV_0", "cmd", "system reboot"),
        (SENTINEL, 2, "time", "sleep", "1")])
    findings = scan_workbook(target, GRAMMAR)
    assert [f["command"] for f in findings] == ["clear config all", "system reboot"]


def test_formula_cells_are_reported_with_the_value_the_framework_would_run(tmp_path):
    """公式看着是 show version，缓存值是 clear config all：框架执行缓存值。"""
    target = _workbook(tmp_path, [(AID1, 2, "APV_0", "cmd_config", None),
                                  (SENTINEL, 2, "time", "sleep", "1")],
                       formula=(0, "G", '="show version"', "clear config all"))
    findings = scan_workbook(target, GRAMMAR)
    formula = [f for f in findings if f["rule"] == FORMULA_RULE]
    assert len(formula) == 1 and formula[0]["command"] == "clear config all"
    assert formula[0]["formula"] == '="show version"'
    assert any(f["rule"] != FORMULA_RULE and f["command"] == "clear config all" for f in findings)


def test_a_formula_in_any_column_is_reported_even_when_harmless(tmp_path):
    target = _workbook(tmp_path, [(AID1, 2, "APV_0", "cmd", "show version"),
                                  (SENTINEL, 2, "time", "sleep", "1")],
                       formula=(0, "D", '="desc"', "desc"))
    findings = scan_workbook(target, GRAMMAR)
    assert [(f["rule"], f["where"].split("!")[1][0]) for f in findings] == [(FORMULA_RULE, "D")]


def test_a_forged_dimension_record_does_not_hide_rows(tmp_path):
    """只读模式按声明的 dimension 截行；框架整本载入照样执行后面的行。"""
    target = _workbook(tmp_path, [(AID1, 2, "APV_0", "cmd_config", "clear config all"),
                                  (SENTINEL, 2, "time", "sleep", "1")], dimension="A1:I29")
    assert [f["command"] for f in scan_workbook(target, GRAMMAR)] == ["clear config all"]


def test_clean_workbook_has_no_findings(tmp_path):
    target = _workbook(tmp_path, [(AID1, 2, "APV_0", "cmd_config", "slb real http r1 10.0.0.1 80"),
                                  (None, 2, "APV_0", "cmd", 'show slb real, prompt="#"'),
                                  (SENTINEL, 2, "time", "sleep", "1")])
    assert scan_workbook(target, GRAMMAR) == []


def test_repo_grammar_is_not_bundled(tmp_path):
    """规则不写在 cex_core 里：没有 grammar 文件就没有规则。"""
    assert not list(Path(REPO_ROOT / "cex_core").rglob("domain_grammar*.json"))


MIRROR_RUNNER = INFOTEST_ROOT / "knowledge" / "framework" / "mirror" / "lib" / "test_xlsx.py"
_G_SAMPLES = [
    "clear config all", '"clear config all"', "'clear config all'", "cmd=clear config all",
    'cmd = "clear config all"', 'show version, "x, y"', "a,,b", " spaced ,  'q' ",
    'cmd="a=b", prompt=#', "1x=clear config all", 'x\\"y, z', "a='b', c=\"d\"",
    "no quotes = here", '"unclosed', "=leading", "k=", "k='v' tail",
]


@pytest.mark.skipif(not MIRROR_RUNNER.is_file(),
                    reason=f"找不到 InfoTest 框架镜像 {MIRROR_RUNNER}（设 INFOTEST_ROOT）")
def test_command_forms_follow_the_framework_parameter_parser():
    """与框架 get_parameter 逐条对拍：扫描查的参数串 = 框架交给设备的位置参数与关键字值。"""
    tree = ast.parse(MIRROR_RUNNER.read_text(encoding="utf-8"))
    wanted = {"_split_parameter_parts", "_unquote_parameter", "_keyword_split", "get_parameter"}
    body = [node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name in wanted]
    assert {node.name for node in body} == wanted
    namespace: dict = {"re": re}
    exec(compile(ast.Module(body=body, type_ignores=[]), str(MIRROR_RUNNER), "exec"),  # noqa: S102
         namespace)
    for sample in _G_SAMPLES:
        try:
            parameters, kwargs = namespace["get_parameter"](sample)
        except ValueError:
            assert command_forms(sample) == [sample.strip()], sample
            continue
        passed = [str(v) for v in [*parameters, *kwargs.values()] if str(v)]
        forms = command_forms(sample)
        assert forms[0] == sample.strip()
        assert set(passed) <= set(forms) and set(forms) <= set(passed) | {sample.strip()}, sample
