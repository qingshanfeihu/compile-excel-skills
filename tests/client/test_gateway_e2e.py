# ruff: noqa: F811 — 从 test_server_e2e 引入的 server 夹具按 pytest 惯例作参数名
"""cex_client 经真网关上机的端到端测试。

真 compile-excel-server（uvicorn 子进程，账号、令牌、数据包、客户端常量）+ 真网关服务（同级 server 仓的
gateway 包，线程里跑）+ 假测试框架（server 仓网关测试夹具里的假 test_xlsx / 假结果库）。
流程：登录 → 取客户端常量（网关地址）→ 租床 → 自检 → 提交 → 轮询 → 结果与回执 → 被拒的自毁用例 →
run_device.py 一次跑完 → 释放租约。找不到 server 仓或缺依赖就跳过。
"""

from __future__ import annotations

import importlib.util
import json
import os
import shutil
import subprocess
import sys
import textwrap
import threading
import time
from pathlib import Path

import pytest

from conftest import REPO_ROOT, SERVER_ROOT

pytestmark = pytest.mark.skipif(
    not (SERVER_ROOT / "gateway" / "service.py").is_file(),
    reason=f"找不到含网关的 compile-excel-server 检出 {SERVER_ROOT}")

from test_server_e2e import BUILD, _login, server  # noqa: E402,F401 — 复用真服务端夹具


def _load_gateway_fixtures():
    spec = importlib.util.spec_from_file_location(
        "ces_gateway_fixtures", SERVER_ROOT / "tests" / "gateway" / "conftest.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _workbook(path: Path, rows: list[tuple[str, str, str, str]]) -> Path:
    from openpyxl import load_workbook

    from cex_core.ist_emit.excel_contract import resolve_execution_sheet

    path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REPO_ROOT / "cex_core" / "templates" / "case_template.xlsx", path)
    wb = load_workbook(path)
    ws, _ = resolve_execution_sheet(wb, allow_legacy=False)
    start = ws.max_row + 1
    for offset, (autoid, device, method, command) in enumerate(rows):
        if autoid:
            ws.cell(row=start + offset, column=1, value=autoid)
        ws.cell(row=start + offset, column=5, value=device)
        ws.cell(row=start + offset, column=6, value=method)
        ws.cell(row=start + offset, column=7, value=command)
    wb.save(path)
    return path


@pytest.fixture(scope="module")
def gateway_up(server, tmp_path_factory):
    pytest.importorskip("openpyxl")
    fx = _load_gateway_fixtures()
    root = tmp_path_factory.mktemp("gw")
    apv = root / "apv_src"
    (apv / "lib").mkdir(parents=True)
    (apv / "conf").mkdir()
    (apv / "smoke_test" / "sdns").mkdir(parents=True)
    (apv / "conftest.py").write_text(fx.FAKE_CONFTEST, encoding="utf-8")
    (apv / "lib" / "test_xlsx.py").write_text(fx.FAKE_TEST_XLSX, encoding="utf-8")
    (apv / "lib" / "mysqldb.py").write_text(fx.FAKE_MYSQLDB, encoding="utf-8")
    (apv / "lib" / "__init__.py").write_text("", encoding="utf-8")
    (apv / "conf" / "bed.conf").write_text(textwrap.dedent("""\
        [comm]
        ssh_ips = 127.0.0.1
        [other]
        mysql_ip = 127.0.0.1
        [array_ustack]
        user = admin
        passwd = devpass
        """), encoding="utf-8")
    secret = root / "gateway.secret"
    proc = subprocess.run(
        [sys.executable, str(SERVER_ROOT / "ces_main.py"), "clients", "add", "gateway",
         "--scopes", "introspect bundles:read", "--data", str(server["data"]),
         "--out", str(secret)], capture_output=True, text=True, timeout=60)
    assert proc.returncode == 0, proc.stdout + proc.stderr
    config = root / "gateway.toml"
    config.write_text(textwrap.dedent(f"""\
        [server]
        url = "{server['base']}"
        client_id = "gateway"
        client_secret_file = "{secret}"
        build = "{BUILD}"
        [listen]
        port = 0
        [framework]
        apv_src = "{apv}"
        py38 = "{sys.executable}"
        conf_name = "bed"
        staging_parent = "{apv / 'smoke_test' / 'sdns'}"
        default_module = "sdns"
        run_max_s = 60
        [state]
        dir = "{root / 'state'}"
        """), encoding="utf-8")
    sys.path.insert(0, str(SERVER_ROOT))
    try:
        import client_config as server_client_config
        from gateway.config import load
        from gateway.service import build_server
        from gateway.tools import Gateway
    finally:
        sys.path.remove(str(SERVER_ROOT))
    httpd = build_server(Gateway(load(config)))
    thread = threading.Thread(target=httpd.serve_forever, daemon=True)
    thread.start()
    url = f"http://127.0.0.1:{httpd.server_address[1]}/mcp"
    server_client_config.set_key(server["data"], "gateway.url", url)
    yield {"url": url, "apv": apv}
    httpd.shutdown()
    httpd.server_close()


def test_folder_to_device_results_through_the_gateway(server, gateway_up, tmp_path,
                                                      monkeypatch):
    from cex_client import tools

    monkeypatch.delenv("CEX_WORKSPACE", raising=False)
    ws_dir = tmp_path / "project"
    ws_dir.mkdir()
    assert tools.call("cex_init", {"workspace": str(ws_dir), "server": server["base"],
                                   "device_build": BUILD})["ok"]
    _login(ws_dir, server)
    assert tools.call("cex_bed_lease", {"workspace": str(ws_dir), "action": "acquire"})["ok"] is False
    config = tools.call("cex_client_config", {"workspace": str(ws_dir)})
    assert config["config"]["gateway"]["url"] == gateway_up["url"]

    lease = tools.call("cex_bed_lease", {"workspace": str(ws_dir), "action": "acquire"})
    assert lease["ok"], lease
    assert stat_mode(ws_dir / ".compile-excel" / "lease.json") == 0o600
    prepared = tools.call("cex_env_prepare", {"workspace": str(ws_dir)})
    assert prepared["ok"] and {c["check"] for c in prepared["checks"]} >= {
        "framework_files", "framework_conf", "destructive_rules", "device_build"}

    xlsx = _workbook(ws_dir / "compile_outputs" / "b1" / "case.xlsx", [
        ("202609240000000101", "APV_1", "cmd_config", "slb real http r1 10.0.0.1 80"),
        ("202609240000000102", "APV_1", "cmd", "show slb real")])
    submitted = tools.call("cex_case_submit", {"workspace": str(ws_dir),
                                               "xlsx": "compile_outputs/b1/case.xlsx"})
    assert submitted["ok"], submitted
    deadline = time.time() + 60
    while tools.call("cex_case_status", {"workspace": str(ws_dir),
                                         "task_id": submitted["task_id"]}).get("state") != "done":
        assert time.time() < deadline
        time.sleep(0.3)
    results = tools.call("cex_case_results", {"workspace": str(ws_dir),
                                              "task_id": submitted["task_id"]})
    assert results["ok"] and results["totals"] == {"cases": 2, "pass": 2, "fail": 0, "not_run": 0}
    receipt = json.loads((xlsx.parent / "run_results.json").read_text(encoding="utf-8"))
    assert receipt["task_id"] == submitted["task_id"] and receipt["result_channel"] == "ready"
    assert (xlsx.parent / "run_receipt.md").is_file()

    bad = _workbook(ws_dir / "compile_outputs" / "b2" / "case.xlsx", [
        ("202609240000000103", "APV_1", "cmd", "system reboot")])
    refused = tools.call("cex_case_submit", {"workspace": str(ws_dir), "xlsx": str(bad)})
    assert refused["ok"] is False and any("destructive" in p for p in refused["problems"])
    outside = tools.call("cex_case_submit", {"workspace": str(ws_dir),
                                             "xlsx": str(tmp_path / "elsewhere.xlsx")})
    assert outside["ok"] is False and "inside the workspace" in outside["error"]

    third = _workbook(ws_dir / "compile_outputs" / "b3" / "case.xlsx", [
        ("202609240000000104", "APV_1", "cmd", "show version")])
    run = subprocess.run(
        [sys.executable, str(REPO_ROOT / "skills" / "compile-excel" / "scripts" / "run_device.py"),
         "--xlsx", str(third), "--poll-s", "0.3", "--max-s", "60"],
        capture_output=True, text=True, timeout=120,
        env={**os.environ, "CEX_HOME": str(REPO_ROOT)})
    assert run.returncode == 0, run.stdout + run.stderr
    assert json.loads(run.stdout)["totals"]["pass"] == 1
    assert (third.parent / "run_receipt.md").is_file()

    released = tools.call("cex_bed_lease", {"workspace": str(ws_dir), "action": "release"})
    assert released["ok"] and not (ws_dir / ".compile-excel" / "lease.json").exists()
    assert tools.call("cex_case_submit", {"workspace": str(ws_dir),
                                          "xlsx": str(xlsx)})["ok"] is False


def stat_mode(path: Path) -> int:
    return path.stat().st_mode & 0o777
