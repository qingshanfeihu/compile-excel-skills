"""cex_client 对真服务端（同级 compile-excel-server 检出，uvicorn 子进程）的端到端测试。

空文件夹 → cex_init → 设备流登录（授权页填用户名 + 访问码）→ cex_sync（逐件校验、增量、篡改拒收、
断网回退）→ cex_cmd_check / cex_scan_destructive 读同步下来的投影与 grammar →
MCP 代理与 cex_tool 走同一套工具 → 登出后服务端令牌失效。找不到服务端仓或缺 fastapi 就跳过。
"""

from __future__ import annotations

import json
import os
import shutil
import socket
import stat
import subprocess
import sys
import time
import urllib.parse
import urllib.request
from pathlib import Path

import pytest

from conftest import REPO_ROOT, SERVER_ROOT

pytestmark = pytest.mark.skipif(
    not (SERVER_ROOT / "server.py").is_file(),
    reason=f"找不到 compile-excel-server 检出 {SERVER_ROOT}（设 CES_SERVER_ROOT）")

BUILD = "E2E_BUILD"
PROJECTION = {"version": "9.9", "device_os_build": "101", "heads": {
    "show version": {"src": "xml", "pmax": 0},
    "slb real http": {"src": "xml", "args": [{"type": "STRING"}, {"type": "IPADDR"},
                                              {"type": "U16"}]},
}}
GRAMMAR = {"destructive_commands": {"patterns": [r"^clear\s+config\s+all\b", r"\breboot\b"]}}


def _free_port() -> int:
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return sock.getsockname()[1]


@pytest.fixture(scope="module")
def server(tmp_path_factory):
    pytest.importorskip("fastapi")
    pytest.importorskip("uvicorn")
    root = tmp_path_factory.mktemp("srv")
    data = root / "data"
    run = lambda *argv: subprocess.run(  # noqa: E731
        [sys.executable, *argv], capture_output=True, text=True, timeout=120, cwd=SERVER_ROOT)
    proc = run(str(SERVER_ROOT / "deploy" / "provision.py"), "--data", str(data), "--sample")
    assert proc.returncode == 0, proc.stdout + proc.stderr
    code_file = root / "alice.code"
    proc = run(str(SERVER_ROOT / "ces_main.py"), "users", "add", "alice", "--data", str(data),
               "--out", str(code_file))
    assert proc.returncode == 0, proc.stdout + proc.stderr

    sys.path.insert(0, str(SERVER_ROOT))
    try:
        from registry import Registry
    finally:
        sys.path.remove(str(SERVER_ROOT))
    reg = Registry(data / "registry")
    src = root / "bundle_src"
    files = {
        "cmdtree/vendor_stdlib_9.9_101.json": json.dumps(PROJECTION).encode(),
        "projections/domain_grammar.json": json.dumps(GRAMMAR).encode(),
        "template/case_template.xlsx":
            (REPO_ROOT / "cex_core" / "templates" / "case_template.xlsx").read_bytes(),
        "spec/docs/规格说明.md": "# 规格\n".encode(),
    }
    entries = []
    for rel, payload in files.items():
        path = src / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(payload)
        blob = reg.put_blob_file(path)
        entries.append({"kind": rel.split("/")[0], "path": rel, "sha256": blob["sha256"],
                        "media_type": "application/octet-stream", "meta": {}})
    result = reg.submit_bundle(BUILD, entries, publisher="test")
    reg.set_channel(BUILD, "stable", result["bundle_id"], "test")

    port = _free_port()
    env = {**os.environ, "CES_DATA_DIR": str(data)}
    proc = subprocess.Popen(
        [sys.executable, "-m", "uvicorn", "server:app", "--host", "127.0.0.1",
         "--port", str(port), "--log-level", "warning"],
        cwd=SERVER_ROOT, env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    base = f"http://127.0.0.1:{port}"
    for _ in range(100):
        try:
            with urllib.request.urlopen(base + "/healthz", timeout=1):
                break
        except OSError:
            time.sleep(0.2)
    else:
        proc.kill()
        pytest.fail("server did not start")
    info = {"base": base, "data": data, "reg": reg, "files": files,
            "code": code_file.read_text(encoding="utf-8").strip(), "proc": proc}
    yield info
    proc.terminate()
    proc.wait(timeout=10)


def _approve(base: str, user_code: str, code: str) -> None:
    body = urllib.parse.urlencode({"user_code": user_code, "username": "alice",
                                   "access_code": code}).encode()
    with urllib.request.urlopen(urllib.request.Request(base + "/activate", data=body),
                                timeout=10) as resp:
        assert resp.status == 200


def _login(ws_dir: Path, server) -> None:
    from cex_client import tools

    started = tools.call("cex_login_start", {"workspace": str(ws_dir)})
    assert started["ok"], started
    pending = tools.call("cex_login_wait", {"workspace": str(ws_dir), "timeout_s": 1})
    assert pending.get("pending") is True
    _approve(server["base"], started["user_code"], server["code"])
    done = tools.call("cex_login_wait", {"workspace": str(ws_dir), "timeout_s": 30})
    assert done["ok"], done


def test_folder_to_synced_bundle_and_tools(server, tmp_path, monkeypatch):
    from cex_client import tools

    monkeypatch.delenv("CEX_WORKSPACE", raising=False)
    ws_dir = tmp_path / "project"
    ws_dir.mkdir()
    assert tools.call("cex_status", {"workspace": str(ws_dir)})["ok"] is False
    init = tools.call("cex_init", {"workspace": str(ws_dir), "server": server["base"],
                                   "device_build": BUILD})
    assert init["ok"], init
    assert tools.call("cex_sync", {"workspace": str(ws_dir)})["ok"] is False  # 未登录

    _login(ws_dir, server)
    token_path = ws_dir / ".compile-excel" / "token.json"
    assert stat.S_IMODE(os.stat(token_path).st_mode) == 0o600
    status = tools.call("cex_status", {"workspace": str(ws_dir)})
    assert status["logged_in"] is True and "bundles:read" in status["scope"]

    first = tools.call("cex_sync", {"workspace": str(ws_dir)})
    assert first["ok"] and first["source"] == "server" and first["downloaded"] == 4, first
    bundle_dir = ws_dir / ".compile-excel" / "bundle" / BUILD
    for rel, payload in server["files"].items():
        assert (bundle_dir / rel).read_bytes() == payload
    again = tools.call("cex_sync", {"workspace": str(ws_dir)})
    assert again["downloaded"] == 0 and again["unchanged"] == 4

    # 本地被改：下次同步重新下载那一个
    (bundle_dir / "spec/docs/规格说明.md").write_text("tampered", encoding="utf-8")
    fixed = tools.call("cex_sync", {"workspace": str(ws_dir)})
    assert fixed["downloaded"] == 1 and fixed["unchanged"] == 3

    check = tools.call("cex_cmd_check", {"workspace": str(ws_dir), "commands": [
        "show version", "slb real http r1 10.0.0.1 80", "slb real http r1 bad 80", "shwo version"]})
    assert check["ok"] and [r["hit"] for r in check["results"]] == [True, True, False, False]

    xlsx = ws_dir / "case.xlsx"
    shutil.copy2(REPO_ROOT / "cex_core" / "templates" / "case_template.xlsx", xlsx)
    from openpyxl import load_workbook

    from cex_core.ist_emit.excel_contract import resolve_execution_sheet

    wb = load_workbook(xlsx)
    ws, _ = resolve_execution_sheet(wb, allow_legacy=False)
    row = ws.max_row + 1
    ws.cell(row=row, column=5, value="APV_1")
    ws.cell(row=row, column=6, value="cmd")
    ws.cell(row=row, column=7, value="system reboot")
    wb.save(xlsx)
    scan = tools.call("cex_scan_destructive", {"workspace": str(ws_dir), "xlsx": "case.xlsx"})
    assert scan["ok"] is False and scan["findings"][0]["command"] == "system reboot"
    outside = tools.call("cex_scan_destructive", {"workspace": str(ws_dir),
                                                  "xlsx": str(tmp_path / "x.xlsx")})
    assert outside["ok"] is False and "inside the workspace" in outside["error"]


def test_server_side_tamper_is_refused(server, tmp_path, monkeypatch):
    from cex_client import tools

    monkeypatch.delenv("CEX_WORKSPACE", raising=False)
    ws_dir = tmp_path / "p2"
    ws_dir.mkdir()
    tools.call("cex_init", {"workspace": str(ws_dir), "server": server["base"],
                            "device_build": BUILD})
    _login(ws_dir, server)
    import hashlib

    sha = hashlib.sha256(server["files"]["spec/docs/规格说明.md"]).hexdigest()
    blob = server["reg"].blob_path(sha)
    original = blob.read_bytes()
    os.chmod(blob, 0o644)
    try:
        blob.write_bytes(b"evil")
        result = tools.call("cex_sync", {"workspace": str(ws_dir)})
        assert result["ok"] is False and "SHA256 mismatch" in result["error"]
        target = ws_dir / ".compile-excel" / "bundle" / BUILD / "spec/docs/规格说明.md"
        assert not target.exists()
        assert not list(target.parent.glob("*.part"))
    finally:
        blob.write_bytes(original)
        os.chmod(blob, 0o444)


def test_mcp_proxy_and_cex_tool_share_the_tool_set(server, tmp_path):
    ws_dir = tmp_path / "p3"
    ws_dir.mkdir()
    env = {**os.environ, "CEX_WORKSPACE": str(ws_dir)}
    cli = subprocess.run(
        [sys.executable, str(REPO_ROOT / "bin" / "cex_tool"), "cex_init",
         json.dumps({"workspace": str(ws_dir), "server": server["base"], "device_build": BUILD})],
        capture_output=True, text=True, timeout=60, env=env)
    assert cli.returncode == 0, cli.stdout + cli.stderr

    messages = [
        {"jsonrpc": "2.0", "id": 1, "method": "initialize",
         "params": {"protocolVersion": "2025-06-18", "capabilities": {},
                    "clientInfo": {"name": "t", "version": "0"}}},
        {"jsonrpc": "2.0", "method": "notifications/initialized"},
        {"jsonrpc": "2.0", "id": 2, "method": "tools/list"},
        {"jsonrpc": "2.0", "id": 3, "method": "tools/call",
         "params": {"name": "cex_status", "arguments": {}}},
        {"jsonrpc": "2.0", "id": 4, "method": "tools/call",
         "params": {"name": "no_such_tool", "arguments": {}}},
        {"jsonrpc": "2.0", "id": 5, "method": "bogus/method"},
    ]
    proc = subprocess.run(
        [sys.executable, str(REPO_ROOT / "bin" / "cex_mcp_proxy.py")],
        input="\n".join(json.dumps(m) for m in messages) + "\n",
        capture_output=True, text=True, timeout=60, env=env)
    replies = {r["id"]: r for r in map(json.loads, proc.stdout.splitlines())}
    assert set(replies) == {1, 2, 3, 4, 5}, proc.stdout + proc.stderr
    assert replies[1]["result"]["serverInfo"]["name"] == "compile-excel"
    listed = {t["name"] for t in replies[2]["result"]["tools"]}
    from cex_client import tools

    assert listed == set(tools.TOOLS) == {s["name"] for s in tools.load_specs()}
    status = replies[3]["result"]["structuredContent"]
    assert status["workspace"] == str(ws_dir.resolve()) and status["logged_in"] is False
    assert replies[4]["result"]["isError"] is True
    assert replies[5]["error"]["code"] == -32601


def test_logout_revokes_on_server_and_offline_falls_back_to_cache(server, tmp_path, monkeypatch):
    from cex_client import tools

    monkeypatch.delenv("CEX_WORKSPACE", raising=False)
    ws_dir = tmp_path / "p4"
    ws_dir.mkdir()
    tools.call("cex_init", {"workspace": str(ws_dir), "server": server["base"],
                            "device_build": BUILD})
    _login(ws_dir, server)
    assert tools.call("cex_sync", {"workspace": str(ws_dir)})["source"] == "server"
    token = json.loads((ws_dir / ".compile-excel" / "token.json").read_text(encoding="utf-8"))
    out = tools.call("cex_logout", {"workspace": str(ws_dir)})
    assert out["ok"] and out["revoked_on_server"] is True
    req = urllib.request.Request(server["base"] + "/v1/whoami",
                                 headers={"Authorization": f"Bearer {token['access_token']}"})
    with pytest.raises(urllib.error.HTTPError) as exc:
        urllib.request.urlopen(req, timeout=10)
    assert exc.value.code == 401

    # 断网回退：换一个不可达的端口，缓存完整就用缓存并明说
    cfg_path = ws_dir / ".compile-excel" / "config.json"
    cfg = json.loads(cfg_path.read_text(encoding="utf-8"))
    cfg["server"] = f"http://127.0.0.1:{_free_port()}"
    cfg_path.write_text(json.dumps(cfg), encoding="utf-8")
    from cex_client import workspace as wsmod

    wsmod.write_private_json(ws_dir / ".compile-excel" / "token.json", {
        "access_token": "x", "refresh_token": "y", "expires_at": time.time() + 3600,
        "server": cfg["server"]})
    offline = tools.call("cex_sync", {"workspace": str(ws_dir)})
    assert offline["ok"] and offline["source"] == "cache" and "unreachable" in offline["note"]
