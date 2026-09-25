"""编写阶段对拍：同一份脑图与机械脑图，客户端编写链与 InfoTest 引擎交付的结果逐行一致。

输入是 InfoTest 已交付的一批（CEX_AUTHOR_REFERENCE_BATCH 指向 workspace/outputs/<批名>）：
- 数据包由 compile-excel-server 的 tools/publish_data_dir.py 从 InfoTest 数据目录现发布
  （命令树脱敏后重推导，判据台账与 SSL 生命周期证据一起发）；
- 客户端：重组（spec='none'）→ 提交该批的机械脑图案并密封 → cex_author_prepare 出契约卡 →
  逐案提交 InfoTest 交付的机械用例（去掉 seal 与 binding）→ 全部封存 → cex_author_emit；
- 断言：每案封存、verify_batch 全过、展开出的每一行（E/F/G）与 InfoTest 交付的工作簿一致。

床事实用 InfoTest 数据目录里的 network_topology.json（与交付那批同一张床）。需要 InfoTest
检出、交付过的批目录与已同步 vendor 的服务端检出；缺任何一个就跳过。
"""

from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from conftest import INFOTEST_ROOT, REPO_ROOT, SERVER_ROOT

REFERENCE = Path(os.environ.get("CEX_AUTHOR_REFERENCE_BATCH") or "/nonexistent")
RAW_BUILD = os.environ.get("CEX_AUTHOR_RAW_BUILD", "")
pytestmark = pytest.mark.skipif(
    not (REFERENCE / "machine_mindmap.json").is_file()
    or not (SERVER_ROOT / "gateway" / "vendor" / "cex_core" / "__init__.py").is_file()
    or not (INFOTEST_ROOT / "knowledge" / "data" / "auto_env" / "network_topology.json").is_file()
    or not RAW_BUILD,
    reason="需要 CEX_AUTHOR_REFERENCE_BATCH（InfoTest 交付过的批目录）、CEX_AUTHOR_RAW_BUILD、"
           "InfoTest 检出与已同步 vendor 的服务端检出")


def _publish(tmp_path: Path) -> Path:
    out = tmp_path / "bundle"
    proc = subprocess.run(
        [sys.executable, str(SERVER_ROOT / "tools" / "publish_data_dir.py"),
         "--data-root", str(INFOTEST_ROOT), "--raw-build", RAW_BUILD,
         "--manual-version", re.search(r"(\d+\.\d+\.\d+)\.\d+$", RAW_BUILD).group(1),
         "--out-dir", str(out)], capture_output=True, text=True, timeout=1800, check=False)
    assert proc.returncode == 0, (proc.stdout + proc.stderr)[-3000:]
    return out


def _workbook_rows(xlsx: Path) -> dict[str, list[tuple[str, str, str]]]:
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    rows: dict[str, list[tuple[str, str, str]]] = {}
    current = None
    for ws in wb.worksheets:
        values = list(ws.iter_rows(values_only=True))
        header = next((i for i, r in enumerate(values) if r and r[0] == "自动化ID"), None)
        if header is None:
            continue
        for row in values[header + 1:]:
            if not row or all(cell is None for cell in row[:9]):
                continue
            if row[0]:
                current = str(row[0])
                rows[current] = []
            if current and any(row[4:7]):
                rows[current].append(tuple(str(cell or "") for cell in row[4:7]))
    wb.close()
    return rows


def _reference_types() -> dict[str, str]:
    """交付那批契约卡里，引擎对每个判据形状的裁定。"""
    types: dict[str, str] = {}
    for card in sorted((REFERENCE / "contracts").glob("*.json")):
        for item in json.loads(card.read_text(encoding="utf-8")).get("expectations") or []:
            claim = (item or {}).get("normalized_claim") or {}
            if claim.get("shape_key") and claim.get("criterion_type"):
                types[str(claim["shape_key"])] = str(claim["criterion_type"])
    return types


def _run_flow(tmp_path: Path, *, drop_ledger: bool) -> dict:
    bundle_dir = _publish(tmp_path)
    if drop_ledger:
        # 不发判据台账种子：每个判据形状都要在客户端裁定一遍
        (bundle_dir / "projections" / "criterion_author_rules.jsonl").unlink()
        manifest = json.loads((bundle_dir / "manifest.json").read_text(encoding="utf-8"))
        manifest["entries"] = [e for e in manifest["entries"]
                               if e["path"] != "projections/criterion_author_rules.jsonl"]
        (bundle_dir / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    types = tmp_path / "types.json"
    types.write_text(json.dumps(_reference_types()), encoding="utf-8")
    ws_root = tmp_path / "ws"
    batch = "parity"
    ws_root.mkdir()
    shutil.copyfile(REFERENCE / "mindmap_source.json", ws_root / "mindmap.json")
    script = f'''
import json, shutil, sys
from pathlib import Path
sys.path.insert(0, {str(REPO_ROOT)!r})
from cex_client import author, engine_env, recompose, workspace as wsmod
ws = wsmod.init(Path({str(ws_root)!r}), server="https://ces.example.test", device_build="B_1")
shutil.copytree({str(bundle_dir)!r}, ws.bundle_dir(), dirs_exist_ok=True)
topology = engine_env.topology_path(ws)
topology.parent.mkdir(parents=True, exist_ok=True)
shutil.copyfile({str(INFOTEST_ROOT / "knowledge/data/auto_env/network_topology.json")!r}, topology)
prep = recompose.prepare(ws, "mindmap.json", out_name={batch!r}, spec="none")
cases = json.loads(Path({str(REFERENCE / "machine_mindmap.json")!r}).read_text())["cases"]
submitted = recompose.submit_cases(ws, {batch!r}, cases)
sealed = recompose.seal(ws, {batch!r})
cards = author.prepare(ws, {batch!r})
types = json.loads(Path({str(types)!r}).read_text())
recorded = []
while cards.get("phase") == "criterion_pending" or cards.get("pending_shapes"):
    shape = cards["pending_shapes"][0]
    judgment = {{"criterion_type": types[shape["shape_key"]],
                 "rationale": "The engine delivery typed this verdict shape the same way.",
                 "disclosure": "按交付批的同一裁定归类。"}}
    cards = author.criterion_record(ws, {batch!r}, shape["shape_key"], judgment)
    recorded.append([shape["shape_key"], cards.get("status")])
results = {{}}
for aid in [c["autoid"] for c in cards.get("cases") or []]:
    body = json.loads((Path({str(REFERENCE)!r}) / "delivered" / aid / "mechanical_case.json").read_text())
    body.pop("seal", None)
    body["binding"] = {{}}
    results[aid] = author.submit_case(ws, {batch!r}, body)
try:
    emitted = author.emit(ws, {batch!r})
except Exception as exc:
    emitted = {{"ok": False, "error": f"{{type(exc).__name__}}: {{exc}}"}}
print(json.dumps({{"prep": prep.get("ok"), "submitted": submitted.get("status"),
                   "sealed": sealed.get("ok"), "phase": cards.get("phase"), "recorded": recorded,
                   "cases": [c["autoid"] for c in cards.get("cases") or []],
                   "results": {{a: [r.get("status"), r.get("violations")] for a, r in results.items()}},
                   "emit": emitted}}, ensure_ascii=False))
'''
    # 引擎在 PYTEST_CURRENT_TEST 下把 runtime/（判据台账）改指临时目录；子进程是真实使用场景
    env = {k: v for k, v in os.environ.items() if k != "PYTEST_CURRENT_TEST"}
    proc = subprocess.run([sys.executable, "-c", script], capture_output=True, text=True,
                          timeout=1800, check=False,
                          env={**env, "CEX_ENGINE_DATA_ROOT": "", "IST_DEVICE_OS_BUILD": ""})
    assert proc.returncode == 0, proc.stderr[-4000:]
    return json.loads(proc.stdout.strip().splitlines()[-1])


def _assert_reproduces_the_delivery(result: dict) -> None:
    assert result["phase"] == "published" and result["cases"], result
    unsealed = {aid: r for aid, r in result["results"].items() if r[0] != "sealed"}
    assert unsealed == {}, json.dumps(unsealed, ensure_ascii=False)[:4000]
    assert result["emit"]["ok"], result["emit"]
    ours = _workbook_rows(Path(result["emit"]["xlsx"]))
    theirs = _workbook_rows(REFERENCE / "case.xlsx")
    for aid in result["cases"]:
        assert ours.get(aid) == theirs.get(aid), aid


def test_client_authoring_reproduces_the_engine_delivery(tmp_path):
    result = _run_flow(tmp_path, drop_ledger=False)
    assert result["recorded"] == [], "every shape is answered by the shipped criterion ledger"
    _assert_reproduces_the_delivery(result)


def test_criteria_typed_through_the_client_reproduce_the_delivery(tmp_path):
    """没有台账种子时每个形状停在 criterion_pending；按交付那批的同一裁定逐个记录，
    引擎复核入账后发布契约卡，交付的机械用例照样逐案封存、出件逐行一致。"""
    result = _run_flow(tmp_path, drop_ledger=True)
    assert result["recorded"], "without the ledger seed the shapes need adjudication"
    assert all(status == "recorded" for _key, status in result["recorded"]), result["recorded"]
    _assert_reproduces_the_delivery(result)
