"""上机：经网关提交工作簿、轮询、取回框架判定，并把回执写回工作簿所在目录。

判定语义归框架：verdict 来自框架结果库；pytest 的 “1 passed” 不代表断言通过。
非 pass 的 case 附框架日志作归因证据：网关给多少留多少（失败断言的实际回显在日志中段，只留尾巴
会切掉它），并把每个失败断言连同它当时的回显摘成 failed_checks；早于投递时间的日志网关已标
stale，这里不拿它当证据。
投递时把工作簿旁 provenance.json 的逐案 check_point 指纹、以及旁边 cases.json 的逐案卷面指纹
（fingerprints.case_fingerprints）记进任务记录，取结果时写进 run_results.json：返工闸以“上一轮真上机
的卷面”为准比对——编写阶段的 cex_author_emit 会在过闸之前重写 provenance.json 与 cases.json。
回执格式沿用原 run_device（ist.excel.device-run-result）：run_results.json 机读、run_receipt.md 人读。
同一工作簿投过多次时，只有最近那次的结果写回执；取较早那次的结果照常返回，但不覆盖回执
（receipt 为 null 并说明），免得旧结果盖掉返工闸要看的最新一轮。回执不跟随符号链接写。

库里的判定要有本案日志里框架的收尾作证（framework_verdict）：框架在 begin case 时就给新案写一行
结果，值是上一个案的结果，到 end case 才改成本案的；跑到一半被杀（超时、OOM、人工）或收尾时登
不上设备，库里那行就停在占位值上。没有收尾、或收尾与库里那行不一致，判 broken，不当 pass/fail。
判 pass 的案再扫一遍执行失败回显（数据包 domain_grammar 的 exec_failure_markers，与 InfoTest 批量
上机同一份规则与豁免）：配置没生效、断言照样命中的空真 pass 判 broken。
runner 死了（channel runner_lost）这一轮不会有判定，不写回执。
"""

from __future__ import annotations

import base64
import datetime as _dt
import hashlib
import json
import re
import time
from pathlib import Path
from typing import Any

from . import bundle, gateway
from .errors import ClientError
from .fingerprints import case_fingerprints
from .workspace import Workspace, write_file_safely

RESULT_SCHEMA = "ist.excel.device-run-result"
MAX_XLSX_BYTES = 32 * 1024 * 1024
SENTINEL_AUTOID = "999999999999999"
_AUTOID_RE = re.compile(r"^\d{12,24}$")
RUNNER_LOST = ("the gateway reports the run as lost: its runner exited without recording an end "
               "(gateway restart, OOM or an operator kill), so this run has no verdicts; "
               "resubmit the workbook")
# 回执里每案日志之外的会话转储（设备 CLI 会话、触发机会话）落在这里，按 autoid 分目录
EVIDENCE_DIR = "evidence"

# 机械归因只判协议级事实；语义归因留给会话（与原 run_device 相同的标记集）
_G_LAYER_MARKERS = (
    "% invalid", "% unrecognized", "% unknown", "% error",
    "syntax error", "invalid input", "command not found",
)
_TRANSIENT_MARKERS = (
    "timeout", "timed out", "read_until", "connection reset",
    "connection closed", "device_busy", "traceback", "not reachable",
)

# 框架对每个失败断言打一行 "#### Fail Num N: fail to find <模式> in:"，其后几行是被检命令与实际回显
_FAIL_HEAD_RE = re.compile(r"#### Fail Num \d+: fail to find .* in:\s*$")
_FAIL_CONTEXT_LINES = 8
_RETURN_TAIL_CHARS = 600


def failed_checks(log: str) -> list[str]:
    """每个失败的 check_point 一段：框架的失败行加它当时比对的回显（到下一条 #### 为止）。"""
    lines = (log or "").splitlines()
    out: list[str] = []
    for i, line in enumerate(lines):
        if not _FAIL_HEAD_RE.search(line):
            continue
        block = [line.strip()]
        for follow in lines[i + 1:i + 1 + _FAIL_CONTEXT_LINES]:
            if "####" in follow:
                break
            if follow.strip():
                block.append(follow.strip())
        out.append("\n".join(block))
    return out


_LOG_STAMP_RE = re.compile(r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2} ")


def check_summary(block: str) -> str:
    """回执一格放得下的一句：期望的模式 → 实际回显的最后一行。"""
    lines = [_LOG_STAMP_RE.sub("", ln).strip() for ln in (block or "").splitlines() if ln.strip()]
    if not lines:
        return ""
    head = re.sub(r"^#### Fail Num \d+: ", "", lines[0])
    return f"{head} → {lines[-1]}" if len(lines) > 1 else head


# 框架收尾（lib/check_point.py close() 与 lib/test_xlsx.py parser_case_id）：跑完的案在自己日志的
# 末尾写失败/通过计数和 PASS/FAIL 横幅，紧跟一行 "end case: <autoid>"
_VERDICT_BANNER_RE = re.compile(r"^#{10,}\s+(PASS|FAIL)\s+#{10,}$")
_FAILED_COUNT_RE = re.compile(r"^#{5,}\s+The failed check point num:\s*(\d+)\s+#{4,}$")
_PASSED_COUNT_RE = re.compile(r"^#{5,}\s+The passed check point num:\s*(\d+)\s+#{4,}$")


def framework_verdict(log: str, autoid: str) -> str | None:
    """本案日志里框架自己收尾写下的判定（"pass"/"fail"）；没收尾或收尾不成形就是 None。

    成形的收尾：恰好一行 end case: <autoid>，紧挨在它前面的是 PASS/FAIL 横幅，再往前没有别的
    横幅（一个案只收尾一次）；PASS 时计数行若在，失败数为 0、通过数大于 0（框架判 PASS 的条件）。
    计数行在网关截的日志尾里可能被截掉，所以只在它在的时候核。"""
    lines = [_LOG_STAMP_RE.sub("", ln).strip() for ln in (log or "").splitlines()]
    lines = [ln for ln in lines if ln]
    end_re = re.compile(rf"^#{{7}}\s+end case:\s*{re.escape(str(autoid))}$")
    ends = [i for i, ln in enumerate(lines) if end_re.match(ln)]
    if len(ends) != 1 or ends[0] == 0:
        return None
    end = ends[0]
    banner = _VERDICT_BANNER_RE.match(lines[end - 1])
    if banner is None or any(_VERDICT_BANNER_RE.match(ln) for ln in lines[:end - 1]):
        return None
    verdict = banner.group(1).lower()
    if verdict == "pass":
        failed = [int(m.group(1)) for ln in lines[:end - 1] if (m := _FAILED_COUNT_RE.match(ln))]
        passed = [int(m.group(1)) for ln in lines[:end - 1] if (m := _PASSED_COUNT_RE.match(ln))]
        if len(failed) > 1 or len(passed) > 1 or any(failed) or (passed and passed[0] <= 0):
            return None
    return verdict


# 执行失败回显的帧与豁免：移植 InfoTest batch_tools 的 _anomaly_frames / _anomalies_are_expected
# （含 2026-09-03 帧正文、09-06 会话分帧、09-21 摘要行三次修正）。pass 案日志里出现失败回显，
# 多半是某条配置没生效、断言空真通过（InfoTest 实跑过：config all tftp 恢复静默失败，后面的
# not_found 照样过）；案自己的 found/abs_found 断言就在等这句话的（负向用例）不算。
MAX_FAILURE_ECHO_LINES = 8
_FRAME_ANNOTATION_RE = re.compile(r"^#{3,}")
_FRAME_STEP_RE = re.compile(r"^#{5,}\s")
_FRAME_HOST_PROMPT_RE = re.compile(r"^[A-Za-z][\w.\-]*[#$](?:\s|$)")
_FRAME_SEGMENT_RE = re.compile(r"^(?:=== .+ ===|\[[^\]]*\])$")
_FRAME_SEND_RE = re.compile(r"^\S+ - sends command in (?:config|enable): (.*)$")
_FRAME_TRIGGER_RE = re.compile(r"^\S+ executes command: (.*)$")
_FRAME_PROMPT_RE = re.compile(r"^APV(?:\(config\))?#(.*)$")
# 框架断言摘要行：编号列（Success/Fail Num）与判定词（successed/fail to find）是两个独立轴
_ASSERTION_SUMMARY_RE = re.compile(
    r"^#{3,}\s*(?:Success|Fail)\s+Num\s+\d+\s*:\s*"
    r"(successed|fail)\s+to\s+find\s*:?\s*(.*?)(?:\s+in\s*:\s*)?$")


def failure_echo_lines(log: str, markers: list[str]) -> list[str]:
    """日志里含执行失败回显的行（至多 MAX_FAILURE_ECHO_LINES 条）。"""
    if not markers:
        return []
    hits = [ln.strip() for ln in (log or "").splitlines() if any(m in ln for m in markers)]
    return hits[:MAX_FAILURE_ECHO_LINES]


def _echo_frames(text: str) -> list[tuple[list[str], str]]:
    """按命令派发、步骤标记、提示符把日志切成帧；帧正文不含起始行和框架注解行。"""
    frames: list[tuple[list[str], str]] = []
    cur: list[str] = []
    head = False

    def flush() -> None:
        if not cur:
            return
        body = "\n".join(ln for ln in (cur[1:] if head else cur)
                         if not _FRAME_ANNOTATION_RE.match(ln))
        frames.append((list(cur), body))

    for raw in (text or "").splitlines():
        line = _LOG_STAMP_RE.sub("", raw).rstrip()
        if (_FRAME_SEND_RE.match(line) or _FRAME_TRIGGER_RE.match(line)
                or _FRAME_STEP_RE.match(line) or _FRAME_PROMPT_RE.match(line)
                or _FRAME_HOST_PROMPT_RE.match(line) or _FRAME_SEGMENT_RE.match(line)):
            flush()
            cur = [line]
            head = True
            continue
        cur.append(line)
    flush()
    return frames


def failure_echo_expected(lines: list[str], patterns: list[tuple[str, str]], log: str) -> bool:
    """失败回显是否都由本案自己的 found/abs_found 断言在等（负向用例）；与 InfoTest 同一判据。"""
    if not lines or not patterns:
        return False

    def covered(text: str) -> bool:
        for method, pattern in patterns:
            try:
                matched = (pattern in text if method == "abs_found"
                           else re.search(pattern, text, re.DOTALL) is not None)
            except re.error:
                matched = False
            if matched:
                return True
        return False

    frames = _echo_frames(log) if log else []
    occurrences: dict[str, int] = {}
    for line in lines:
        probe = _LOG_STAMP_RE.sub("", line).strip()
        # 摘要行逐字复读断言的 pattern：只有“成功找到”且 pattern 就是本案某条断言时才算在等它
        summary = _ASSERTION_SUMMARY_RE.match(probe)
        if summary is not None:
            verdict, summary_pattern = summary.group(1), summary.group(2).strip()
            if verdict == "successed" and summary_pattern and any(
                    method in {"found", "abs_found"} and summary_pattern == pattern
                    for method, pattern in patterns):
                continue
            return False
        # 按行及其出现次序绑定帧：同一句失败出现在不同命令帧里，不能借第一帧的证据
        hits = [(body, bool(_FRAME_ANNOTATION_RE.match(raw)))
                for raw_lines, body in frames for raw in raw_lines if raw.strip() == probe]
        occurrence = occurrences.get(probe, 0)
        occurrences[probe] = occurrence + 1
        if frames and occurrence >= len(hits):
            return False
        frame_body, matched_annotation = hits[occurrence] if hits else ("", False)
        if frame_body:
            if covered(frame_body):
                continue
            return False
        if frames and matched_annotation:
            # 帧里只有框架注解、没有设备正文：注解复读 pattern 不能自证
            return False
        if covered(probe or line):
            continue
        return False
    return True


def failure_markers(ws: Workspace) -> list[str] | None:
    """数据包 domain_grammar 里的执行失败回显；读不到就是 None（扫描停用，回执里注明）。"""
    try:
        path = bundle.entry_path(ws, "projections", "domain_grammar")
        grammar = json.loads(path.read_text(encoding="utf-8")) if path else None
    except (OSError, ValueError, ClientError):
        return None
    patterns = ((grammar or {}).get("exec_failure_markers") or {}).get("patterns")
    if not isinstance(patterns, list):
        return None
    return [str(p) for p in patterns if str(p)]


def workbook_facts(data: bytes) -> dict[str, Any]:
    """投递的工作簿里每个真案的 found/abs_found 断言（失败回显豁免用）与有没有卷尾哨兵案。"""
    import io

    from openpyxl import load_workbook

    from cex_core.ist_emit.excel_contract import resolve_execution_sheet

    patterns: dict[str, list[list[str]]] = {}
    sentinel = False
    wb = load_workbook(io.BytesIO(data), read_only=False, data_only=True)
    try:
        sheet, layout = resolve_execution_sheet(wb, allow_legacy=False)
        current = ""
        for row in sheet.iter_rows(min_row=layout.data_start, values_only=True):
            cells = [str(value if value is not None else "").strip() for value in row[:9]]
            cells += [""] * (9 - len(cells))
            if cells[0] == SENTINEL_AUTOID:
                sentinel = True
                current = ""
                continue
            if _AUTOID_RE.match(cells[0]):
                current = cells[0]
                patterns.setdefault(current, [])
            if current and cells[4] == "check_point" and cells[5] in {"found", "abs_found"} \
                    and cells[6]:
                patterns[current].append([cells[5], cells[6]])
    finally:
        wb.close()
    return {"assertion_patterns": patterns, "sentinel": sentinel}


def provenance_fingerprints(prov_path: Path) -> dict[str, str]:
    """provenance.json 的逐案 check_point 指纹；与 rework_gate 的 _entries_fp 同一算法。"""
    try:
        prov = json.loads(prov_path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    out: dict[str, str] = {}
    for autoid, entries in ((prov or {}).get("cases") or {}).items():
        if isinstance(entries, list):
            blob = json.dumps(entries, ensure_ascii=False, sort_keys=True)
            out[str(autoid)] = hashlib.sha256(blob.encode("utf-8")).hexdigest()
    return out


def attribute_fail(detail_tail: str) -> dict[str, str]:
    text = (detail_tail or "").lower()
    for marker in _G_LAYER_MARKERS:
        if marker in text:
            return {"layer": "G", "evidence": marker, "note": "命令/语法层——回显含 CLI 错误标记"}
    for marker in _TRANSIENT_MARKERS:
        if marker in text:
            return {"layer": "transient?", "evidence": marker,
                    "note": "疑似瞬态（超时/连接/忙）——同签名复发则非瞬态，不升格"}
    return {"layer": "undetermined", "evidence": "",
            "note": "机械标记无法裁决，需会话按 detail_tail 语义归因（E/V/产品缺陷）"}


def _now() -> str:
    return _dt.datetime.now().astimezone().isoformat(timespec="seconds")


def cases_fingerprints(cases_path: Path) -> dict[str, str]:
    """工作簿旁 cases.json 的逐案卷面指纹；没有或读不了就是空表。"""
    try:
        doc = json.loads(cases_path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    return case_fingerprints(doc) if isinstance(doc, dict) else {}


def resolve_xlsx(ws: Workspace, xlsx: str) -> Path:
    path = Path(str(xlsx or "")).expanduser()
    if not path.is_absolute():
        path = ws.root / path
    path = path.resolve()
    if ws.root.resolve() not in path.parents or not path.is_file():
        raise ClientError("xlsx must be an existing file inside the workspace")
    if path.stat().st_size > MAX_XLSX_BYTES:
        raise ClientError("workbook is larger than the gateway accepts")
    return path


def submit(ws: Workspace, xlsx: str, module: str | None = None) -> dict[str, Any]:
    path = resolve_xlsx(ws, xlsx)
    data = path.read_bytes()
    args: dict[str, Any] = {**gateway.lease_args(ws),
                            "xlsx_b64": base64.b64encode(data).decode("ascii")}
    if module:
        args["module"] = module
    out = gateway.call_tool(ws, "case_submit", args)
    if out.get("ok"):
        local_sha = hashlib.sha256(data).hexdigest()
        if out.get("sha256") != local_sha:
            raise ClientError("gateway staged a different workbook than the one sent")
        try:
            facts = workbook_facts(data)
        except Exception:  # noqa: BLE001 — 网关已收下这份工作簿；读不出断言只是失败回显扫描少了豁免
            facts = {}
        gateway.remember_task(ws, out["task_id"], {
            "xlsx": path.relative_to(ws.root).as_posix(), "sha256": local_sha,
            "submitted_at": _now(), "case_ids": out.get("case_ids", []),
            **{key: out[key] for key in ("submit_autoid", "module") if out.get(key)},
            **facts,
            "provenance_fp": provenance_fingerprints(path.parent / "provenance.json"),
            "case_fingerprints": cases_fingerprints(path.parent / "cases.json")})
    return out


def status(ws: Workspace, task_id: str) -> dict[str, Any]:
    return gateway.call_tool(ws, "case_status", {"task_id": task_id})


def _submitted_patterns(ws: Workspace, record: dict[str, Any]) -> dict[str, list] | None:
    """这次投递的逐案 found/abs_found 断言：投递时记下的；较早的记录没有，就读工作簿（sha 还对得上
    时）；都没有就是 None。"""
    if isinstance(record.get("assertion_patterns"), dict):
        return record["assertion_patterns"]
    try:
        data = (ws.root / str(record.get("xlsx") or "")).read_bytes()
        if record.get("xlsx") and hashlib.sha256(data).hexdigest() == record.get("sha256"):
            return workbook_facts(data)["assertion_patterns"]
    except Exception:  # noqa: BLE001 — 读不出就当不知道
        pass
    return None


def case_entry(case: dict[str, Any], markers: list[str] | None,
               patterns: dict[str, list]) -> dict[str, Any]:
    """一个案的判定与证据。库里的行要有本案日志里框架的收尾作证；pass 再扫一遍执行失败回显
    （markers 为 None 时不扫）。"""
    autoid = str(case["case_id"])
    recorded = str(case.get("result") or "").lower()
    log = "" if case.get("log_stale") else str(case.get("log") or "")
    entry: dict[str, Any] = {"autoid": autoid}
    reason = None
    if recorded in ("pass", "fail"):
        closing = framework_verdict(log, autoid)
        if closing == recorded:
            verdict = recorded
        elif closing is None:
            verdict = "broken"
            reason = (f"结果库记的是 {recorded}，但本案日志里没有框架的收尾（PASS/FAIL 横幅紧跟 end case）："
                      "这一轮停在了本案里（超时被杀、崩溃或收尾时登不上设备），库里那行是本案开跑时"
                      "写下的占位值（上一个案的结果），不是本案的判定")
        else:
            verdict = "broken"
            reason = f"结果库记的是 {recorded}，本案日志里框架收尾判的是 {closing}，两边对不上"
    else:
        verdict = "not_run"
    if verdict == "pass" and markers:
        echo = failure_echo_lines(log, markers)
        expected = [(str(p[0]), str(p[1])) for p in patterns.get(autoid, [])
                    if isinstance(p, (list, tuple)) and len(p) == 2]
        if echo and not failure_echo_expected(echo, expected, log):
            verdict = "broken"
            reason = ("框架判 pass，但本案日志里有执行失败回显，而本案的断言并不在等它："
                      "多半某条配置没生效、断言空真通过")
            entry["failure_echo"] = echo
    entry["verdict"] = verdict
    if recorded and recorded != verdict:
        entry["recorded_result"] = recorded
    if reason:
        entry["broken_reason"] = reason
    if verdict != "pass":
        entry["detail_tail"] = log
        entry["failed_checks"] = failed_checks(log)
        if case.get("log_stale"):
            entry["note"] = "框架日志早于本次投递（上一轮残留），不作本次证据"
        entry["attribution"] = attribute_fail(log)
    return entry


_SESSION_NAME_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.:-]{0,120}\.txt$")


def write_sessions(ws: Workspace, batch_dir: Path, task_id: str, autoid: str,
                   sessions: Any) -> dict[str, str]:
    """网关带回的会话转储（设备 CLI 会话 apv_*.txt、触发机会话）写到
    <batch>/evidence/<task_id>/<autoid>/，返回 名字 → 工作区相对路径。名字不像文件名的不写。"""
    written: dict[str, str] = {}
    if not isinstance(sessions, dict):
        return written
    for name, text in sessions.items():
        if not _SESSION_NAME_RE.match(str(name)) or not isinstance(text, str):
            continue
        path = batch_dir / EVIDENCE_DIR / task_id / autoid / str(name)
        write_file_safely(ws.root, path, text.encode("utf-8"))
        written[str(name)] = path.relative_to(ws.root).as_posix()
    return written


def results(ws: Workspace, task_id: str) -> dict[str, Any]:
    """取结果；运行结束且结果通道就绪时，把回执写回工作簿目录。runner 死了的那一轮没有判定，
    不写回执。"""
    out = gateway.call_tool(ws, "case_results", {"task_id": task_id})
    if not out.get("ok") or out.get("channel") in ("not_completed", "runner_lost"):
        return out
    record = gateway.task_record(ws, task_id) or {}
    markers = failure_markers(ws)
    patterns = _submitted_patterns(ws, record)
    scan_note = None
    if markers is None:
        scan_note = ("数据包里读不到执行失败回显规则（domain_grammar.exec_failure_markers），"
                     "pass 案没有扫空真")
    elif patterns is None:
        scan_note = "不知道这次投递的断言（工作簿投递后改过），pass 案没有扫空真"
    batch_dir = (ws.root / record["xlsx"]).parent if record.get("xlsx") else None
    cases = []
    for case in out.get("cases") or []:
        entry = case_entry(case, None if scan_note else markers, patterns or {})
        if batch_dir is not None and case.get("sessions"):
            entry["sessions"] = write_sessions(ws, batch_dir, task_id, entry["autoid"],
                                               case["sessions"])
        cases.append(entry)
    verdicts = [c["verdict"] for c in cases]
    totals = {"cases": len(verdicts), "pass": verdicts.count("pass"),
              "fail": verdicts.count("fail"), "broken": verdicts.count("broken"),
              "not_run": sum(1 for v in verdicts if v not in ("pass", "fail", "broken"))}
    written = None
    note = None
    newer = gateway.newer_submissions(ws, task_id) if record.get("xlsx") else []
    if record.get("xlsx") and newer:
        note = (f"{record['xlsx']} was submitted again after this run (task {newer[-1]}); these "
                "are the older run's results, so the batch receipt was not overwritten")
    elif record.get("xlsx"):
        xlsx = ws.root / record["xlsx"]
        result = {
            "schema": RESULT_SCHEMA, "xlsx": record["xlsx"], "xlsx_sha256": out.get("xlsx_sha256"),
            "batch": xlsx.parent.name, "task_id": task_id, "result_channel": out.get("channel"),
            "rc": out.get("rc"), "run_dir": out.get("run_dir"),
            "submit_autoid": out.get("submit_autoid") or record.get("submit_autoid"),
            "module": out.get("module") or record.get("module"),
            "report_dir": out.get("report_dir"),
            "sentinel": record.get("sentinel"),
            "failure_echo_scan": scan_note or "on",
            "submitted": record.get("submitted_at"), "finished": _now(), "cases": cases,
            "totals": totals,
        }
        if record.get("provenance_fp"):
            result["provenance_fingerprints"] = record["provenance_fp"]
        if record.get("case_fingerprints"):
            result["case_fingerprints"] = record["case_fingerprints"]
        write_receipts(result, xlsx.parent, ws.root)
        written = (xlsx.parent / "run_receipt.md").relative_to(ws.root).as_posix()
    # 整段日志只进 run_results.json；回给会话的只带失败断言和一截尾巴，免得几段日志挤占上下文
    compact = [{**case, "detail_tail": case["detail_tail"][-_RETURN_TAIL_CHARS:]}
               if case.get("detail_tail") else case for case in cases]
    reply = {**out, "totals": totals, "cases": compact, "receipt": written}
    if note:
        reply["note"] = note
    return reply


def write_receipts(result: dict[str, Any], out_dir: Path, root: Path | None = None) -> None:
    """run_results.json + run_receipt.md；root（工作区根）之下原子写、不跟随符号链接。"""
    root = Path(root) if root is not None else Path(out_dir)

    def put(name: str, text: str) -> None:
        write_file_safely(root, Path(out_dir) / name, text.encode("utf-8"))

    put("run_results.json", json.dumps(result, ensure_ascii=False, indent=1))
    t = result["totals"]
    rc = result.get("rc")
    lines = [
        "# 上机回执（run_receipt）", "",
        f"- batch: `{result['batch']}`",
        f"- xlsx: `{result['xlsx']}` (sha256 `{str(result.get('xlsx_sha256') or '')[:16]}…`)",
        f"- task_id: `{result.get('task_id', '')}`",
        f"- 时间: {result.get('submitted')} → {result['finished']}",
        f"- 总数: {t['cases']} · pass {t['pass']} · fail {t['fail']} · broken {t.get('broken', 0)}"
        f" · not_run {t['not_run']}",
        f"- result channel: {result.get('result_channel')} · 框架进程退出码 rc: {rc}",
    ]
    if result.get("submit_autoid") or result.get("run_dir"):
        lines.append(f"- 跳板机上：落位 `ist_staging_{result.get('module') or '?'}/"
                     f"{result.get('submit_autoid') or '?'}`，报告目录 "
                     f"`{result.get('report_dir') or ('report/' + str(result.get('run_dir')))}`")
    if rc not in (0, None):
        lines.append(f"- 注意：框架进程退出码 {rc} 不是 0，这一轮没有正常跑完（崩溃或被杀）："
                     "没有收尾的案判 broken，没开跑的案是 not_run")
    if result.get("failure_echo_scan") not in (None, "on"):
        lines.append(f"- 注意：{result['failure_echo_scan']}")
    if result.get("sentinel"):
        lines.append("- 卷尾的 `999999999999999` 是出件时垫的哨兵案，不计入总数；原始日志末尾它的 "
                     "FAIL 横幅不代表任何用例失败")
    lines += ["", "| autoid | 判定 | 证据摘录 |", "|---|---|---|"]
    for case in result["cases"]:
        checks = case.get("failed_checks") or []
        if case.get("broken_reason"):
            echo = (case.get("failure_echo") or [""])[0]
            evidence = case["broken_reason"] + (f" ⏎ {echo}" if echo else "")
        else:
            evidence = check_summary(checks[0]) if checks else (case.get("detail_tail") or "")[-160:]
        note = (evidence or case.get("note") or "").replace("\n", " ⏎ ")[:240]
        layer = (case.get("attribution") or {}).get("layer")
        if layer:
            note = f"[{layer}] {note}"
        if case.get("sessions"):
            note += " · 会话转储：" + "、".join(f"`{p}`" for p in case["sessions"].values())
        lines.append(f"| {case['autoid']} | {case['verdict']} | {note or '-'} |")
    put("run_receipt.md", "\n".join(lines) + "\n")


def run_and_wait(ws: Workspace, xlsx: str, *, module: str | None = None, poll_s: float = 10,
                 max_s: float = 2400, heartbeat_s: float = 300) -> dict[str, Any]:
    """提交 → 轮询到结束（期间续租约）→ 取结果写回执。给 run_device.py 这类一次跑完的入口用。"""
    submitted = submit(ws, xlsx, module)
    if not submitted.get("ok"):
        return submitted
    task_id = submitted["task_id"]
    deadline = time.time() + max_s
    last_beat = time.time()
    while time.time() < deadline:
        state = status(ws, task_id)
        if state.get("state") == "done":
            return {**results(ws, task_id), "task_id": task_id}
        if state.get("state") == "lost":
            return {"ok": False, "task_id": task_id, "state": "lost", "error": RUNNER_LOST}
        if time.time() - last_beat > heartbeat_s:
            gateway.lease(ws, "heartbeat")
            last_beat = time.time()
        time.sleep(poll_s)
    return {"ok": False, "task_id": task_id,
            "error": f"run did not finish within {int(max_s)}s; poll cex_case_status later"}
