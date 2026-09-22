#!/usr/bin/env python3
"""compile_excel: 从用例步骤 JSON 产出结构正确的 case.xlsx。

分工：结构由本脚本保证（执行页定位、表头、列语义、契约 marker、原子写盘），
内容由调用方（agent）决定。对应 InfoTest emit_xlsx_tool.compile_emit 的
"结构由工具保证，内容由 agent 决定"哲学（InfoTest_Engine 同步来源：
main/case_compiler/xlsx_emit.py + excel_contract.py）。

用法：
    python compile_excel.py --cases cases.json [--out DIR] [--template PATH]

cases.json 契约：
    {
      "batch": "批次名",               # 必填，决定输出目录名
      "init_commands": ["...","..."],  # 可选，文件级前置命令
      "cases": [
        {
          "autoid": "自动化ID",
          "priority": "P1",            # 可选
          "stmt_type": "功能",          # 可选
          "description": "用例描述",
          "steps": [
            {"e":"被测设备","f":"cmd_config","g":"show version",
             "h":"","i":""}
          ]
        }
      ]
    }
列语义见 reference/column-semantics.md。
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

# 与 InfoTest case_compiler/excel_contract.py 的 EXECUTION_HEADERS 保持一致
EXECUTION_HEADERS = (
    "自动化ID",
    "优先级",
    "语句类型",
    "描述",
    "测试对象",
    "方法",
    "数据",
    "临时保存期望结果",
    "输入变量",
)

EXECUTION_SHEET_MARKER = "IST_EXECUTION_SHEET"
CONTRACT_MARKER = "IST_EXCEL_CONTRACT"

# 与 emit_xlsx_tool 认可的 F 列方法族保持一致；收紧而不是放开
ALLOWED_METHODS = frozenset({
    "cmd_config", "cmds_config", "found", "not_found", "found_times",
    "事实源主机名", "sleep",
})

DEFAULT_TEMPLATE = Path(__file__).resolve().parent.parent / "templates" / "case_template.v2.xlsx"

# E=操作对象  F=方法  G=数据  H=临时保存期望结果  I=输入变量
STEP_KEYS = ("e", "f", "g", "h", "i")


class CompileError(Exception):
    """产出前的所有校验失败都走这个异常，禁止带病出盘。"""


def _validate_cases(doc: dict) -> list[dict]:
    if not isinstance(doc, dict):
        raise CompileError("cases JSON 顶层必须是对象")
    batch = str(doc.get("batch") or "").strip()
    if not batch:
        raise CompileError("缺少 batch（批次名）")
    if "/" in batch or batch.startswith("."):
        raise CompileError(f"batch 名不安全: {batch!r}")
    cases = doc.get("cases")
    if not isinstance(cases, list) or not cases:
        raise CompileError("cases 必须是非空数组")

    for idx, case in enumerate(cases):
        label = case.get("autoid") or f"#{idx + 1}"
        for key in ("autoid", "description"):
            if not str(case.get(key) or "").strip():
                raise CompileError(f"用例 {label}: 缺少 {key}")
        steps = case.get("steps")
        if not isinstance(steps, list) or not steps:
            raise CompileError(f"用例 {label}: steps 必须是非空数组")
        for sidx, step in enumerate(steps):
            method = str(step.get("f") or "").strip()
            if method not in ALLOWED_METHODS:
                raise CompileError(
                    f"用例 {label} 第 {sidx + 1} 步: 方法 {method!r} 不在允许集 "
                    f"{sorted(ALLOWED_METHODS)} 内"
                )
            for key in STEP_KEYS:
                if key not in step:
                    raise CompileError(f"用例 {label} 第 {sidx + 1} 步: 缺列 {key}")
    return cases


def _locate_execution_sheet(wb) -> object:
    """按 IST_EXECUTION_SHEET marker 定位执行页，找不到则拒绝。"""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value == EXECUTION_SHEET_MARKER:
                    return ws
    raise CompileError(
        f"模板中没有 {EXECUTION_SHEET_MARKER} marker，无法定位执行页——"
        "拒绝回退猜测（空真结果风险）"
    )


def _check_contract(wb) -> None:
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value == CONTRACT_MARKER:
                    return
    raise CompileError(
        f"模板中没有 {CONTRACT_MARKER} marker（markerless legacy 一律拒绝）"
    )


def _next_empty_row(ws, header_row: int) -> int:
    row = header_row + 1
    while any(ws.cell(row=row, column=c).value not in (None, "") for c in range(1, 10)):
        row += 1
    return row


def compile_excel(cases_path: str, out_dir: str, template: Path) -> dict:
    doc = json.loads(Path(cases_path).read_text(encoding="utf-8"))
    cases = _validate_cases(doc)
    batch = str(doc["batch"]).strip()

    if not template.is_file():
        raise CompileError(f"模板不存在: {template}")
    wb = load_workbook(template)
    _check_contract(wb)
    ws = _locate_execution_sheet(wb)

    # 表头必须在第 1 行且逐字一致，防止错位模板带病通过
    headers = tuple(str(ws.cell(row=1, column=c + 1).value or "") for c in range(9))
    if headers != EXECUTION_HEADERS:
        raise CompileError(f"模板表头与契约不一致: {headers}")

    cursor = _next_empty_row(ws, 1)
    init_commands = [str(c) for c in (doc.get("init_commands") or [])]
    total_steps = 0
    for case in cases:
        autoid = str(case["autoid"]).strip()
        shared = (
            autoid,
            str(case.get("priority") or ""),
            str(case.get("stmt_type") or ""),
            str(case["description"]).strip(),
        )
        for step in case["steps"]:
            for col, key in enumerate(STEP_KEYS, start=5):
                ws.cell(row=cursor, column=col, value=str(step.get(key) or ""))
            for col, value in enumerate(shared, start=1):
                ws.cell(row=cursor, column=col, value=value)
            cursor += 1
            total_steps += 1
        # A-D 列在步骤行重复填充，保证框架按行扫描时每行自洽

    out_root = Path(out_dir)
    out_dir_final = out_root / batch
    out_dir_final.mkdir(parents=True, exist_ok=True)
    target = out_dir_final / "case.xlsx"

    # 原子写盘：tmp + rename，避免半成品文件被误读
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=str(out_dir_final))
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, target)
    except BaseException:
        Path(tmp).unlink(missing_ok=True)
        raise

    return {
        "ok": True,
        "batch": batch,
        "cases": len(cases),
        "steps": total_steps,
        "init_commands": len(init_commands),
        "output": str(target),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="compile_excel: 用例 JSON → case.xlsx")
    parser.add_argument("--cases", required=True, help="cases JSON 路径")
    parser.add_argument("--out", default="compile_outputs", help="产物根目录")
    parser.add_argument("--template", default=str(DEFAULT_TEMPLATE), help="契约模板路径")
    args = parser.parse_args()

    try:
        result = compile_excel(args.cases, args.out, Path(args.template))
    except CompileError as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False))
        return 1
    except FileNotFoundError as exc:
        print(json.dumps({"ok": False, "error": f"文件不存在: {exc}"}, ensure_ascii=False))
        return 1

    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
